#!/usr/bin/env python3
"""make_layouts.py -- a corpus of REALISTIC statement LAYOUTS, grouped into BANK
FAMILIES, each file paired with the ground truth it was drawn from.

WHAT THIS IS FOR, and how it differs from make_corpus.py.

make_corpus.py is adversarial: ONE layout (the shipped ANZ bands) pushed until it
breaks. This corpus is the opposite: ~40 plausible statement DESIGNS per split,
organised as 8 bank families (ANZ, ASB, BNZ, Westpac, Kiwibank, TSB, Co-operative
Bank, and the fictional Rimu Bank), 4-6 designs per bank (everyday, savings, credit
card, business, loan), 3-5 statements per design (different months, page counts and
transaction mixes) -- a pile of each bank's statements, so a tool that learns a
bank's layouts from its own pile can be measured. It was written WITHOUT reading any
column-finding code, and it is not tuned against one.

REAL BANK NAMES, FAKE EVERYTHING ELSE. The bank names are real so bank
identification can be tested; there are no logos and no brand colours, every page
says "SYNTHETIC TEST DOCUMENT - NOT A REAL STATEMENT", and the people, addresses,
account digits, card digits, phone numbers and figures are invented. Account
numbers have the NZ shape BB-bbbb-AAAAAAA-SSS with the issuing bank's real 2-digit
code (ANZ 01/06, BNZ 02, Westpac 03, ASB 12, TSB 15, Kiwibank 38, Rimu 99). The
Co-operative Bank is printed with 02, ASSUMED to share BNZ's code, which makes its
statements the ones where the masthead AND the code are both needed.

TWO SPLITS. `dev` is for looking at; `holdout` is for scoring once. Different seeds,
and every design draws its positions, widths, fonts, sizes, heading wordings and
date formats from a disjoint holdout menu; each split also has designs the other
never sees. Within a split, all statements of one design share one geometry.

WHAT IS WRITTEN, per split directory:
    <case>.pdf        + <case>.truth.json     ~130 text-layer statements
    <case>_scan.pdf   + <case>_scan.truth.json ~15 image-only rasterised copies
    <case>.csv / .xlsx + <case>.truth.json     ~12 internet-banking exports
    index.json                                 one entry per truth file

THE TRUTH FILE is make_corpus.py's format, plus keys:
    case, generator, note, opening_balance, closing_balance, row_count,
    rows: [{date "YYYY-MM-DD", description, debit, credit, balance, ...}],
    bank, layout, product, account_bank_code, account_number, account_redaction,
    features, row_order ("oldest_first" | "newest_first"), removed_rows,
    source_format ("pdf" | "scan" | "csv" | "xlsx")
    accounts   (several accounts in one statement; rows carry account_index;
                the top-level opening/closing are then null)
    statements (several statements bundled in one PDF; rows carry statement_index;
                the balances chain from one statement into the next)
  debit = money OUT and credit = money IN, both positive, from the ACCOUNT HOLDER'S
  side: on a card a purchase is a debit and a payment a credit, whatever sign the
  card prints, and a card's opening/closing balance is the amount owed NEGATED. A
  zero-value row ("FEE WAIVED 0.00") is debit 0.0.
  balance = the running balance printed ON THAT ROW, else null.
  description = every text cell that is not a date or a figure, in reading order
  (line by line, left to right), joined by ONE space. Foreign-currency code and
  amount columns are figures, not description.
  date = the transaction date column (a card's processed date or an account's value
  date is printed beside it and is not the truth date).
  Opening, brought/carried-forward, page-total, sub-total, cardholder, account and
  closing lines are NOT rows. Rows are in PRINTED order.
  REDACTIONS. A row may carry redacted: [...] -- the text was removed and a black
  box drawn where it was, so the value is gone and is null in the truth -- or
  overlay_redacted: [...] -- a black box drawn OVER text still in the text layer,
  which the policy ("if it is readable, we use it") keeps. A whole removed row is
  not in the truth; removed_rows counts them. A scan cannot read under a box, so
  in a _scan truth every overlay becomes a removal.

EVERY FILE IS CHECKED BEFORE ITS TRUTH IS WRITTEN, and the run aborts loudly on the
first failure rather than write a truth that lies:
  * arithmetic: per account, in date order, opening - debits + credits = every
    balance and the closing, both on the generator's own data and on the truth as
    written (through every redaction it can see through);
  * every printed date parses back to exactly the truth date, given the period;
  * every printed figure parses back (sign, CR/DR/OD, parentheses, trailing minus,
    $, thousands) to the truth amount or balance;
  * every printed description joins to the truth description;
  * every cell lies inside its column, no two strings collide, nothing is off the
    page; a scan has no text layer; an export reads back to its truth.

PYTHON 3.9+, reportlab; pymupdf + numpy + Pillow for the scans; openpyxl for the
.xlsx exports (skipped, with a message, if it is missing). Byte-identical output
across runs: crc32 seeds, reportlab invariant mode, fixed PDF and zip metadata.

Run:  python3 tools/synth/make_layouts.py --out /tmp/zoo/dev --split dev
      python3 tools/synth/make_layouts.py --out /tmp/zoo/holdout --split holdout
      python3 tools/synth/make_layouts.py --split dev --list
      python3 tools/synth/make_layouts.py --out /tmp/zoo/dev --split dev --only anz_

Dev-time only. Nothing here ships to the server.
"""

import argparse
import calendar
import csv
import datetime as dt
import io
import json
import os
import random
import re
import sys
import zipfile
import zlib

MIN_PYTHON = (3, 9)
if sys.version_info < MIN_PYTHON:
    sys.exit("make_layouts.py needs Python %d.%d or newer (this is %s)"
             % (MIN_PYTHON[0], MIN_PYTHON[1], sys.version.split()[0]))

try:
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import A4, letter, landscape
    from reportlab.pdfbase.pdfmetrics import stringWidth
except ImportError as e:            # noqa: BLE001 -- the message IS the handling
    sys.exit("make_layouts.py needs reportlab (%s).\n"
             "  python3 -m pip install reportlab pymupdf numpy pillow openpyxl\n"
             "Nothing in the app needs it: this is a dev-time tool." % e)

GENERATOR = "tools/synth/make_layouts.py"
SPLIT_SEED = {"dev": 20260117, "holdout": 77031}
SYNTHETIC = "SYNTHETIC TEST DOCUMENT - NOT A REAL STATEMENT"

PAGES = {"A4": A4, "A4L": landscape(A4), "Letter": letter}
FONTS = {"Helvetica": ("Helvetica", "Helvetica-Bold"),
         "Times": ("Times-Roman", "Times-Bold"),
         "Courier": ("Courier", "Courier-Bold")}
MONEY_KINDS = ("debit", "credit", "amount", "balance")
TEXT_KINDS = ("type", "desc", "payee", "part", "code", "ref")
FX_KINDS = ("fxcur", "fxamt")
PAD_T = 2.0          # a text cell starts this far inside its column
PAD_M = 3.0          # a right-aligned figure ends this far inside its column

MON = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
       "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
MONTH = ["January", "February", "March", "April", "May", "June", "July",
         "August", "September", "October", "November", "December"]
DOW = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]

# Real banks are drawn in neutral greys only: names for identification, never
# trade dress. Colour is kept for the fictional bank.
DARK = (0.22, 0.24, 0.28)
MID = (0.42, 0.44, 0.48)
LIGHT = (0.86, 0.86, 0.86)
TEAL = (0.0, 0.40, 0.44)
GREEN = (0.06, 0.38, 0.24)
BLACK = (0.0, 0.0, 0.0)
WHITE = (1.0, 1.0, 1.0)
GREY = (0.90, 0.90, 0.90)
PALE = (0.95, 0.95, 0.95)

BANKS = {
    "anz": dict(name="ANZ", mast="ANZ", legal="ANZ Bank New Zealand Limited",
                codes=("01", "06")),
    "asb": dict(name="ASB", mast="ASB", legal="ASB Bank Limited", codes=("12",)),
    "bnz": dict(name="BNZ", mast="BNZ", legal="Bank of New Zealand", codes=("02",)),
    "westpac": dict(name="Westpac", mast="Westpac", legal="Westpac New Zealand Limited",
                    codes=("03",)),
    "kiwibank": dict(name="Kiwibank", mast="Kiwibank", legal="Kiwibank Limited",
                     codes=("38",)),
    "tsb": dict(name="TSB", mast="TSB", legal="TSB Bank Limited", codes=("15",)),
    "coop": dict(name="Co-operative Bank", mast="The Co-operative Bank",
                 legal="The Co-operative Bank Limited", codes=("02",)),
    "rimu": dict(name="Rimu Bank", mast="Rimu Bank", legal="Rimu Bank of Aotearoa Limited",
                 codes=("99",), fictional=True),
}
OTHER_CODES = ["01", "02", "03", "06", "12", "15", "38"]


class GenError(Exception):
    """A self-check failed. Never caught quietly: a truth file that disagrees with
    its own document is worse than no corpus, because it accuses the reader."""


# ---------------------------------------------------------------------------
# Dates and money, printed the ways NZ statements and exports print them.
# ---------------------------------------------------------------------------

DATE_FMTS = {
    "dd Mon":       lambda d: "%02d %s" % (d.day, MON[d.month - 1]),
    "d Mon":        lambda d: "%d %s" % (d.day, MON[d.month - 1]),
    "dd MON":       lambda d: "%02d %s" % (d.day, MON[d.month - 1].upper()),
    "Dow dd Mon":   lambda d: "%s %02d %s" % (DOW[d.weekday()], d.day, MON[d.month - 1]),
    "Dow d Mon":    lambda d: "%s %d %s" % (DOW[d.weekday()], d.day, MON[d.month - 1]),
    "d Mon yyyy":   lambda d: "%d %s %d" % (d.day, MON[d.month - 1], d.year),
    "dd Mon yyyy":  lambda d: "%02d %s %d" % (d.day, MON[d.month - 1], d.year),
    "dd Mon yy":    lambda d: "%02d %s %02d" % (d.day, MON[d.month - 1], d.year % 100),
    "dd MON yy":    lambda d: "%02d %s %02d" % (d.day, MON[d.month - 1].upper(), d.year % 100),
    "d Month yyyy": lambda d: "%d %s %d" % (d.day, MONTH[d.month - 1], d.year),
    "dd/mm/yyyy":   lambda d: "%02d/%02d/%d" % (d.day, d.month, d.year),
    "d/mm/yyyy":    lambda d: "%d/%02d/%d" % (d.day, d.month, d.year),
    "dd/mm/yy":     lambda d: "%02d/%02d/%02d" % (d.day, d.month, d.year % 100),
    "dd-mm-yy":     lambda d: "%02d-%02d-%02d" % (d.day, d.month, d.year % 100),
    "dd-mm-yyyy":   lambda d: "%02d-%02d-%d" % (d.day, d.month, d.year),
    "dd-Mon-yy":    lambda d: "%02d-%s-%02d" % (d.day, MON[d.month - 1], d.year % 100),
    "dd-Mon-yyyy":  lambda d: "%02d-%s-%d" % (d.day, MON[d.month - 1], d.year),
    "yyyy-mm-dd":   lambda d: d.isoformat(),
    "yyyy/mm/dd":   lambda d: "%d/%02d/%02d" % (d.year, d.month, d.day),
    "d/m/yyyy":     lambda d: "%d/%d/%d" % (d.day, d.month, d.year),
    "d/m/yy":       lambda d: "%d/%d/%02d" % (d.day, d.month, d.year % 100),
    "yyyymmdd":     lambda d: "%d%02d%02d" % (d.year, d.month, d.day),
}
# A year-less date is only readable beside a printed statement period, which every
# layout that uses one prints in its page-1 header.
YEARLESS = {"dd Mon", "d Mon", "dd MON", "Dow dd Mon", "Dow d Mon"}


def fmt_date(d, f):
    return DATE_FMTS[f](d)


def period_text(f, start, end):
    """The statement period, always WITH the year."""
    def one(d):
        if f == "long":
            return "%d %s %d" % (d.day, MONTH[d.month - 1], d.year)
        if f == "short":
            return "%d %s %d" % (d.day, MON[d.month - 1], d.year)
        if f == "num":
            return "%02d/%02d/%d" % (d.day, d.month, d.year)
        if f == "upper":
            return "%02d %s %d" % (d.day, MON[d.month - 1].upper(), d.year)
        raise GenError("unknown period format %r" % f)
    return one(start), one(end)


def add_months(d, k):
    m0 = d.month - 1 + k
    y, m = d.year + m0 // 12, m0 % 12 + 1
    return dt.date(y, m, min(d.day, calendar.monthrange(y, m)[1]))


def daterange(a, b):
    d = a
    while d <= b:
        yield d
        d += dt.timedelta(days=1)


def mag(c, thousands=True, dollar=False):
    """|c| cents as printed: 1,234.56. Integer arithmetic, never float formatting."""
    d, r = divmod(abs(int(c)), 100)
    s = ("{:,}".format(d) if thousands else str(d)) + ".%02d" % r
    return "$" + s if dollar else s


def fmt_money(v, style, L):
    """A signed figure (holder's side: negative = money out, or overdrawn) as the
    page prints it. Returns (number, token); token is "" or CR / DR / OD.

    THE SIGN MUST SURVIVE THIS FUNCTION. make_corpus.py's money() once printed
    abs(x) and blamed the reader for 221 wrong balances; every style here is
    parsed back by parse_money() before a truth file is written."""
    m = mag(v, L["thousands"], L["dollar"])
    neg = v < 0
    if style == "plain":
        if v < 0:
            raise GenError("plain style asked to print a negative figure")
        return m, ""
    if style == "lead":
        return ("-" + m if neg else m), ""
    if style == "plus":
        return (("-" if neg else "+") + m), ""
    if style == "trail":
        return (m + "-" if neg else m), ""
    if style == "paren":
        return ("(" + m + ")" if neg else m), ""
    if style == "drcr":
        return m, ("DR" if neg else "CR")
    if style == "od":
        return m, ("OD" if neg else "")
    if style == "dr_only":
        return m, ("DR" if neg else "")
    # Credit-card conventions: printed from the CARD's side, so a purchase (money
    # out for the holder) is a plain figure and a payment carries the mark.
    if style == "card_cr":
        return m, ("" if neg else "CR")
    if style == "card_minus":
        return (m if neg else "-" + m), ""
    raise GenError("unknown money style %r" % style)


def joined(num, tok):
    return num + (" " + tok if tok else "")


def parse_money(num, tok, card=False):
    """Read a printed figure back, independently of how fmt_money built it."""
    s = num
    neg = False
    if s.startswith("(") and s.endswith(")"):
        neg, s = True, s[1:-1]
    if s.startswith("-"):
        neg, s = True, s[1:]
    elif s.startswith("+"):
        s = s[1:]
    if s.endswith("-"):
        neg, s = True, s[:-1]
    if s.startswith("$"):
        s = s[1:]
    if not re.fullmatch(r"\d{1,3}(,\d{3})*\.\d\d|\d+\.\d\d", s):
        raise GenError("printed figure %r does not parse" % num)
    c = int(s.replace(",", "").replace(".", ""))
    if tok not in ("", "CR", "DR", "OD"):
        raise GenError("unknown token %r" % tok)
    if card:
        return c if (neg or tok == "CR") else -c
    if tok in ("DR", "OD"):
        neg = True
    return -c if neg else c


def cents(x):
    return None if x is None else int(round(x * 100))


# ---------------------------------------------------------------------------
# Drawing, in TOP-DOWN coordinates, with the collision check built in.
# ---------------------------------------------------------------------------

class Sheet:
    """One PDF being drawn. Every string drawn is recorded, and at the end of each
    page the strings are checked for collisions and for leaving the page: a corpus
    whose text overlaps measures the overlap, not the reader."""

    def __init__(self, path, pagesize, family):
        self.c = canvas.Canvas(path, pagesize=pagesize, invariant=1)
        self.w, self.h = pagesize
        self.reg, self.bold = FONTS[family]
        self.dx = 0.0
        self.stack = []
        self.items = []
        self.pageno = 0

    def width(self, s, size, bold=False):
        return stringWidth(s, self.bold if bold else self.reg, size)

    def begin_page(self, dx=0.0):
        self.pageno += 1
        self.items = []
        self.dx = dx
        self.stack = []
        if dx:
            self.c.translate(dx, 0)

    def push_dx(self, d):
        """Shift what follows (the table) without moving the header or footer."""
        self.c.saveState()
        if d:
            self.c.translate(d, 0)
        self.dx += d
        self.stack.append(d)

    def pop_dx(self):
        d = self.stack.pop()
        self.c.restoreState()
        self.dx -= d

    def place(self, x, s, size, bold=False, align="left"):
        w = self.width(s, size, bold)
        x0 = x - w if align == "right" else (x - w / 2.0 if align == "center" else x)
        return x0, x0 + w

    def text(self, x, y, s, size, bold=False, align="left", color=None):
        if s == "":
            return (x, x)
        fn = self.bold if bold else self.reg
        x0, x1 = self.place(x, s, size, bold, align)
        self.c.setFont(fn, size)
        if color is not None:
            self.c.setFillColorRGB(*color)
        self.c.drawString(x0, self.h - y, s)
        if color is not None:
            self.c.setFillColorRGB(0, 0, 0)
        self.items.append((y, x0 + self.dx, x1 + self.dx, size, s))
        return (x0, x1)

    def blackout(self, x0, x1, y, size):
        """A redaction box over the glyph box of a string at baseline y."""
        self.rect(x0 - 1.5, y - 0.85 * size, (x1 - x0) + 3.0, 1.15 * size, fill=BLACK)

    def rect(self, x, y_top, w, h, fill=None, stroke=None, width=0.6):
        if fill is not None:
            self.c.setFillColorRGB(*fill)
        if stroke is not None:
            self.c.setStrokeColorRGB(*stroke)
            self.c.setLineWidth(width)
        self.c.rect(x, self.h - y_top - h, w, h,
                    fill=1 if fill is not None else 0,
                    stroke=1 if stroke is not None else 0)
        self.c.setFillColorRGB(0, 0, 0)
        self.c.setStrokeColorRGB(0, 0, 0)

    def line(self, x0, y0, x1, y1, width=0.6, color=None):
        self.c.setLineWidth(width)
        if color is not None:
            self.c.setStrokeColorRGB(*color)
        self.c.line(x0, self.h - y0, x1, self.h - y1)
        self.c.setStrokeColorRGB(0, 0, 0)

    def end_page(self):
        if self.stack:
            raise GenError("unbalanced table shift")
        self.check()
        self.c.showPage()

    def check(self):
        its = sorted(self.items)
        for a in its:
            if a[1] < 4 or a[2] > self.w - 4 or a[0] < 8 or a[0] > self.h - 4:
                raise GenError("page %d: %r drawn off the page (x %.1f-%.1f, y %.1f)"
                               % (self.pageno, a[4], a[1], a[2], a[0]))
        # Same baseline: at least 2.5pt apart, or a text extractor runs the two
        # strings together into one word. Different baselines: the glyph boxes
        # (cap height 0.75em above the baseline, descenders 0.22em below) must not
        # overlap where the strings overlap horizontally.
        for i, a in enumerate(its):
            for b in its[i + 1:]:
                dy = b[0] - a[0]
                if dy > 30:
                    break
                if dy < 0.01:
                    hit = b[1] < a[2] + 2.5 and a[1] < b[2] + 2.5
                else:
                    vert = (b[0] - 0.75 * b[3]) < (a[0] + 0.22 * a[3])
                    hit = vert and b[1] < a[2] - 0.5 and a[1] < b[2] - 0.5
                if hit:
                    raise GenError("page %d: %r and %r collide (y %.1f / %.1f)"
                                   % (self.pageno, a[4], b[4], a[0], b[0]))

    def save(self):
        self.c.save()


class Col:
    """One table column: a band [x0, x1] and how its cells sit in it."""

    def __init__(self, kind, x0, x1, align, head, head_align):
        self.kind, self.x0, self.x1 = kind, float(x0), float(x1)
        self.align, self.head, self.head_align = align, tuple(head), head_align

    def anchor(self, align=None):
        a = align or self.align
        if a == "right":
            return self.x1 - PAD_M
        if a == "center":
            return (self.x0 + self.x1) / 2.0
        return self.x0 + PAD_T

    def __repr__(self):
        return "%s[%.0f-%.0f]" % (self.kind, self.x0, self.x1)


def mkcols(left, right, spec):
    """Columns laid side by side from `left` to `right`. spec entries are
    (kind, width or None for the one flexible column, heading lines[, align[, head_align]])."""
    fixed = sum(e[1] for e in spec if e[1])
    flex = right - left - fixed
    if flex < 60:
        raise GenError("no room left for the flexible column (%.0f)" % flex)
    out, x = [], float(left)
    for e in spec:
        kind, w, head = e[0], (e[1] or flex), e[2]
        align = e[3] if len(e) > 3 and e[3] else (
            "right" if kind in MONEY_KINDS or kind == "fxamt" else "left")
        hal = e[4] if len(e) > 4 and e[4] else align
        out.append(Col(kind, x, x + w, align, head, hal))
        x += w
    return out


def fit(measure, s, maxw):
    """Shorten s (whole words first) until it fits in maxw points."""
    s = s.strip()
    while s and measure(s) > maxw:
        if " " in s:
            s = s.rsplit(" ", 1)[0].rstrip(" -,@/")
        else:
            s = s[:-1]
    return s


# ---------------------------------------------------------------------------
# What the transactions SAY. Invented merchants, people and numbers; the shapes
# (card numbers, times, foreign amounts, invoice numbers that look like money,
# dates inside descriptions) are the real ones.
# ---------------------------------------------------------------------------

NAMES = {
    "grocer": ["HARBOUR FOODMARKET", "VALLEY FRESH GROCER", "KAURI SUPERSTORE",
               "MAIN ST MARKET", "RIVERSIDE DAIRY", "CORNER FOODS 24/7"],
    "fuel": ["TOTARA FUELS", "SOUTHGATE FUEL STOP 114", "KEA GAS AND SERVICE"],
    "cafe": ["MAIN ST CAFE", "BEAN THERE ESPRESSO", "THE DAILY GRIND"],
    "online": ["STREAMFLIX.COM", "APPSTORE ONLINE", "CLOUDBOX STORAGE", "WEBSHOP INTL"],
    "city": ["AUCKLAND", "WELLINGTON", "CHRISTCHURCH", "HAMILTON", "DUNEDIN",
             "NELSON", "TAURANGA", "NAPIER"],
    "utility": ["KAURI ENERGY", "TASMAN POWER", "CITYGAS LTD", "FIBRENET BROADBAND"],
    "council": ["CITY COUNCIL RATES", "DISTRICT COUNCIL WATER", "REGIONAL TRANSPORT"],
    "landlord": ["EXAMPLE RENTALS LTD", "SAMPLE PROPERTY TRUST", "B SAMPLE"],
    "employer": ["EXAMPLE HOLDINGS LTD", "SAMPLE ENGINEERING", "TEST LOGISTICS NZ",
                 "DEMO HEALTH SERVICES"],
    "insurer": ["SOUTHERN INSURANCE", "HARBOUR LIFE COVER"],
    "telco": ["HARBOUR TELECOM", "TUI MOBILE"],
    "atm": ["QUEEN ST", "LAMBTON QUAY", "COLOMBO ST", "VICTORIA ST", "THE MALL"],
    "person": ["A EXAMPLE", "B SAMPLE", "C TESTER", "D DEMO"],
    "supplier": ["SAMPLE SUPPLIES LTD", "EXAMPLE TIMBER CO", "DEMO FREIGHT LTD",
                 "TEST PACKAGING"],
    "customer": ["EXAMPLE CAFE LTD", "SAMPLE MOTORS", "DEMO PLUMBING", "TEST & CO LTD"],
    "airline": ["SKYWAYS AIR", "TASMAN AIRLINES"],
    "branch": ["SAMPLETOWN", "EXAMPLE CENTRAL", "HILLSIDE"],
}


def T(d, w, amt, *lines, cat=""):
    """A catalogue entry printed as one or two description lines."""
    return {"dir": d, "w": w, "amt": amt, "lines": lines, "cat": cat}


def TY(d, w, amt, typ, det, cat=""):
    """An entry for a layout with a separate Type column."""
    return {"dir": d, "w": w, "amt": amt, "type": typ, "lines": (det,), "cat": cat}


def P(d, w, amt, part, code, ref, payee="", cat=""):
    """An entry for an NZ Particulars / Code / Reference layout."""
    return {"dir": d, "w": w, "amt": amt, "part": part, "code": code, "ref": ref,
            "payee": payee, "lines": ("",), "cat": cat}


def FX(d, w, cur, *lines):
    """A purchase in a foreign currency, for a layout with currency columns."""
    return {"dir": d, "w": w, "amt": ("fx",), "lines": lines, "cat": "", "fx": cur}


ZERO_T = T("D", 0, ("fixed", 0.0), "MONTHLY FEE WAIVED", cat="zero")

BANK = [
    T("D", 6, ("r", 12, 310), "EFTPOS {grocer}"),
    T("D", 3, ("r", 35, 160), "EFTPOS {fuel} {card} {time}"),
    T("D", 3, ("r", 4, 38), "POS W/D {cafe}-{time}"),
    T("D", 1, ("r", 9, 14), "2 x COFFEE {cafe}"),
    T("D", 2, ("fx",), "VISA DEBIT {online} USD {fx}"),
    T("D", 2, ("r", 380, 760), "AP {landlord} RENT REF {ref}"),
    T("D", 2, ("r", 60, 340), "DD {utility} {custno}"),
    T("D", 2, ("r", 90, 1650), "BILL PAYMENT {council} INV {inv}"),
    T("D", 2, ("mult", 20, 400, 20), "ATM WITHDRAWAL {atm} {time}"),
    T("D", 1, ("fixed", 5.00), "MONTHLY ACCOUNT FEE", cat="fee"),
    T("D", 1, ("r", 0.25, 4.5), "OVERSEAS TXN FEE", cat="fee"),
    T("D", 2, ("r", 50, 2500), "TFR TO {acct}"),
    T("D", 1, ("r", 28, 260), "DD {insurer} POLICY {invno}"),
    T("D", 1, ("r", 3, 24), "PARKING {ddmm} {time} {city}"),
    T("D", 1, ("r", 150, 900), "LOAN REPAYMENT {n} OF 52"),
    T("D", 2, ("r", 35, 190), "ONLINE PAYMENT {telco} REF {ref}"),
    T("D", 4, ("r", 12, 310), "EFTPOS PURCHASE", "{grocer} {card}"),
    T("D", 2, ("r", 380, 1600), "AUTOMATIC PAYMENT", "{landlord} RENT {n} WKS"),
    T("D", 2, ("fx",), "VISA PURCHASE", "{online} USD {fx}"),
    T("D", 2, ("r", 60, 340), "DIRECT DEBIT", "{utility} {custno}"),
    T("D", 2, ("r", 90, 1650), "BILL PAYMENT", "{council} INV {inv}"),
    T("D", 1, ("mult", 20, 400, 20), "ATM WITHDRAWAL", "{atm} {ddmm} {time}"),
    T("C", 3, ("r", 1800, 6200), "SALARY {employer}", cat="salary"),
    T("C", 2, ("r", 850, 2400), "DIRECT CREDIT {employer} WAGES WK {wk}", cat="wages"),
    T("C", 1, ("r", 0.15, 24), "CREDIT INTEREST", cat="interest"),
    T("C", 2, ("r", 50, 1500), "TFR FROM {acct}"),
    T("C", 1, ("r", 20, 900), "DEPOSIT {branch} BRANCH"),
    T("C", 1, ("r", 5, 180), "REFUND {grocer} INV {inv}"),
    T("C", 1, ("r", 100, 2400), "IRD REFUND {person}", cat="ird"),
    T("C", 3, ("r", 1800, 6200), "DIRECT CREDIT", "{employer} SALARY", cat="salary"),
    T("C", 2, ("r", 50, 2500), "TRANSFER", "FROM {acct}"),
    T("C", 1, ("r", 20, 1800), "DEPOSIT", "MOBILE CHEQUE {chq}"),
    ZERO_T,
]

CARD = [
    T("D", 5, ("r", 12, 320), "{grocer} {city}"),
    T("D", 3, ("r", 35, 160), "{fuel} {city}"),
    T("D", 3, ("r", 4, 40), "{cafe} {city}"),
    T("D", 2, ("r", 6, 120), "{online} {city}"),
    T("D", 1, ("r", 180, 1450), "{airline} {city} TKT {invno}"),
    T("D", 1, ("r", 0.2, 6), "OFFSHORE SERVICE MARGIN", cat="fee"),
    T("D", 3, ("fx",), "{online} SAN FRANCISCO", "USD {fx} @ {rate}"),
    T("D", 2, ("fx",), "{online} SYDNEY", "AUD {fx} @ {arate}"),
    T("D", 1, ("r", 40, 400), "{airline}", "{city} TKT {invno} 2 x ADULT"),
    T("C", 3, ("pay",), "PAYMENT RECEIVED - THANK YOU", cat="payment"),
    T("C", 1, ("pay",), "AUTOMATIC PAYMENT", "THANK YOU", cat="payment"),
    T("C", 1, ("r", 5, 120), "REFUND {online} {city}"),
    T("C", 1, ("r", 5, 120), "REFUND", "{grocer} {city}"),
]

# The same card, for a layout with separate currency columns: the foreign amount
# is not in the description, it is in its own two columns.
CARD_FX = [
    T("D", 5, ("r", 12, 320), "{grocer} {city}"),
    T("D", 3, ("r", 35, 160), "{fuel} {city}"),
    T("D", 3, ("r", 4, 40), "{cafe} {city}"),
    T("D", 1, ("r", 180, 1450), "{airline} {city} TKT {invno}"),
    T("D", 1, ("r", 0.2, 6), "OFFSHORE SERVICE MARGIN", cat="fee"),
    FX("D", 3, "USD", "{online} SAN FRANCISCO"),
    FX("D", 2, "AUD", "{online} SYDNEY"),
    FX("D", 1, "GBP", "{online} LONDON"),
    T("C", 3, ("pay",), "PAYMENT RECEIVED - THANK YOU", cat="payment"),
    T("C", 1, ("r", 5, 120), "REFUND {online} {city}"),
]

BIZ = [
    T("D", 3, ("r", 120, 9800), "AP {supplier} INV {invno}"),
    T("D", 1, ("r", 800, 14000), "IRD GST PERIOD {ddmm}"),
    T("D", 1, ("r", 2500, 18000), "WAGES BATCH {chq}"),
    T("D", 1, ("r", 12, 85), "BANK FEES", cat="fee"),
    T("D", 2, ("r", 40, 180), "EFTPOS {fuel} {card}"),
    T("D", 1, ("r", 600, 4800), "DD {insurer} POLICY {invno}"),
    T("D", 2, ("r", 120, 9800), "AUTOMATIC PAYMENT", "{supplier} INV {inv}"),
    T("D", 1, ("r", 800, 14000), "DIRECT DEBIT", "IRD PAYE {ddmm}"),
    T("C", 4, ("r", 90, 12500), "BILL PAYMENT {customer} INV {invno}"),
    T("C", 2, ("r", 300, 8800), "MERCHANT SETTLEMENT {chq}"),
    T("C", 1, ("r", 100, 3000), "CASH DEPOSIT {branch}"),
    T("C", 2, ("r", 90, 12500), "DIRECT CREDIT", "{customer} INV {inv}"),
    # Very large figures: a property settlement through a business account.
    T("C", 0, ("big", 1000000, 1900000), "SALE PROCEEDS {customer} SETTLEMENT", cat="large"),
    T("D", 0, ("big", 900000, 1700000), "PROPERTY PURCHASE SETTLEMENT {invno}", cat="large"),
    ZERO_T,
]

SAV = [
    T("C", 2, ("r", 0.5, 40), "CREDIT INTEREST", cat="interest"),
    T("C", 2, ("r", 100, 3000), "TFR FROM {acct}"),
    T("D", 1, ("r", 0.1, 9), "RWT ON INTEREST", cat="rwt"),
    T("D", 3, ("r", 100, 1400), "TFR TO {acct}"),
    T("D", 1, ("r", 20, 800), "WITHDRAWAL {branch} BRANCH"),
    T("D", 2, ("r", 100, 1400), "TRANSFER", "TO {acct} {person}"),
    T("C", 1, ("r", 100, 3000), "TRANSFER", "FROM {acct} {person}"),
]

LOAN = [
    T("C", 5, ("r", 1100, 2600), "REPAYMENT THANK YOU"),
    T("C", 1, ("r", 2000, 15000), "LUMP SUM PAYMENT"),
    T("D", 3, ("r", 700, 2200), "INTEREST"),
    T("D", 1, ("fixed", 10.00), "LOAN SERVICE FEE", cat="fee"),
    T("C", 2, ("r", 1100, 2600), "AUTOMATIC PAYMENT", "FROM {acct}"),
    T("D", 1, ("r", 700, 2200), "INTEREST CHARGED", "RATE 6.24% P.A."),
]

TYPED = [
    TY("D", 6, ("r", 12, 310), "EFTPOS", "{grocer} {card}"),
    TY("D", 3, ("r", 35, 160), "EFTPOS", "{fuel} {time}"),
    TY("D", 2, ("r", 4, 38), "POS", "{cafe} {time}"),
    TY("D", 2, ("fx",), "VISA", "{online} USD {fx}"),
    TY("D", 2, ("r", 380, 760), "AP", "{landlord} RENT REF {ref}"),
    TY("D", 2, ("r", 60, 340), "DD", "{utility} {custno}"),
    TY("D", 2, ("r", 90, 1650), "BP", "{council} INV {inv}"),
    TY("D", 2, ("mult", 20, 400, 20), "ATM", "{atm} {time}"),
    TY("D", 1, ("fixed", 5.00), "FEE", "MONTHLY ACCOUNT FEE", cat="fee"),
    TY("D", 2, ("r", 50, 2500), "TFR", "TO {acct}"),
    TY("C", 3, ("r", 1800, 6200), "DC", "{employer} SALARY", cat="salary"),
    TY("C", 1, ("r", 0.15, 24), "INT", "CREDIT INTEREST", cat="interest"),
    TY("C", 2, ("r", 50, 2500), "TFR", "FROM {acct}"),
    TY("C", 1, ("r", 20, 1800), "DEP", "{branch} BRANCH"),
]

TYPED_BIZ = [
    TY("D", 3, ("r", 120, 9800), "AP", "{supplier} INV {invno}"),
    TY("D", 1, ("r", 800, 14000), "DD", "IRD GST PERIOD {ddmm}"),
    TY("D", 1, ("r", 2500, 18000), "BP", "WAGES BATCH {chq}"),
    TY("D", 1, ("r", 12, 85), "FEE", "ACCOUNT FEES", cat="fee"),
    TY("D", 2, ("r", 40, 180), "EFTPOS", "{fuel} {card}"),
    TY("D", 1, ("r", 600, 4800), "DD", "{insurer} POLICY {invno}"),
    TY("C", 4, ("r", 90, 12500), "BP", "{customer} INV {invno}"),
    TY("C", 2, ("r", 300, 8800), "MS", "MERCHANT SETTLEMENT {chq}"),
    TY("C", 1, ("r", 100, 3000), "DEP", "CASH {branch}"),
    TY("C", 2, ("r", 90, 12500), "DC", "{customer} INV {inv}"),
]

PCR = [
    P("D", 5, ("r", 12, 310), "EFTPOS", "{card4}", "{grocer}"),
    P("D", 2, ("r", 380, 760), "{landlord}", "RENT", "WK {wk}"),
    P("D", 2, ("r", 60, 340), "{utility}", "{custno}", "INV {inv}"),
    P("D", 2, ("mult", 20, 400, 20), "ATM", "{atm}", "{time}"),
    P("D", 1, ("fixed", 5.00), "ACCOUNT FEE", "", "", cat="fee"),
    P("D", 2, ("r", 35, 190), "{telco}", "MOBILE", "REF {ref}"),
    P("D", 2, ("fx",), "VISA {online}", "USD {fx}", "{card4}"),
    P("C", 3, ("r", 1800, 6200), "SALARY", "{employer}", "PAY {ddmm}", cat="salary"),
    P("C", 1, ("r", 0.15, 24), "INTEREST", "", "", cat="interest"),
    P("C", 2, ("r", 50, 2500), "TRANSFER", "FROM SAVINGS", "{ref}"),
    P("C", 1, ("r", 5, 180), "REFUND", "{invno}", "{grocer}"),
]

BIZ_PCR = [
    P("D", 3, ("r", 120, 9800), "INV {invno}", "SUPPLIES", "{ref}", payee="{supplier}"),
    P("D", 1, ("r", 800, 14000), "GST", "{ddmm}", "", payee="INLAND REVENUE"),
    P("D", 1, ("r", 2500, 18000), "WAGES", "BATCH {chq}", "", payee="PAYROLL"),
    P("D", 1, ("r", 12, 85), "FEES", "", "", payee="BANK FEES", cat="fee"),
    P("D", 2, ("r", 40, 180), "EFTPOS", "{card4}", "{time}", payee="{fuel}"),
    P("C", 4, ("r", 90, 12500), "INV {invno}", "{ref}", "THANK YOU", payee="{customer}"),
    P("C", 2, ("r", 300, 8800), "SETTLEMENT", "{chq}", "", payee="MERCHANT SVCS"),
    P("C", 1, ("r", 100, 3000), "CASH", "", "{branch}", payee="DEPOSIT"),
]

# A personal account with a payee/other-party field, for exports.
PAYEE_PCR = [
    P("D", 5, ("r", 12, 310), "EFTPOS", "{card4}", "{time}", payee="{grocer}"),
    P("D", 2, ("r", 380, 760), "RENT", "WK {wk}", "{ref}", payee="{landlord}"),
    P("D", 2, ("r", 60, 340), "{custno}", "POWER", "INV {inv}", payee="{utility}"),
    P("D", 2, ("mult", 20, 400, 20), "ATM", "{atm}", "{time}", payee="ATM WITHDRAWAL"),
    P("D", 1, ("fixed", 5.00), "", "", "", payee="ACCOUNT FEE", cat="fee"),
    P("D", 2, ("r", 35, 190), "MOBILE", "", "REF {ref}", payee="{telco}"),
    P("C", 3, ("r", 1800, 6200), "SALARY", "", "PAY {ddmm}", payee="{employer}", cat="salary"),
    P("C", 1, ("r", 0.15, 24), "", "", "", payee="CREDIT INTEREST", cat="interest"),
    P("C", 2, ("r", 50, 2500), "TRANSFER", "SAVINGS", "{ref}", payee="{person}"),
]

CATALOGS = {"bank": BANK, "card": CARD, "card_fx": CARD_FX, "biz": BIZ, "savings": SAV,
            "loan": LOAN, "typed": TYPED, "typed_biz": TYPED_BIZ, "pcr": PCR,
            "biz_pcr": BIZ_PCR, "payee_pcr": PAYEE_PCR}

# At most this many of a category per statement: four salaries in one month is not
# a statement anybody receives.
CAPS = {"salary": 2, "wages": 4, "payment": 2, "interest": 2, "rwt": 1, "ird": 1}

FX_RATE = {"USD": (5600, 6300), "AUD": (8800, 9300), "GBP": (4400, 4800)}

CUSTOMERS = {
    "personal": [
        ("J SAMPLE", ["12 EXAMPLE STREET", "SAMPLEVILLE", "TESTBURY 9010"]),
        ("MS A EXAMPLE", ["FLAT 3", "45 SPECIMEN ROAD", "DEMOTOWN 7020"]),
        ("J SAMPLE & K SAMPLE", ["PO BOX 1234", "EXAMPLE BAY 3050"]),
        ("MR P TESTER", ["7 PLACEHOLDER LANE", "RD 2", "MOCKBURN 9310"]),
        ("R DEMO", ["101 TEMPLATE TERRACE", "ILLUSTRATION HEIGHTS", "TESTBURY 9011"]),
    ],
    "biz": [
        ("SAMPLE TRADING LIMITED", ["UNIT 4, 88 FICTION AVENUE", "SAMPLEVILLE 0610"]),
        ("EXAMPLE PLUMBING LTD", ["PO BOX 4321", "DEMOTOWN 7022"]),
        ("DEMO FARMS PARTNERSHIP", ["1450 MOCK VALLEY ROAD", "RD 5", "TESTBURY 9071"]),
    ],
}

# What a lines-mode layout prints in an extra text column (a Reference column
# beside the details): some rows carry one, some do not.
REF_T = ["INV {inv}", "{chq}", "REF {ref}", "{invno}", "", "", "WK {wk}"]

# Same glyph widths as "dd/mm" in every font used, so fitting is unaffected.
DDMM_MARK = "##/##"

EXTEND = ["REF 88412-00", "PARTICULARS RENT", "CODE 0114", "BRANCH 0114",
          "TRACE 556201", "NZ", "AUCKLAND", "WK 12"]


def acct_number(rng, code):
    return "%s-%04d-%07d-%03d" % (code, rng.randint(1, 9999), rng.randint(0, 9999999),
                                  rng.randint(0, 30))


def fill(rng, tmpl, st, vals):
    """Fill a description template. vals collects the foreign amount and rate so
    the NZ$ figure can be derived from them, as a real conversion would be."""
    def sub(m):
        k = m.group(1)
        if k in NAMES:
            return rng.choice(NAMES[k])
        if k == "card":
            return "%04d****%04d" % (rng.randint(4000, 5599), rng.randint(0, 9999))
        if k == "card4":
            return "****%04d" % rng.randint(0, 9999)
        if k == "time":
            return "%02d:%02d" % (rng.randint(6, 22), rng.randint(0, 59))
        if k == "fx":
            vals["fx"] = rng.randint(300, 24000)
            return mag(vals["fx"], thousands=False)
        if k in ("rate", "arate"):
            vals["rate"] = rng.randint(5600, 6300) if k == "rate" else rng.randint(8800, 9300)
            return "0.%04d" % vals["rate"]
        if k == "inv":
            return "%d.%02d" % (rng.randint(1000, 9999), rng.randint(0, 99))
        if k == "invno":
            return str(rng.randint(1000, 99999))
        if k == "ref":
            return str(rng.randint(1, 99))
        if k == "acct":
            return acct_number(rng, rng.choice(OTHER_CODES))
        if k == "custno":
            return str(rng.randint(10 ** 7, 10 ** 9 - 1))
        if k == "ddmm":
            # A date inside a description is on or just before the transaction's
            # own date, which is not known yet: a same-width marker now, the real
            # dd/mm once the dates are drawn (DDMM_MARK).
            return DDMM_MARK
        if k == "wk":
            return str(rng.randint(1, 52))
        if k == "n":
            return str(rng.randint(2, 9))
        if k == "chq":
            return "%06d" % rng.randint(1, 999999)
        raise GenError("unknown placeholder {%s}" % k)
    return re.sub(r"\{(\w+)\}", sub, tmpl)


def amount_for(rng, spec, vals, scale):
    kind = spec[0]
    if kind == "fx":
        rate = vals.get("rate") or rng.randint(5600, 6300)
        return max(1, int(round(vals["fx"] * 10000.0 / rate)))
    if kind == "fixed":
        return int(round(spec[1] * 100))
    if kind == "mult":
        lo, hi, step = spec[1:]
        return rng.randrange(lo, hi + step, step) * 100
    if kind == "pay":
        return None                      # set once the opening amount owed is known
    if kind == "big":
        return rng.randint(spec[1] * 100, spec[2] * 100)
    lo, hi = spec[1], spec[2]
    return rng.randint(int(lo * 100 * scale), int(hi * 100 * scale))


# ---------------------------------------------------------------------------
# The layout record and its defaults.
# ---------------------------------------------------------------------------

DEFAULTS = dict(
    page="A4", font="Helvetica", size=9.0, pitch=None, cpitch=None, head_size=None,
    cols=None, date_fmt="dd Mon", period_fmt="long", period_label="Statement period",
    period_sep=" to ", thousands=True, dollar=False, amount_style="lead",
    balance_style="lead", hang=(), bal_mode="every", head_every_page=True,
    head_style="rule", head_fill=GREY, opening=None, trailer=(),
    closing_label="Closing balance", totals_label="Totals", open_trailer_label="Opening balance",
    cf_bf=None, page_totals=None, hdr=None, footer=None, page_dx=None, table_dx=None,
    head_shift=0.0, head_dx=None, wrap=0.0, wrap_indent=0.0, long_desc=0.0,
    date_once=False, zebra=False, row_rules=False, col_rules=False, table_top1=260.0,
    table_top2=90.0, bottom=72.0, max_rows=None, catalog="bank", credits="normal",
    credit_p=0.17, balance_kind="positive", open_range=(300, 9000), sections=None,
    after=None, stmts=(), note="", customers="personal", scale=1.0, card=None,
    rate_text="", head_gap=None, product="", bank=None, newest_first=False,
    money_line="first", segments=None, segblock=None, seg_gap=None, bundle=None,
    large=0, zero_rows=0,
)


class Split:
    """The split-specific choices for one design. v() is a fixed dev/holdout pair;
    ch() and r() draw from DISJOINT dev and holdout menus, so a holdout layout is a
    parameter combination dev never used."""

    def __init__(self, split, arch_id):
        self.split = split
        self.dev = split == "dev"
        self.rng = random.Random(SPLIT_SEED[split] * 1000003
                                 + zlib.crc32(arch_id.encode("utf-8")))

    def v(self, d, h):
        return d if self.dev else h

    def ch(self, d, h):
        return self.rng.choice(d if self.dev else h)

    def r(self, d, h, step=0.5):
        lo, hi = d if self.dev else h
        return round(self.rng.uniform(lo, hi) / step) * step

    def n(self, lo, hi):
        return self.rng.randint(lo, hi)


def layout(S, lid, **kw):
    unknown = set(kw) - set(DEFAULTS)
    if unknown:
        raise GenError("%s: unknown layout keys %s" % (lid, sorted(unknown)))
    L = dict(DEFAULTS)
    L.update(kw)
    L["id"], L["split"] = lid, S.split
    if L["bank"] not in BANKS:
        raise GenError("%s: unknown bank %r" % (lid, L["bank"]))
    L["bank_id"] = L["bank"]
    L["bank"] = BANKS[L["bank_id"]]["mast"]
    L["size"] = float(L["size"])
    if L["pitch"] is None:
        L["pitch"] = round(L["size"] * 1.7, 2)
    if L["cpitch"] is None:
        L["cpitch"] = round(L["size"] * 1.2, 2)
    if L["head_size"] is None:
        L["head_size"] = L["size"]
    if L["head_gap"] is None:
        L["head_gap"] = L["pitch"] + 6 + (L["size"] + 2 if L["head_style"] == "dashes" else 0)
    if L["seg_gap"] is None:
        L["seg_gap"] = 2 * L["pitch"]
    L["head_dx"] = dict(L["head_dx"] or {})
    L["trailer"] = list(L["trailer"])
    kinds = [c.kind for c in L["cols"]]
    L["kinds"] = kinds
    if "date" not in kinds:
        raise GenError("%s: no date column" % lid)
    for a, b in zip(L["cols"], L["cols"][1:]):
        if b.x0 < a.x1 - 0.01:
            raise GenError("%s: columns %r and %r overlap" % (lid, a, b))
    if L["date_fmt"] not in DATE_FMTS:
        raise GenError("%s: unknown date format" % lid)
    if ("totals" in L["trailer"] or "totals_close" in L["trailer"] or L["page_totals"]) \
            and not ("debit" in kinds and "credit" in kinds):
        raise GenError("%s: a totals line needs debit and credit columns" % lid)
    if L["bal_mode"] != "none" and "balance" not in kinds:
        raise GenError("%s: balance mode %s with no balance column" % (lid, L["bal_mode"]))
    if L["cf_bf"] and "balance" not in kinds:
        raise GenError("%s: carried forward needs a balance column" % lid)
    if L["newest_first"] and (L["cf_bf"] or L["bal_mode"] == "last_of_day" or L["sections"]):
        raise GenError("%s: newest-first with forward lines or day balances" % lid)
    if L["segments"] and (L["sections"] or L["bundle"]):
        raise GenError("%s: accounts cannot also have card sections or bundles" % lid)
    return L


def acct_printed(L):
    return "acct" in L["hdr"].get("det_rows", ["acct"]) or bool(L["segments"])


# ---------------------------------------------------------------------------
# Pagination: pure geometry, so a statement can be planned before it is filled.
# A statement is a list of SLOTS (one per transaction, in display order), each
# belonging to an account SEGMENT; a segment opens with an optional title, the
# column heading and an opening line, and closes with its trailer lines.
# ---------------------------------------------------------------------------

def heading_height(L):
    nl = max(len(c.head) for c in L["cols"])
    return (nl - 1) * (L["head_size"] + 1.5) + L["head_gap"]


def after_height(L):
    if not L["after"]:
        return 0.0
    return 10 + (len(L["after"]["rows"]) + 1) * (L["size"] + 4)


def segblock_height(L):
    if not L["segblock"]:
        return 0.0
    return 8 + (len(L["segblock"]["rows"]) + 1) * (L["size"] + 4)


def title_height(L):
    return 1.5 * L["size"] + 8 if L["segments"] else 0.0


def slot_height(L, sl):
    p, cp = L["pitch"], L["cpitch"]
    return sl["pre"] * p + (sl["nlines"] - 1) * cp + p + sl["post"] * p


def paginate(L, slots):
    """Lay the rows out. Returns pages, each a list of (kind, index, y), where index
    is a slot index for a "txn" item and a segment index for everything else."""
    ph = PAGES[L["page"]][1]
    p = L["pitch"]
    bottom = ph - L["bottom"]
    tail = p * ((1 if L["cf_bf"] else 0) + (1 if L["page_totals"] else 0))
    seg_trail = p * len(L["trailer"]) + segblock_height(L)
    lead = title_height(L) + heading_height(L) + (p if L["opening"] else 0)
    pages = []
    st = {"items": None, "y": 0.0, "count": 0}

    def start_page(cont_seg):
        st["items"] = []
        pages.append(st["items"])
        st["count"] = 0
        st["y"] = L["table_top1"] if len(pages) == 1 else L["table_top2"]
        if cont_seg is not None:
            if L["head_every_page"]:
                st["items"].append(("head", cont_seg, st["y"]))
                st["y"] += heading_height(L)
            if L["cf_bf"]:
                st["items"].append(("bf", cont_seg, st["y"]))
                st["y"] += p

    def full(h, extra):
        return (st["y"] + h - p + extra > bottom
                or (L["max_rows"] and st["count"] >= L["max_rows"]))

    start_page(None)
    n = len(slots)
    for i, sl in enumerate(slots):
        h = slot_height(L, sl)
        extra = (seg_trail + (after_height(L) if i == n - 1 else 0)) if sl["seg_last"] else tail
        s = sl["seg"]
        if sl["seg_first"]:
            fresh = not any(k in ("txn", "seg_title") for k, _, _ in st["items"])
            gap = 0.0 if fresh else L["seg_gap"]
            st["y"] += gap
            if full(lead + h, extra) and not fresh:
                start_page(None)
            if L["segments"]:
                st["items"].append(("seg_title", s, st["y"]))
                st["y"] += title_height(L)
            st["items"].append(("head", s, st["y"]))
            st["y"] += heading_height(L)
            if L["opening"]:
                st["items"].append(("open", s, st["y"]))
                st["y"] += p
            if st["y"] + h - p + extra > PAGES[L["page"]][1] - L["bottom"]:
                raise GenError("%s: a sub-table does not fit on an empty page" % L["id"])
        elif full(h, extra):
            if st["count"] == 0:
                raise GenError("%s: a row does not fit on an empty page" % L["id"])
            if L["page_totals"]:
                st["items"].append(("ptot", s, st["y"]))
                st["y"] += p
            if L["cf_bf"]:
                st["items"].append(("cf", s, st["y"]))
                st["y"] += p
            start_page(s)
        st["items"].append(("txn", i, st["y"]))
        st["y"] += h
        st["count"] += 1
        if sl["seg_last"]:
            for t in L["trailer"]:
                st["items"].append((t, s, st["y"]))
                st["y"] += p
            if L["segblock"]:
                st["items"].append(("segblock", s, st["y"] + 6))
                st["y"] += segblock_height(L)
            if i == n - 1 and L["after"]:
                st["items"].append(("after", None, st["y"] + 10))
    return pages


def page_rows(pages):
    return [[idx for k, idx, _ in items if k == "txn"] for items in pages]


def col_of(L, kind):
    for c in L["cols"]:
        if c.kind == kind:
            return c
    return None


def desc_col(L):
    return col_of(L, "desc")


def seg_cfgs(L):
    """One config per account segment: the layout's own settings, overridden by
    the segment's (a multi-account statement mixes an everyday and a saver)."""
    keys = ("catalog", "credits", "credit_p", "balance_kind", "open_range", "scale",
            "large", "zero_rows", "product", "card", "rate_text")
    base = {k: L[k] for k in keys}
    base["kinds"] = set(L["kinds"])
    out = []
    for seg in (L["segments"] or [{}]):
        c = dict(base)
        c.update(seg)
        out.append(c)
    return out


# ---------------------------------------------------------------------------
# One account's rows: direction, words, amounts, dates, balances.
# ---------------------------------------------------------------------------

def gen_rows(rng, L, cfg, st, nls, page_of, section_of, width_of, measure, forced_open=None):
    """Invent one account's rows in DISPLAY order. width_of(kind) is the room in a
    column (None: unlimited, as in an export). Returns (rows, opening, closing, limit)."""
    n = len(nls)
    cat = CATALOGS[cfg["catalog"]]
    newest = L["newest_first"]
    kinds = cfg["kinds"]

    def fitw(kind, s, indent=0.0):
        w = width_of(kind)
        return s.strip() if w is None else fit(measure, s, w - indent)

    # Direction of each row, which is where sparse deposits are arranged.
    plan = cfg["credits"]
    if plan in ("normal", "loan"):
        p_c = cfg["credit_p"] if plan == "normal" else 0.6
        dirs = ["C" if rng.random() < p_c else "D" for _ in range(n)]
        if "C" not in dirs:
            dirs[rng.randrange(n)] = "C"
    elif plan in ("sparse", "sparse_none_p1", "none_p1", "card"):
        allowed = [i for i in range(n) if plan in ("sparse", "card") or page_of[i] >= 1]
        if not allowed:
            raise GenError("%s: %s needs a second page" % (L["id"], plan))
        if plan == "none_p1":
            dirs = ["C" if (page_of[i] >= 1 and rng.random() < cfg["credit_p"]) else "D"
                    for i in range(n)]
            if "C" not in dirs:
                dirs[rng.choice(allowed)] = "C"
        else:
            pick = set(rng.sample(allowed, min(rng.choice([1, 2]), len(allowed))))
            dirs = ["C" if i in pick else "D" for i in range(n)]
    else:
        raise GenError("unknown credit plan %r" % plan)
    if "D" not in dirs:
        dirs[0] = "D"

    # Forced special rows: a very large settlement pair, and fee-waived zero rows.
    special = {}
    singles = [i for i in range(1, n - 1) if nls[i] == 1]
    if cfg["large"]:
        a, b = sorted(rng.sample(singles, 2))
        first, second = (b, a) if newest else (a, b)      # money arrives, then leaves
        special[first], special[second] = ("C", "large"), ("D", "large")
    if cfg["zero_rows"]:
        for i in rng.sample([i for i in singles if i not in special], cfg["zero_rows"]):
            special[i] = ("D", "zero")
    for i, (d, _) in special.items():
        dirs[i] = d

    rows = []
    used = {}
    first_credit = True
    for i in range(n):
        d = dirs[i]
        if i in special:
            cands = [e for e in cat if e["cat"] == special[i][1] and e["dir"] == d]
        else:
            cands = [e for e in cat if e["dir"] == d and len(e["lines"]) == nls[i] and e["w"] > 0]
            cands = [e for e in cands if used.get(e["cat"], 0) < CAPS.get(e["cat"], 10 ** 6)] or cands
            if cfg["catalog"].startswith("card") and d == "C" and first_credit:
                cands = [e for e in cands if e["cat"] == "payment"] or cands
        if d == "C":
            first_credit = False
        if not cands:
            raise GenError("%s: catalogue has no %s entry with %d line(s)" % (L["id"], d, nls[i]))
        e = rng.choices(cands, weights=[max(c["w"], 1) for c in cands])[0]
        used[e["cat"]] = used.get(e["cat"], 0) + 1
        vals = {}
        r = {"dir": d, "cat": e["cat"], "pre": 0, "post": 0, "nlines": nls[i],
             "page": page_of[i], "section": section_of[i], "type": "", "payee": "",
             "part": "", "code": "", "ref": "", "lines": [], "fxcur": "", "fxamt": ""}
        if e.get("fx"):
            vals["fx"] = rng.randint(300, 24000)
            vals["rate"] = rng.randint(*FX_RATE[e["fx"]])
            r["fxcur"], r["fxamt"] = e["fx"], mag(vals["fx"], L["thousands"])
        if "type" in e:
            r["type"] = e["type"]
            r["lines"] = [fitw("desc", fill(rng, e["lines"][0], st, vals))]
        elif "part" in e:
            for k in ("payee", "part", "code", "ref"):
                if k not in kinds:
                    continue
                txt = fill(rng, e[k], st, vals)
                if k != "payee":
                    txt = txt[:12].strip()          # the NZ 12-character field
                r[k] = fitw(k, txt)
            if "desc" in kinds:
                raise GenError("%s: a P/C/R catalogue in a layout with a description" % L["id"])
        else:
            for j, t in enumerate(e["lines"]):
                r["lines"].append(fitw("desc", fill(rng, t, st, vals),
                                       L["wrap_indent"] if j else 0.0))
            if len(r["lines"]) == 1 and width_of("desc") and rng.random() < L["long_desc"]:
                # A LONG description that runs up to the first money column: the
                # common real shape where a reference is tacked on the end.
                maxw = width_of("desc")
                s = r["lines"][0]
                for tok in rng.sample(EXTEND, len(EXTEND)):
                    if measure(s + " " + tok) <= maxw:
                        s = s + " " + tok
                r["lines"][0] = s
                r["long"] = maxw - measure(s) < 14
            for k in ("payee", "part", "code", "ref"):
                if k in kinds:
                    r[k] = fitw(k, fill(rng, rng.choice(REF_T), st, vals)[:12].strip())
        if any(not t for t in r["lines"][1:]):
            raise GenError("%s: an empty continuation line" % L["id"])
        r["amt"] = amount_for(rng, e["amt"], vals, cfg["scale"])
        rows.append(r)

    # Dates: sorted within each section; newest-first reverses them on the page.
    # A card prints in processed-date order, with the transaction a day or so before.
    days = (st["end"] - st["start"]).days + 1
    for s in sorted(set(section_of)):
        idx = [i for i in range(n) if section_of[i] == s]
        ds = sorted(st["start"] + dt.timedelta(days=rng.randrange(days)) for _ in idx)
        if newest:
            ds = ds[::-1]
        for i, d in zip(idx, ds):
            if "pdate" in kinds:
                rows[i]["pdate"] = d
                rows[i]["date"] = max(st["start"],
                                      d - dt.timedelta(days=rng.choice([0, 0, 1, 1, 2, 3])))
            else:
                rows[i]["date"] = d

    for r in rows:
        if any(DDMM_MARK in x for x in r["lines"] + [r["part"], r["code"], r["ref"]]):
            d = max(st["start"], r["date"] - dt.timedelta(days=rng.choice([0, 1, 2])))
            ddmm = "%02d/%02d" % (d.day, d.month)
            r["lines"] = [x.replace(DDMM_MARK, ddmm) for x in r["lines"]]
            for k in ("part", "code", "ref"):
                r[k] = r[k].replace(DDMM_MARK, ddmm)
        r["money_like"] = bool(re.search(r"\d+\.\d\d|\d\d:\d\d|\*\*\*\*|\d\d/\d\d|\b\d+ x ",
                                         " ".join(r["lines"] + [r["part"], r["code"], r["ref"]])))

    # Balances, in DATE order whatever the printed order. Integer cents.
    order = list(range(n))[::-1] if newest else list(range(n))
    kind = cfg["balance_kind"]
    limit = None
    if kind == "card":
        owed = rng.randint(30000, 480000)
        opening = -owed
        remaining = owed
        for i in order:
            if rows[i]["amt"] is None:
                rows[i]["amt"] = max(500, rng.randint(remaining // 4,
                                                      max(remaining // 2, remaining // 4 + 1)))
                remaining -= rows[i]["amt"]
                if remaining < 0:
                    raise GenError("%s: card payments exceed the amount owed" % L["id"])
    for r in rows:
        if r["amt"] is None or r["amt"] < 0 or (r["amt"] == 0 and r["cat"] != "zero"):
            raise GenError("%s: a row with no amount" % L["id"])
    signed = {i: (-rows[i]["amt"] if rows[i]["dir"] == "D" else rows[i]["amt"]) for i in range(n)}
    pref, acc = [], 0
    for i in order:
        acc += signed[i]
        pref.append(acc)
    lo_p, hi_p = min(pref), max(pref)
    if forced_open is not None:
        opening = forced_open
    elif kind == "positive":
        floor = rng.randint(4000, 90000)
        base = rng.randint(cfg["open_range"][0] * 100, cfg["open_range"][1] * 100)
        opening = max(base, floor - lo_p)
    elif kind == "overdraft":
        q = (hi_p - lo_p) // 4
        lo, hi = -hi_p + q, -lo_p - q
        if lo >= hi:
            raise GenError("%s: cannot make the balance cross zero" % L["id"])
        opening = rng.randint(lo, hi)
    elif kind == "loan":
        opening = -rng.randint(15000000, 65000000)
    elif kind != "card":
        raise GenError("unknown balance kind %r" % kind)
    if forced_open is None:
        while opening == 0 or any(opening + x == 0 for x in pref):
            opening += 7
    bal = opening
    for i in order:
        bal += signed[i]
        r = rows[i]
        r["bal"] = bal
        r["debit"] = r["amt"] if r["dir"] == "D" else None
        r["credit"] = r["amt"] if r["dir"] == "C" else None
    closing = bal
    if kind == "positive" and min(opening + x for x in pref) < 100:
        raise GenError("%s: an account meant to stay in credit went below $1" % L["id"])
    if kind == "loan" and (opening >= 0 or any(r["bal"] >= 0 for r in rows)):
        raise GenError("%s: a loan balance went into credit" % L["id"])
    if kind == "overdraft" and not (any(r["bal"] < 0 for r in rows)
                                    and any(r["bal"] > 0 for r in rows)):
        raise GenError("%s: the overdraft statement never crossed zero" % L["id"])
    if kind == "card":
        # A card's limit has to cover what is owed on it, or the summary prints an
        # impossible "available credit 0.00" on a card 1,400 over its limit.
        limit = (cfg["card"] or {}).get("limit", 500000)
        peak = max([-opening] + [-r["bal"] for r in rows])
        while peak > 0.8 * limit:
            limit += 100000

    # WHICH balances are printed.
    for pos, i in enumerate(order):
        r = rows[i]
        if L["bal_mode"] == "every":
            r["bal_printed"] = True
        elif L["bal_mode"] == "last_of_day":
            nxt = order[pos + 1] if pos + 1 < n else None
            r["bal_printed"] = nxt is None or rows[nxt]["date"] != r["date"]
        else:
            r["bal_printed"] = False

    # The description the truth records: printed pieces in reading order.
    order_src = cfg["text_order"] if "text_order" in cfg else [c.kind for c in L["cols"]]
    text_order = [k for k in order_src if k in TEXT_KINDS]
    for r in rows:
        pieces = []
        for k in text_order:
            if k == "desc":
                pieces.append(r["lines"][0] if r["lines"] else "")
            elif r[k]:
                pieces.append(r[k])
        pieces += r["lines"][1:]
        r["desc"] = " ".join(p for p in pieces if p)
    return rows, opening, closing, limit


def dx_series(mode, npages, rng):
    if not mode:
        return [0.0] * npages
    kind, a = mode
    if kind == "gutter":
        return [0.0 if p % 2 == 0 else a for p in range(npages)]
    if kind == "jitter":
        return [round(rng.uniform(-a, a) * 2) / 2.0 for _ in range(npages)]
    if kind == "drift":
        return [p * a for p in range(npages)]
    if kind == "list":
        return [float(a[p % len(a)]) for p in range(npages)]
    raise GenError("unknown offset mode %r" % kind)


def build_statement(L, spec, seed, ident=None, forced_open=None):
    """Plan, fill and balance one statement. Returns the statement record."""
    rng = random.Random(seed)
    measure = lambda s: stringWidth(s, FONTS[L["font"]][0], L["size"])   # noqa: E731

    def width_of(kind):
        c = col_of(L, kind)
        return None if c is None else c.x1 - c.x0 - PAD_T - 2

    ym = spec["ym"]
    start = dt.date(ym[0], ym[1], ym[2] if len(ym) > 2 else 1)
    end = add_months(start, spec.get("months", 1)) - dt.timedelta(days=1)
    st = {"start": start, "end": end, "spec": spec}
    cfgs = seg_cfgs(L)
    nseg = len(cfgs)

    def can_wrap(cfg):
        cat = CATALOGS[cfg["catalog"]]
        return ("type" not in cat[0] and "part" not in cat[0]
                and any(len(e["lines"]) == 2 for e in cat))

    def lines_for(cfg, k):
        return [2 if (can_wrap(cfg) and rng.random() < L["wrap"]) else 1 for _ in range(k)]

    # 1. PLAN the shape. Nothing about direction or amount affects the geometry, so
    #    the page breaks are known before the data, which is how a short last page
    #    or "no deposits on page 1" is arranged.
    n = spec["n"]
    if nseg == 1:
        N = n + 40
        nl = lines_for(cfgs[0], N)
        if spec.get("short_last") and not L["sections"]:
            big = [{"nlines": k, "pre": 0, "post": 0, "seg": 0, "seg_first": j == 0,
                    "seg_last": j == N - 1} for j, k in enumerate(nl)]
            starts = [pr[0] for pr in page_rows(paginate(L, big))[1:]]
            cands = [s + j for s in starts for j in (1, 2) if 8 <= s + j <= 60 and s + j < N]
            if not cands:
                raise GenError("%s: no way to leave a short last page" % L["id"])
            n = min(cands, key=lambda c: (abs(c - spec["n"]), c))
        ns = [n]
        nls = [nl[:n]]
    else:
        shares = [seg.get("share", 1.0) for seg in L["segments"]]
        ns = [max(5, int(round(n * s / sum(shares)))) for s in shares]
        nls = [lines_for(c, k) for c, k in zip(cfgs, ns)]
    slots = []
    for s in range(nseg):
        for j in range(ns[s]):
            slots.append({"nlines": nls[s][j], "pre": 0, "post": 0, "seg": s,
                          "seg_first": j == 0, "seg_last": j == ns[s] - 1})
    section_of = [0] * len(slots)
    if L["sections"]:
        k = len(L["sections"])
        m = len(slots)
        cuts = sorted(rng.sample(range(4, m - 3), k - 1))
        bounds = [0] + cuts + [m]
        for s in range(k):
            for i in range(bounds[s], bounds[s + 1]):
                section_of[i] = s
            slots[bounds[s]]["pre"] = 1
            slots[bounds[s + 1] - 1]["post"] = 1
    pages = paginate(L, slots)
    if len(pages) > 4:
        raise GenError("%s: %d pages (the brief says 1-4)" % (L["id"], len(pages)))
    page_of = [0] * len(slots)
    for pi, pr in enumerate(page_rows(pages)):
        for i in pr:
            page_of[i] = pi

    # 2. WHO and WHICH account. A pile of one bank's statements comes from several
    #    customers; a bundle is one customer's consecutive months.
    B = BANKS[L["bank_id"]]
    code = spec.get("code") or B["codes"][0]
    if code not in B["codes"]:
        raise GenError("%s: bank code %s is not %s's" % (L["id"], code, B["name"]))
    if ident is None:
        who = rng.choice(CUSTOMERS[L["customers"]])
        ident = {"cust": who[0], "addr": who[1], "acct": acct_number(rng, code),
                 "card_no": "4%03d **** **** %04d" % (rng.randint(0, 999), rng.randint(0, 9999)),
                 "stno": str(rng.randint(3, 140)),
                 "branch": "%04d %s" % (rng.randint(100, 9999), rng.choice(NAMES["branch"]))}
    st.update(ident)
    st["ident"] = ident
    st["code"] = ident["acct"][:2]
    st["sec_cards"] = ["4%03d **** **** %04d" % (rng.randint(0, 999), rng.randint(0, 9999))
                       for _ in (L["sections"] or [])]

    # 3. FILL each account.
    rows, segs = [], []
    for s, cfg in enumerate(cfgs):
        idx = [i for i, sl in enumerate(slots) if sl["seg"] == s]
        seg_rows, op, cl, limit = gen_rows(
            rng, L, cfg, st, [slots[i]["nlines"] for i in idx], [page_of[i] for i in idx],
            [section_of[i] for i in idx], width_of, measure,
            forced_open=forced_open if s == 0 else None)
        for r, i in zip(seg_rows, idx):
            r["seg"] = s
            r["pre"], r["post"] = slots[i]["pre"], slots[i]["post"]
        base = ident["acct"]
        acct = base if s == 0 else "%s-%03d" % (base[:-4], (int(base[-3:]) + s) % 1000)
        segs.append({"opening": op, "closing": cl, "limit": limit, "acct": acct,
                     "product": cfg["product"], "cfg": cfg})
        rows += seg_rows
    st.update(rows=rows, pages=pages, segs=segs, opening=segs[0]["opening"],
              closing=segs[-1]["closing"], limit=segs[0]["limit"])

    # 4. REDACTIONS and IDENTITY. At most one redaction per row; whole-row removal
    #    only ever removes (a box over a whole row that left the text readable would
    #    just be an overlay of every field).
    st["acct_mode"] = spec.get("acct")
    red = spec.get("redact")
    if red:
        cand = [i for i, r in enumerate(rows)
                if 0 < i < len(rows) - 1 and not r["pre"] and not r["post"]
                and r["cat"] not in ("large", "zero")]
        chosen = set()
        for field in ("row", "desc", "amount", "balance"):
            k = red.get(field, 0)
            if field == "row" and red["mode"] != "remove":
                continue
            pool = [i for i in cand if i not in chosen and i - 1 not in chosen
                    and i + 1 not in chosen and (field != "balance" or rows[i]["bal_printed"])]
            for i in rng.sample(pool, min(k, len(pool))):
                rows[i]["redact"] = {field: red["mode"]}
                chosen.add(i)

    st["dx"] = dx_series(L["page_dx"], len(pages), rng)
    st["tdx"] = dx_series(L["table_dx"], len(pages), rng)
    return st


# ---------------------------------------------------------------------------
# Rendering.
# ---------------------------------------------------------------------------

def tok_width(sh, L):
    return max(sh.width(t, L["size"]) for t in ("CR", "DR", "OD"))


def cell(sh, L, c, y, s, bold=False, indent=0.0, mode=None):
    """One text cell. mode None draws it; "overlay" draws it and boxes it over;
    "remove" draws only the box where it would have been."""
    if c.align == "right":
        x = c.x1 - PAD_M
    elif c.align == "center":
        x = (c.x0 + c.x1) / 2.0
    else:
        x = c.x0 + PAD_T + indent
    a, b = sh.place(x, s, L["size"], bold, c.align)
    if a < c.x0 - 0.01 or b > c.x1 + 0.01:
        raise GenError("%s: %r (%.1f-%.1f) spills out of column %r" % (L["id"], s, a, b, c))
    if mode != "remove":
        sh.text(x, y, s, L["size"], bold=bold, align=c.align)
    if mode:
        sh.blackout(a, b, y, L["size"])


def money_cell(sh, L, c, y, num, tok, bold=False, mode=None):
    if c.kind in L["hang"] and c.align == "right":
        # The token sits in its own slot at the right, so the figures stay aligned
        # whether or not a row carries one.
        xr = c.x1 - PAD_M
        xn = xr - tok_width(sh, L) - 3
        a, b = sh.place(xn, num, L["size"], bold, "right")
        if a < c.x0 - 0.01:
            raise GenError("%s: %r spills out of column %r" % (L["id"], num, c))
        if mode != "remove":
            if tok:
                sh.text(xr, y, tok, L["size"], bold=bold, align="right")
            sh.text(xn, y, num, L["size"], bold=bold, align="right")
        if mode:
            sh.blackout(a, xr if tok else b, y, L["size"])
    else:
        cell(sh, L, c, y, joined(num, tok), bold=bold, mode=mode)


def holder_style(L, kind):
    if kind in ("debit", "credit"):
        return "plain"
    if kind == "amount":
        return L["amount_style"]
    return L["balance_style"]


def fmt_cell(L, kind, v):
    return fmt_money(v, holder_style(L, kind), L)


def label_region(L):
    """Where a non-transaction label ("Opening balance") may run: from the first
    text column to the next column that is not text."""
    cols = L["cols"]
    k = next(i for i, c in enumerate(cols) if c.kind in TEXT_KINDS)
    stop = next((c.x0 for c in cols[k + 1:] if c.kind not in TEXT_KINDS), cols[-1].x1)
    return cols[k].x0 + PAD_T, stop - 3


def draw_label(sh, L, y, s, bold=False, size=None):
    x0, xmax = label_region(L)
    _, b = sh.text(x0, y, s, size or L["size"], bold=bold)
    if b > xmax:
        raise GenError("%s: label %r runs into the next column" % (L["id"], s))


def figure_col(L):
    """Where an opening/closing figure goes: the balance column, or on a layout
    without one (a card), the amount column."""
    return col_of(L, "balance") or col_of(L, "amount")


def draw_heading(sh, L, y):
    s = L["head_size"]
    cols = L["cols"]
    nl = max(len(c.head) for c in cols)
    lh = s + 1.5
    x0, x1 = cols[0].x0, cols[-1].x1
    top = y - s - 3
    bot = y + (nl - 1) * lh + 4
    style = L["head_style"]
    ink = None
    if style == "bar":
        sh.rect(x0 - 2, top, x1 - x0 + 4, bot - top, fill=L["head_fill"])
    elif style == "dark":
        sh.rect(x0 - 2, top, x1 - x0 + 4, bot - top, fill=L["head_fill"])
        ink = WHITE
    elif style == "box":
        sh.rect(x0 - 2, top, x1 - x0 + 4, bot - top, stroke=BLACK, width=0.5)
    for c in cols:
        # A heading printed a few points off its own column, or the whole heading
        # row set at a different x from the data: both happen on real statements.
        off = L["head_shift"] + L["head_dx"].get(c.kind, 0.0)
        for k, t in enumerate(c.head):
            yy = y + (nl - len(c.head) + k) * lh
            a, b = sh.text(c.anchor(c.head_align) + off, yy, t, s, bold=True,
                           align=c.head_align, color=ink)
            if a < c.x0 + off - 0.01 or b > c.x1 + off + 0.01:
                raise GenError("%s: heading %r spills out of %r" % (L["id"], t, c))
    if style in ("rule", "bar"):
        sh.line(x0, bot, x1, bot, width=0.6)
    if style == "rule2":
        sh.line(x0, top - 1, x1, top - 1, width=0.9)
        sh.line(x0, bot, x1, bot, width=0.5)
    if style == "dashes":
        dy = y + (nl - 1) * lh + s + 1
        for c in cols:
            w = c.x1 - c.x0 - PAD_T - PAD_M
            sh.text(c.x0 + PAD_T, dy, "-" * max(1, int(w / sh.width("-", s))), s)


def text_mode(r, field):
    return (r.get("redact") or {}).get(field)


def draw_txn(sh, L, st, i, y, show_date):
    r = st["rows"][i]
    p, cp, s = L["pitch"], L["cpitch"], L["size"]
    rec = {"date": None, "pdate": None, "money": {}, "text": []}
    if r["pre"]:
        draw_label(sh, L, y, L["sections"][r["section"]].format(card=st["sec_cards"][r["section"]]),
                   bold=True)
        y += p
    y_last = y + (r["nlines"] - 1) * cp
    # Money on the LAST line of a two-line transaction: the staggered shape.
    ym = y_last if L["money_line"] == "last" else y
    x0, x1 = L["cols"][0].x0, L["cols"][-1].x1
    if L["zebra"] and r["zebra"]:
        sh.rect(x0 - 2, y - 0.35 * s - p / 2.0, x1 - x0 + 4, (y_last - y) + p, fill=PALE)
    if text_mode(r, "row") == "remove":
        # The whole transaction is gone: one bar per printed line.
        for yy in sorted({y, y_last, ym}):
            sh.rect(x0, yy - 0.85 * s, x1 - x0, 1.15 * s, fill=BLACK)
        r["rec"] = None
        return y_last
    dmode = text_mode(r, "desc")
    for c in L["cols"]:
        k = c.kind
        if k == "date":
            if show_date:
                rec["date"] = fmt_date(r["date"], L["date_fmt"])
                cell(sh, L, c, y, rec["date"])
        elif k == "pdate":
            rec["pdate"] = fmt_date(r["pdate"], L["date_fmt"])
            cell(sh, L, c, y, rec["pdate"])
        elif k == "desc":
            cell(sh, L, c, y, r["lines"][0], mode=dmode)
            if dmode != "remove":
                rec["text"].append(r["lines"][0])
        elif k in TEXT_KINDS:
            if r[k]:
                cell(sh, L, c, y, r[k], mode=dmode)
                if dmode != "remove":
                    rec["text"].append(r[k])
        elif k in FX_KINDS:
            if r[k]:
                cell(sh, L, c, ym, r[k])
        elif k in ("debit", "credit"):
            if r[k] is not None:
                num, tok = fmt_cell(L, k, r[k])
                m = text_mode(r, "amount")
                money_cell(sh, L, c, ym, num, tok, mode=m)
                if m != "remove":
                    rec["money"][k] = (num, tok)
        elif k == "amount":
            v = -r["amt"] if r["dir"] == "D" else r["amt"]
            num, tok = fmt_cell(L, k, v)
            m = text_mode(r, "amount")
            money_cell(sh, L, c, ym, num, tok, mode=m)
            if m != "remove":
                rec["money"][k] = (num, tok)
        elif k == "balance":
            if r["bal_printed"]:
                num, tok = fmt_cell(L, k, r["bal"])
                m = text_mode(r, "balance")
                money_cell(sh, L, c, ym, num, tok, mode=m)
                if m != "remove":
                    rec["money"][k] = (num, tok)
    dc = desc_col(L)
    for j, t in enumerate(r["lines"][1:], 1):
        cell(sh, L, dc, y + j * cp, t, indent=L["wrap_indent"], mode=dmode)
        if dmode != "remove":
            rec["text"].append(t)
    if L["row_rules"]:
        yy = y_last + p / 2.0 - 0.35 * s
        sh.line(x0, yy, x1, yy, width=0.25, color=(0.6, 0.6, 0.6))
    if r["post"]:
        yt = y_last + p
        sec = [x for x in st["rows"] if x["section"] == r["section"]]
        tot = sum((-x["amt"] if x["dir"] == "D" else x["amt"]) for x in sec)
        draw_label(sh, L, yt, "Total for card ending %s" % st["sec_cards"][r["section"]][-4:],
                   bold=True)
        num, tok = fmt_cell(L, "amount", tot)
        money_cell(sh, L, col_of(L, "amount"), yt, num, tok, bold=True)
    r["rec"] = rec
    return y_last


def acct_shown(st, acct):
    if st["acct_mode"] == "masked":
        return "%s-XXXX-XXXXXXX-%s" % (acct[:2], acct[-3:])
    return acct


def draw_acct(sh, x, y, size, st, acct, bold=False):
    """An account number, as this statement's identity case prints it."""
    s = acct_shown(st, acct)
    mode = st["acct_mode"]
    if mode == "removed":
        a, b = sh.place(x, s, size, bold)
        sh.blackout(a, b, y, size)
        return b
    a, b = sh.text(x, y, s, size, bold=bold)
    if mode == "overlay":
        sh.blackout(a, b, y, size)
    return b


def summary_rows(L, st, keys):
    rows = st["rows"]
    card = L["balance_kind"] == "card"

    def bal(v):
        if card:
            num, tok = fmt_money(v, L["amount_style"], L)
        else:
            num, tok = fmt_money(v, L["balance_style"] if "balance" in L["kinds"] else "lead", L)
        return joined(num, tok)

    out = []
    for label, key in keys:
        seg = 0
        if key and ":" in key:
            key, seg = key.split(":")
            seg = int(seg)
        sr = [r for r in rows if r["seg"] == seg]
        if key == "open":
            v = bal(st["segs"][seg]["opening"])
        elif key == "close":
            v = bal(st["segs"][seg]["closing"])
        elif key == "tot_d":
            v = mag(sum(r["debit"] or 0 for r in sr), L["thousands"], L["dollar"])
        elif key == "tot_c":
            v = mag(sum(r["credit"] or 0 for r in sr), L["thousands"], L["dollar"])
        elif key == "fees":
            v = mag(sum(r["amt"] for r in sr if r["cat"] == "fee"), L["thousands"], L["dollar"])
        elif key == "rate":
            v = st["segs"][seg]["cfg"]["rate_text"] or L["rate_text"]
        elif key == "limit":
            v = mag(st["limit"], L["thousands"], L["dollar"])
        elif key == "avail":
            v = mag(max(0, st["limit"] + st["closing"]), L["thousands"], L["dollar"])
        elif key == "minpay":
            v = mag(max(1000, (-st["closing"] * 3 // 100) // 100 * 100), L["thousands"], L["dollar"])
        elif key == "due":
            v = fmt_date(st["end"] + dt.timedelta(days=25), "d Month yyyy")
        elif key == "count":
            v = str(len(sr))
        elif key == "odlimit":
            v = mag(100000, L["thousands"], L["dollar"])
        else:
            raise GenError("unknown summary key %r" % key)
        out.append((label, v))
    return out


def clear_of(xs, anchors, d=10.0):
    return all(abs(x - a) >= d for x in xs for a in anchors)


def nudge(L, xs, lo, hi, what):
    """Shift a header box so none of its figures lines up with a table money
    column: the brief is a summary box that is NOT a fifth column in disguise."""
    anchors = [c.anchor() for c in L["cols"] if c.kind in MONEY_KINDS or c.kind == "fxamt"]
    for t in [0, -6, 6, -12, 12, -18, 18, -24, 24, -30, 30, -36, 36]:
        if clear_of([x + t for x in xs], anchors) and lo + t >= 12 \
                and hi + t <= PAGES[L["page"]][0] - 12:
            return t
    raise GenError("%s: cannot place the %s clear of the table columns" % (L["id"], what))


def draw_box(sh, L, st, B, what):
    """A summary or sidebar box: title, then label ... value rows."""
    sz = B.get("size", L["size"])
    style = B.get("style", "box")
    rows = summary_rows(L, st, B["rows"])
    x, y, w = B["x"], B["y"], B["w"]
    if style == "band":
        cw = w / float(len(rows))
        xs = [x + (k + 1) * cw - 8 for k in range(len(rows))]
        x += nudge(L, xs, x, x + w, what)
        for k, (lab, v) in enumerate(rows):
            cx = x + k * cw
            sh.rect(cx, y, cw - 4, 2 * sz + 16, fill=B.get("fill", PALE))
            sh.text(cx + 6, y + sz + 4, lab, sz - 0.5)
            sh.text(cx + cw - 8, y + 2 * sz + 9, v, sz + 1, bold=True, align="right")
        return
    lh = sz + B.get("lead", 5)
    title = B.get("title")
    h = lh * (len(rows) + (1 if title else 0)) + 7
    x += nudge(L, [x + w - 6], x, x + w, what)
    if style == "box":
        sh.rect(x, y, w, h, stroke=B.get("stroke", BLACK), width=0.6)
    elif style == "shaded":
        sh.rect(x, y, w, h, fill=B.get("fill", PALE))
    yy = y + lh + 1
    if title:
        sh.text(x + 6, yy, title, sz + 0.5, bold=True)
        yy += lh
    for lab, v in rows:
        strong = lab.lower().startswith(("closing", "new balance", "amount due"))
        sh.text(x + 6, yy, lab, sz, bold=strong)
        if v:
            sh.text(x + w - 6, yy, v, sz, bold=strong, align="right")
        yy += lh
    for extra in B.get("notes", []):
        sh.text(x + 6, yy, extra, sz - 1)
        yy += sz + 1


def details_rows(L, st, keys):
    p0, p1 = period_text(L["period_fmt"], st["start"], st["end"])
    out = []
    for k in keys:
        if k == "product":
            out.append(("Account", L["product"]))
        elif k == "acct":
            out.append(("Account number", None))          # drawn by draw_acct
        elif k == "period":
            out.append((L["period_label"], p0 + L["period_sep"] + p1))
        elif k == "stno":
            out.append(("Statement number", st["stno"]))
        elif k == "branch":
            out.append(("Branch", st["branch"]))
        elif k == "card":
            out.append(("Card number", st["card_no"]))
        elif k == "issued":
            out.append(("Date issued", fmt_date(st["end"] + dt.timedelta(days=2), "d Month yyyy")))
        elif k == "printed":
            out.append(("Printed", fmt_date(st["end"] + dt.timedelta(days=3), "dd/mm/yyyy")
                        + " 10:15"))
        elif k == "name":
            out.append(("Account name", st["cust"]))
        else:
            raise GenError("unknown details key %r" % k)
    return out


def draw_first_header(sh, L, st):
    H = L["hdr"]
    pw = sh.w
    left, right = H["left"], H["right"]
    mast = H["mast"]
    color = H.get("color", DARK)
    fictional = BANKS[L["bank_id"]].get("fictional")
    if mast == "bar":
        sh.rect(0, 0, pw, H.get("bar_h", 58), fill=color)
        sh.text(left, 38, L["bank"], 18, bold=True, color=H.get("ink", WHITE))
        sh.text(right, 38, H["title"], 11, align="right", color=H.get("ink", WHITE))
    elif mast == "plain":
        sh.text(left, 50, L["bank"], 20, bold=True)
        if H.get("tagline"):
            sh.text(left, 63, H["tagline"], 7.5)
        sh.text(right, 50, H["title"], 12, bold=True, align="right")
    elif mast == "right":
        sh.text(right, 46, L["bank"], 17, bold=True, align="right")
        sh.text(right, 60, H["title"], 9, align="right")
    elif mast == "logo":
        # A coloured mark for the fictional bank only; real banks get no logo.
        if fictional:
            sh.rect(left, 26, 22, 22, fill=color)
            sh.text(left + 30, 44, L["bank"], 16, bold=True)
        else:
            sh.text(left, 44, L["bank"], 16, bold=True)
        sh.text(right, 44, H["title"], 11, bold=True, align="right")
    elif mast == "courier":
        mid = (left + right) / 2.0
        sh.text(mid, 40, L["bank"].upper(), 11, bold=True, align="center")
        sh.text(mid, 54, H["title"].upper(), 9, align="center")
        sh.line(left, 62, right, 62, width=1.0)
    elif mast == "ib":
        # An internet-banking printout: no masthead, no bank name, no address.
        sh.text(left, 40, H["title"], 13, bold=True)
        sh.line(left, 48, right, 48, width=0.4, color=MID)
    else:
        raise GenError("unknown masthead %r" % mast)
    if H.get("addr"):
        ax, ay = H["addr"]
        asz = H.get("addr_size", 9)
        for k, s in enumerate([st["cust"]] + st["addr"]):
            sh.text(ax, ay + k * (asz + 2.5), s, asz, bold=(k == 0 and H.get("addr_bold", False)))
    dx_, dy_, lw = H["det"]
    dsz = H.get("det_size", 8.5)
    det = details_rows(L, st, H.get("det_rows", ["acct", "period", "stno"]))
    lw = max(lw, max(sh.width(lab, dsz, H.get("det_bold", False)) for lab, _ in det) + 8)
    for k, (lab, v) in enumerate(det):
        yy = dy_ + k * (dsz + 4)
        sh.text(dx_, yy, lab, dsz, bold=H.get("det_bold", False))
        if v is None:
            draw_acct(sh, dx_ + lw, yy, dsz, st, st["segs"][0]["acct"])
        else:
            sh.text(dx_ + lw, yy, v, dsz)
    if H.get("summary"):
        draw_box(sh, L, st, H["summary"], "summary box")
    if H.get("sidebar"):
        draw_box(sh, L, st, H["sidebar"], "sidebar")
    if H.get("table_title"):
        sh.text(L["cols"][0].x0, L["table_top1"] - L["size"] - 12, H["table_title"],
                L["size"] + 2, bold=True)


def draw_cont_header(sh, L, st, pno, npg):
    H = L["hdr"]
    left, right = H["left"], H["right"]
    p0, p1 = period_text(L["period_fmt"], st["start"], st["end"])
    period = "%s%s%s" % (p0, L["period_sep"], p1)
    cont = H.get("cont", "line")
    named = H["mast"] != "ib"
    show_acct = acct_printed(L)
    card = "card" in H.get("det_rows", [])

    def ident_line(x, y, size, upper=False):
        if show_acct:
            lab = "ACCOUNT" if upper else "Account number"
            sh.text(x, y, lab, size)
            draw_acct(sh, x + sh.width(lab, size) + 5, y, size, st, st["segs"][0]["acct"])
        elif card:
            sh.text(x, y, "Card number %s" % st["card_no"], size)

    if cont == "line":
        sh.text(left, 40, "%s  %s" % (L["bank"], L["product"]) if named else L["product"], 9,
                bold=True)
        ident_line(left, 52, 8)
        sh.text(right, 40, "%s %s" % (L["period_label"], period), 8, align="right")
    elif cont == "bar":
        sh.rect(0, 0, sh.w, 34, fill=H.get("color", DARK))
        sh.text(left, 22, L["bank"], 12, bold=True, color=H.get("ink", WHITE))
        ident_line(left, 50, 8)
        sh.text(right, 50, period, 8, align="right")
    elif cont == "courier":
        mid = (left + right) / 2.0
        sh.text(mid, 40, L["bank"].upper(), 10, bold=True, align="center")
        ident_line(left, 54, 8, upper=True)
        sh.text(right, 54, period, 8, align="right")
    elif cont == "minimal":
        sh.text(left, 44, "%s - continued" % L["product"], 8, bold=True)
    else:
        raise GenError("unknown continuation header %r" % cont)
    if H.get("top_page"):
        sh.text(right, 66 if cont != "minimal" else 44, "Page %d of %d" % (pno, npg), 8,
                align="right")


def small_print(L, which):
    B = BANKS[L["bank_id"]]
    check = ("Please check this statement carefully and tell us about anything that looks "
             "wrong within 30 days. Call 0800 000 000.")
    if B.get("fictional"):
        about = "%s is a fictional bank. This statement is synthetic test data." % B["name"]
    else:
        about = "Synthetic test data for software testing. Not issued by %s." % B["legal"]
    return {"auto": [check, about], "first": [check], "about": [about], "none": []}[which]


def draw_footer(sh, L, st, pno, npg):
    F = L["footer"]
    H = L["hdr"]
    left, right = H["left"], H["right"]
    ph = sh.h
    sz = F.get("size", 7)
    y = ph - F.get("y", 28)
    small = small_print(L, F.get("small", "auto"))
    label = F.get("fmt", "Page {p} of {n}").format(p=pno, n=npg)
    pos = F.get("pos", "right")
    if label:
        x = {"left": left, "center": (left + right) / 2.0, "right": right}[pos]
        sh.text(x, y, label, sz, align=pos)
    other = F.get("other")
    if other:
        other = other.format(branch=st["branch"], stno=st["stno"])
        if pos == "left":
            sh.text(right, y, other, sz, align="right")
        else:
            sh.text(left, y, other, sz)
    for k, s in enumerate(small):
        sh.text(left, y - (k + 1) * 8 - 2, s, 6)
    # On every page, in plain sight: this is not a real statement.
    sh.text(sh.w / 2.0 - sh.dx, ph - 12, SYNTHETIC, 7, bold=True, align="center")
    top = y - len(small) * 8 - 2 - sz
    if top < ph - L["bottom"] + 4:
        raise GenError("%s: footer reaches into the table area" % L["id"])


def draw_after(sh, L, st, y):
    A = L["after"]
    sz = L["size"]
    rows = st["rows"]
    vals = {
        "fees": sum(r["amt"] for r in rows if r["cat"] == "fee"),
        "interest": sum(r["amt"] for r in rows if r["cat"] == "interest"),
        "rwt": sum(r["amt"] for r in rows if r["cat"] == "rwt"),
        "count": len(rows),
    }
    x, w = A["x"], A["w"]
    x += nudge(L, [x + w - 4], x, x + w, "after-table block")
    sh.text(x, y + sz, A["title"], sz + 0.5, bold=True)
    for k, (lab, key) in enumerate(A["rows"]):
        yy = y + (k + 2) * (sz + 4)
        sh.text(x, yy, lab, sz)
        v = str(vals[key]) if key == "count" else mag(vals[key], L["thousands"], L["dollar"])
        sh.text(x + w - 4, yy, v, sz, align="right")


def draw_segblock(sh, L, st, seg, y):
    """An interest-and-fees table printed between two accounts' transactions."""
    A = L["segblock"]
    sz = L["size"]
    sr = [r for r in st["rows"] if r["seg"] == seg]
    cfg = st["segs"][seg]["cfg"]
    vals = {
        "fees": mag(sum(r["amt"] for r in sr if r["cat"] == "fee"), L["thousands"], L["dollar"]),
        "interest": mag(sum(r["amt"] for r in sr if r["cat"] == "interest"), L["thousands"],
                        L["dollar"]),
        "rate": cfg["rate_text"] or "0.00% p.a.",
    }
    x, w = A["x"], A["w"]
    x += nudge(L, [x + w - 4], x, x + w, "interest table")
    sh.text(x, y + sz, A["title"].format(product=cfg["product"]), sz, bold=True)
    sh.line(x, y + sz + 3, x + w, y + sz + 3, width=0.3)
    for k, (lab, key) in enumerate(A["rows"]):
        yy = y + (k + 2) * (sz + 4)
        sh.text(x, yy, lab, sz)
        sh.text(x + w - 4, yy, vals[key], sz, align="right")


def render_statement(sh, L, st):
    rows = st["rows"]
    npg = len(st["pages"])
    for k, r in enumerate(rows):
        r["zebra"] = (k % 2 == 1)
    bal_col = col_of(L, "balance")
    last_bal = {s: seg["opening"] for s, seg in enumerate(st["segs"])}
    for pi, items in enumerate(st["pages"]):
        sh.begin_page(st["dx"][pi])
        if pi == 0:
            draw_first_header(sh, L, st)
        else:
            draw_cont_header(sh, L, st, pi + 1, npg)
        sh.push_dx(st["tdx"][pi])
        prev_date = None
        y_min, y_max = None, None
        for kind, idx, y in items:
            if kind == "seg_title":
                seg = st["segs"][idx]
                _, b = sh.text(L["cols"][0].x0, y, seg["product"], L["size"] + 1, bold=True)
                if acct_printed(L):
                    draw_acct(sh, b + 12, y, L["size"] + 1, st, seg["acct"], bold=True)
                y_min = None
                prev_date = None
            elif kind == "head":
                draw_heading(sh, L, y)
                y_min = y - L["head_size"] - 3
            elif kind == "txn":
                r = rows[idx]
                show = not (L["date_once"] and prev_date == r["date"])
                yl = draw_txn(sh, L, st, idx, y, show)
                prev_date = r["date"]
                last_bal[r["seg"]] = r["bal"]
                y_max = yl
            elif kind in ("open", "bf", "cf"):
                seg = st["segs"][idx]
                if kind == "open":
                    closing_top = L["opening"].get("value") == "closing"
                    if L["opening"].get("date"):
                        cell(sh, L, col_of(L, "date"), y,
                             fmt_date(st["end"] if closing_top else st["start"], L["date_fmt"]))
                    draw_label(sh, L, y, L["opening"]["label"])
                    v = seg["closing"] if closing_top else seg["opening"]
                    fc = figure_col(L)
                else:
                    draw_label(sh, L, y, L["cf_bf"][0] if kind == "cf" else L["cf_bf"][1])
                    v = last_bal[idx]
                    fc = bal_col
                if fc.kind == "amount" and L["balance_kind"] != "card":
                    raise GenError("%s: an opening figure in a signed amount column" % L["id"])
                num, tok = fmt_cell(L, fc.kind, v)
                money_cell(sh, L, fc, y, num, tok)
                y_max = y
            elif kind == "ptot":
                draw_label(sh, L, y, L["page_totals"], bold=True)
                mine = [rows[j] for k2, j, _ in items if k2 == "txn" and rows[j]["seg"] == idx]
                for kk in ("debit", "credit"):
                    num, tok = fmt_cell(L, kk, sum(r[kk] or 0 for r in mine))
                    money_cell(sh, L, col_of(L, kk), y, num, tok, bold=True)
                y_max = y
            elif kind in ("totals", "totals_close"):
                seg = st["segs"][idx]
                sr = [r for r in rows if r["seg"] == idx]
                draw_label(sh, L, y, L["totals_label"], bold=True)
                for kk in ("debit", "credit"):
                    num, tok = fmt_cell(L, kk, sum(r[kk] or 0 for r in sr))
                    money_cell(sh, L, col_of(L, kk), y, num, tok, bold=True)
                if kind == "totals_close":
                    num, tok = fmt_cell(L, "balance", seg["closing"])
                    money_cell(sh, L, bal_col, y, num, tok, bold=True)
                y_max = y
            elif kind in ("close", "open_bal"):
                seg = st["segs"][idx]
                draw_label(sh, L, y, L["closing_label"] if kind == "close"
                           else L["open_trailer_label"], bold=(kind == "close"))
                fc = figure_col(L)
                num, tok = fmt_cell(L, fc.kind, seg["closing"] if kind == "close"
                                    else seg["opening"])
                money_cell(sh, L, fc, y, num, tok, bold=(kind == "close"))
                y_max = y
            elif kind == "segblock":
                draw_segblock(sh, L, st, idx, y)
            elif kind == "after":
                draw_after(sh, L, st, y)
            else:
                raise GenError("unknown page item %r" % kind)
            if kind == "head" and L["col_rules"]:
                st.setdefault("_rule_top", {})[pi] = y_min
        if L["col_rules"] and y_max is not None:
            top = st.get("_rule_top", {}).get(pi, y_min)
            if top is not None:
                for c in L["cols"][1:]:
                    sh.line(c.x0, top, c.x0, y_max + 4, width=0.3, color=(0.55, 0.55, 0.55))
        sh.pop_dx()
        draw_footer(sh, L, st, pi + 1, npg)
        sh.end_page()


# ---------------------------------------------------------------------------
# The truth, and the checks it must pass before it is written.
# ---------------------------------------------------------------------------

RED_NAME = {"desc": "description", "amount": "amount", "balance": "balance"}


def truth_rows(st, multi=False, stmt_index=None):
    """The rows a reader should find, in printed order, with what each redaction
    took away. Also returns, per kept row, whether a removed row came before it in
    date order (the arithmetic check cannot see through that)."""
    out, gap = [], []
    rows = st["rows"]
    for r in rows:
        red = r.get("redact") or {}
        if red.get("row") == "remove":
            continue
        t = {"date": r["date"].isoformat(), "description": r["desc"],
             "debit": None if r["debit"] is None else r["debit"] / 100.0,
             "credit": None if r["credit"] is None else r["credit"] / 100.0,
             "balance": r["bal"] / 100.0 if r["bal_printed"] else None}
        removed = [RED_NAME[f] for f, m in red.items() if m == "remove"]
        overlay = [RED_NAME[f] for f, m in red.items() if m == "overlay"]
        if "description" in removed:
            t["description"] = None
        if "amount" in removed:
            t["debit"] = t["credit"] = None
        if "balance" in removed:
            t["balance"] = None
        if removed:
            t["redacted"] = removed
        if overlay:
            t["overlay_redacted"] = overlay
        if multi:
            t["account_index"] = r["seg"]
        if stmt_index is not None:
            t["statement_index"] = stmt_index
        out.append(t)
    return out


def chain_check(name, rows, opening, closing, newest, gaps):
    """The truth's own arithmetic, read the way a scorer would read it: in date
    order, through every null it cannot see past (a removed amount, a removed row)
    the running balance is unknown until the next printed balance."""
    seq = list(zip(rows, gaps))
    if newest:
        seq = seq[::-1]
    bal = cents(opening)
    for k, (t, gap_before) in enumerate(seq):
        d, c = cents(t["debit"]), cents(t["credit"])
        if d is not None and c is not None:
            raise GenError("%s: a row with both a debit and a credit" % name)
        if (d is not None and d < 0) or (c is not None and c < 0):
            raise GenError("%s: a negative debit or credit" % name)
        if gap_before or bal is None or (d is None and c is None):
            bal = None
        else:
            bal = bal - (d or 0) + (c or 0)
        if t["balance"] is not None:
            if bal is not None and cents(t["balance"]) != bal:
                raise GenError("%s: printed balance %s, arithmetic says %s"
                               % (name, t["balance"], bal / 100.0))
            bal = cents(t["balance"])
        if not re.fullmatch(r"\d{4}-\d\d-\d\d", t["date"]):
            raise GenError("%s: bad date %r" % (name, t["date"]))
    if bal is not None and bal != cents(closing):
        raise GenError("%s: rows do not reach the closing balance" % name)


def gaps_for(st, seg):
    """For each kept row of one account, in printed order: was a removed row
    immediately before it in date order?"""
    rows = [r for r in st["rows"] if r["seg"] == seg]
    newest = st["newest"]
    seq = rows[::-1] if newest else rows
    flags, pending = {}, False
    for r in seq:
        if (r.get("redact") or {}).get("row") == "remove":
            pending = True
            continue
        flags[id(r)] = pending
        pending = False
    return [flags[id(r)] for r in rows if id(r) in flags]


def verify_internal(L, st, name):
    for s, seg in enumerate(st["segs"]):
        rows = [r for r in st["rows"] if r["seg"] == s]
        seq = rows[::-1] if st["newest"] else rows
        bal = seg["opening"]
        for r in seq:
            if (r["debit"] is None) == (r["credit"] is None):
                raise GenError("%s: row needs exactly one of debit/credit" % name)
            bal = bal - (r["debit"] or 0) + (r["credit"] or 0)
            if r["bal"] != bal:
                raise GenError("%s: internal balance drift" % name)
        if bal != seg["closing"]:
            raise GenError("%s: internal closing mismatch" % name)


def verify_printed(L, st, name):
    """What was drawn must read back to the truth."""
    card = L["balance_kind"] == "card"
    lo = st["start"] - dt.timedelta(days=45)
    hi = st["end"] + dt.timedelta(days=10)
    window = list(daterange(lo, hi))
    for k, r in enumerate(st["rows"]):
        red = r.get("redact") or {}
        rec = r.get("rec")
        if red.get("row") == "remove":
            if rec is not None:
                raise GenError("%s row %d: a removed row was drawn" % (name, k))
            continue
        if rec is None:
            raise GenError("%s row %d: never drawn" % (name, k))
        for key in ("date", "pdate"):
            s = rec[key]
            if s is None:
                continue
            hits = [d for d in window if fmt_date(d, L["date_fmt"]) == s]
            want = r["date"] if key == "date" else r["pdate"]
            if hits != [want]:
                raise GenError("%s row %d: printed %s %r reads as %s, truth %s"
                               % (name, k, key, s, hits, want))
        want_text = None if red.get("desc") == "remove" else r["desc"]
        got = " ".join(rec["text"]) if rec["text"] else None
        if got != want_text:
            raise GenError("%s row %d: printed %r, truth %r" % (name, k, got, want_text))
        m = rec["money"]
        for kk in ("debit", "credit"):
            if kk in m and parse_money(*m[kk]) != r[kk]:
                raise GenError("%s row %d: %s cell %r" % (name, k, kk, m[kk]))
        if "amount" in m:
            v = parse_money(*m["amount"], card=card)
            if v != (-r["amt"] if r["dir"] == "D" else r["amt"]):
                raise GenError("%s row %d: amount cell %r" % (name, k, m["amount"]))
        printed = [kk for kk in ("debit", "credit", "amount") if kk in m]
        if len(printed) != (0 if red.get("amount") == "remove" else 1):
            raise GenError("%s row %d: %d amount cells printed" % (name, k, len(printed)))
        want_bal = r["bal_printed"] and red.get("balance") != "remove"
        if ("balance" in m) != want_bal:
            raise GenError("%s row %d: balance printed/truth disagree" % (name, k))
        if "balance" in m and parse_money(*m["balance"]) != r["bal"]:
            raise GenError("%s row %d: balance cell %r" % (name, k, m["balance"]))
    if L["date_fmt"] in YEARLESS and "period" not in L["hdr"].get("det_rows", ["period"]):
        raise GenError("%s: a year-less date with no printed period" % name)


def features(L, st):
    f = ["bank:" + L["bank_id"], "cols:" + "|".join(L["kinds"]), "date:" + L["date_fmt"],
         "font:%s" % L["font"], "size:%g" % L["size"], "page:" + L["page"],
         "pages:%d" % len(st["pages"])]
    if L["date_fmt"] in YEARLESS:
        f.append("date_yearless")
    if st["start"].year != st["end"].year:
        f.append("period_crosses_year")
    f.append("money:thousands" if L["thousands"] else "money:no_thousands")
    if L["dollar"]:
        f.append("money:dollar")
    if "amount" in L["kinds"]:
        f.append("sign:" + L["amount_style"])
    if "balance" in L["kinds"]:
        f.append("bal:" + L["balance_style"])
    if L["hang"]:
        f.append("token_column")
    if L["bal_mode"] == "last_of_day":
        f.append("balance_last_of_day")
    if "balance" not in L["kinds"]:
        f.append("no_balance_column")
    f.append("heading_every_page" if L["head_every_page"] else "heading_page1_only")
    if any(len(c.head) > 1 for c in L["cols"]):
        f.append("two_line_heading")
    if any(c.align == "center" for c in L["cols"] if c.kind in MONEY_KINDS):
        f.append("money_centred")
    if L["opening"]:
        f.append("opening_line")
    if "close" in L["trailer"]:
        f.append("closing_line")
    if "totals" in L["trailer"] or "totals_close" in L["trailer"]:
        f.append("totals_line")
    if L["cf_bf"]:
        f.append("carried_forward")
    if L["page_totals"]:
        f.append("page_totals")
    if L["hdr"].get("summary"):
        f.append("summary_box")
    if L["hdr"].get("sidebar"):
        f.append("sidebar")
    if L["hdr"]["mast"] == "ib":
        f.append("id:no_masthead")
    if L["after"]:
        f.append("after_table_block")
    if L["sections"]:
        f.append("card_sections")
    if any(st["dx"]):
        f.append("page_x_offset")
    if any(st["tdx"]):
        f.append("table_x_offset")
    if L["head_shift"] or L["head_dx"]:
        f.append("heading_offset")
    if L["date_once"]:
        f.append("date_once_per_day")
    if L["zebra"]:
        f.append("zebra")
    if L["col_rules"] or L["row_rules"]:
        f.append("rules")
    if L["head_style"] in ("dark", "bar"):
        f.append("shaded_heading")
    if L["newest_first"]:
        f.append("newest_first")
    if L["money_line"] == "last":
        f.append("staggered_amounts")
    if L["kinds"][0] != "date":
        f.append("desc_before_date")
    if "balance" in L["kinds"] and "amount" in L["kinds"] \
            and L["kinds"].index("balance") < L["kinds"].index("amount"):
        f.append("balance_left_of_amount")
    if "fxamt" in L["kinds"]:
        f.append("fx_columns")
    if L["segments"]:
        f.append("multi_account")
    if L["segblock"]:
        f.append("mid_statement_table")
    if L["page"] == "A4L":
        f.append("landscape")
    rows = st["rows"]
    if any(r["nlines"] > 1 for r in rows):
        f.append("multiline_desc")
    if any(r.get("long") for r in rows):
        f.append("long_desc")
    if any(r["money_like"] for r in rows):
        f.append("money_like_desc")
    if "type" in L["kinds"]:
        f.append("type_column")
    if "code" in L["kinds"]:
        f.append("particulars_code_reference")
    if any(r["cat"] == "zero" for r in rows):
        f.append("zero_amounts")
    if any(abs(r["bal"]) >= 100000000 or r["amt"] >= 100000000 for r in rows):
        f.append("large_amounts")
    pr = page_rows(st["pages"])
    if len(pr) > 1 and min(len(x) for x in pr) <= 2:
        f.append("short_page")
    n_c = sum(1 for r in rows if r["dir"] == "C")
    if L["balance_kind"] != "card" and n_c <= 2:
        f.append("sparse_credits")
    if len(pr) > 1 and not any(rows[i]["dir"] == "C" for i in pr[0]):
        f.append("no_credits_page1")
    if any(r["bal"] < 0 for r in rows) and L["balance_kind"] != "card":
        f.append("negative_balance")
    reds = [m for r in rows for m in (r.get("redact") or {}).items()]
    if any(m == ("row", "remove") for m in reds):
        f.append("redacted_row")
    if any(m[1] == "remove" and m[0] != "row" for m in reds):
        f.append("redacted_removed")
    if any(m[1] == "overlay" for m in reds):
        f.append("redacted_overlay")
    if st["acct_mode"]:
        f.append("id:account_" + st["acct_mode"])
    if L["bank_id"] == "coop" or st["code"] == "06":
        f.append("id:code_shared_or_legacy")
    return f


def identity_keys(L, st):
    """What a reader can learn about the account from the page."""
    if not acct_printed(L):
        return None, None
    a = st["segs"][0]["acct"]
    mode = st["acct_mode"]
    if mode == "removed":
        return None, None
    return a[:2], acct_shown(st, a)


def write_json(path, obj):
    with open(path, "w") as f:
        json.dump(obj, f, indent=1, sort_keys=True)


def build_case(L, k, spec, out_dir):
    """Draw one statement (or one bundle of statements) and write its truth."""
    name = "%s_%d" % (L["id"], k + 1)
    seed = SPLIT_SEED[L["split"]] + zlib.crc32(("%s:%d" % (L["id"], k)).encode("utf-8"))
    pdf = os.path.join(out_dir, name + ".pdf")
    sh = Sheet(pdf, PAGES[L["page"]], L["font"])
    multi = bool(L["segments"])
    sts = []
    if L["bundle"]:
        count = random.Random(seed).randint(*L["bundle"])
        ident, opening = None, None
        y, m = spec["ym"][0], spec["ym"][1]
        for j in range(count):
            sub = dict(spec, ym=(y, m), n=max(8, spec["n"] + 3 * j - 4))
            st = build_statement(L, sub, seed + 7919 * j, ident=ident, forced_open=opening)
            ident = dict(st["ident"], stno=str(int(st["ident"]["stno"]) + 1))
            opening = st["closing"]
            m += 1
            if m > 12:
                y, m = y + 1, 1
            sts.append(st)
    else:
        sts.append(build_statement(L, spec, seed))
    rows, total_pages = [], 0
    for j, st in enumerate(sts):
        st["newest"] = L["newest_first"]
        render_statement(sh, L, st)
        total_pages += len(st["pages"])
        verify_internal(L, st, name)
        verify_printed(L, st, name)
        rows += truth_rows(st, multi=multi, stmt_index=j if L["bundle"] else None)
    sh.save()
    st0 = sts[0]
    code, shown = identity_keys(L, st0)
    feats = features(L, st0)
    if L["bundle"]:
        feats += ["bundle", "bundle:%d" % len(sts)]
    feats = [f for f in feats if not f.startswith("pages:")] + ["pages:%d" % total_pages]
    truth = {
        "case": name,
        "generator": GENERATOR,
        "note": "%s split. %s, %s (synthetic). %s" % (L["split"], L["bank"], L["product"], L["note"]),
        "bank": BANKS[L["bank_id"]]["name"],
        "layout": L["id"],
        "product": L["product"],
        "source_format": "pdf",
        "account_bank_code": code,
        "account_number": shown,
        "account_redaction": st0["acct_mode"],
        "features": feats,
        "row_order": "newest_first" if L["newest_first"] else "oldest_first",
        "opening_balance": None if multi else st0["opening"] / 100.0,
        "closing_balance": None if multi else sts[-1]["closing"] / 100.0,
        "removed_rows": sum(1 for st in sts for r in st["rows"]
                            if (r.get("redact") or {}).get("row") == "remove"),
        "row_count": len(rows),
        "rows": rows,
    }
    if multi:
        truth["accounts"] = [
            {"account_index": s, "product": seg["product"],
             "account_number": None if st0["acct_mode"] == "removed" else acct_shown(st0, seg["acct"]),
             "opening_balance": seg["opening"] / 100.0, "closing_balance": seg["closing"] / 100.0}
            for s, seg in enumerate(st0["segs"])]
    if L["bundle"]:
        truth["statements"] = [
            {"statement_index": j, "period_start": st["start"].isoformat(),
             "period_end": st["end"].isoformat(), "opening_balance": st["opening"] / 100.0,
             "closing_balance": st["closing"] / 100.0} for j, st in enumerate(sts)]
    verify_truth(truth, sts, L)
    write_json(os.path.join(out_dir, name + ".truth.json"), truth)
    return name, truth, total_pages


def verify_truth(truth, sts, L):
    name = truth["case"]
    if not 8 <= truth["row_count"] + truth["removed_rows"] <= 60 * (len(sts)):
        raise GenError("%s: row count %d" % (name, truth["row_count"]))
    if truth["row_count"] != len(truth["rows"]):
        raise GenError("%s: row_count disagrees with rows" % name)
    newest = truth["row_order"] == "newest_first"
    if truth.get("accounts"):
        st = sts[0]
        for s, acc in enumerate(truth["accounts"]):
            rows = [t for t in truth["rows"] if t["account_index"] == s]
            chain_check(name, rows, acc["opening_balance"], acc["closing_balance"], newest,
                        gaps_for(st, s))
    else:
        gaps = []
        for st in sts:
            gaps += gaps_for(st, 0)
        chain_check(name, truth["rows"], truth["opening_balance"], truth["closing_balance"],
                    newest, gaps)
    for t in truth["rows"]:
        if t["debit"] == 0.0 and t["description"] is not None \
                and "FEE WAIVED" not in t["description"]:
            raise GenError("%s: a zero debit that is not a waived fee" % name)


# ---------------------------------------------------------------------------
# Scans: the same statement with no text layer.
# ---------------------------------------------------------------------------

def make_scan(src, dst, seed):
    """Rasterise every page, add mild blur, noise and (sometimes) a skew under one
    degree, and wrap the images as a PDF with NO text layer. Returns the settings."""
    import pymupdf
    import numpy as np
    from PIL import Image, ImageFilter
    rng = random.Random(seed)
    dpi = rng.choice([200, 200, 240, 300])
    angle = rng.choice([0.0, 0.0, round(rng.uniform(-1.0, 1.0), 2)])
    blur = rng.choice([0.0, 0.4, 0.7])
    sigma = rng.choice([3.0, 5.0, 7.0])
    doc = pymupdf.open(src)
    out = pymupdf.open()
    for pno, page in enumerate(doc):
        pix = page.get_pixmap(dpi=dpi, colorspace=pymupdf.csGRAY)
        img = Image.frombytes("L", (pix.width, pix.height), pix.samples)
        if angle:
            img = img.rotate(angle + 0.1 * (pno % 2), resample=Image.BICUBIC, fillcolor=255)
        if blur:
            img = img.filter(ImageFilter.GaussianBlur(blur))
        arr = np.asarray(img, dtype=np.float32) * 0.96 + 6.0      # a little grey paper
        arr += np.random.default_rng(seed + pno).normal(0.0, sigma, arr.shape)
        img = Image.fromarray(np.clip(arr, 0, 255).astype(np.uint8), "L")
        buf = io.BytesIO()
        img.save(buf, "JPEG", quality=72)
        p = out.new_page(width=page.rect.width, height=page.rect.height)
        p.insert_image(p.rect, stream=buf.getvalue())
    out.set_metadata({"creator": GENERATOR, "producer": GENERATOR, "title": "",
                      "author": "", "subject": "", "keywords": "",
                      "creationDate": "D:20260101000000Z", "modDate": "D:20260101000000Z"})
    out.save(dst, garbage=3, deflate=True, no_new_id=True)
    out.close()
    doc.close()
    chk = pymupdf.open(dst)
    if any(pg.get_text().strip() for pg in chk):
        raise GenError("%s: the scan has a text layer" % dst)
    chk.close()
    return {"dpi": dpi, "skew": angle, "blur": blur, "noise": sigma}


def scan_truth(truth, settings):
    """A scan cannot read under a box: every overlay becomes a removal."""
    t = json.loads(json.dumps(truth))
    t["case"] = truth["case"] + "_scan"
    t["source_format"] = "scan"
    t["note"] = truth["note"] + " IMAGE-ONLY SCAN of %s.pdf (%d dpi, skew %.2f deg)." % (
        truth["case"], settings["dpi"], settings["skew"])
    t["features"] = truth["features"] + ["image_only_scan", "scan_dpi:%d" % settings["dpi"]] + (
        ["scan_skew"] if settings["skew"] else []) + (["scan_blur"] if settings["blur"] else [])
    for r in t["rows"]:
        ov = r.pop("overlay_redacted", [])
        for f in ov:
            if f == "description":
                r["description"] = None
            elif f == "amount":
                r["debit"] = r["credit"] = None
            elif f == "balance":
                r["balance"] = None
        if ov:
            r["redacted"] = sorted(set(r.get("redacted", []) + ov))
    if t["account_redaction"] == "overlay":
        t["account_bank_code"] = None
        t["account_number"] = None
        for a in t.get("accounts", []):
            a["account_number"] = None
    return t


# ---------------------------------------------------------------------------
# Exports: what internet banking hands you as CSV or Excel.
# ---------------------------------------------------------------------------

EXPORTS = []


def export_design(splits="both"):
    def deco(fn):
        EXPORTS.append((fn.__name__, splits, fn))
        return fn
    return deco


def export_cell(E, kind, r):
    if kind == "date":
        return fmt_date(r["date"], E["date_fmt"])
    if kind == "pdate":
        return fmt_date(r["pdate"], E["date_fmt"])
    if kind == "desc":
        return r["lines"][0]
    if kind in TEXT_KINDS:
        return r[kind]
    v = -r["amt"] if r["dir"] == "D" else r["amt"]
    if kind == "amount":
        return joined(*fmt_money(v, E["amount_style"], E))
    if kind == "amount_u":
        return mag(r["amt"], E["thousands"])
    if kind == "drcr":
        return "DR" if r["dir"] == "D" else "CR"
    if kind in ("debit", "credit"):
        return "" if r[kind] is None else mag(r[kind], E["thousands"])
    if kind == "balance":
        return joined(*fmt_money(r["bal"], E["balance_style"], E))
    if kind == "blank":
        return ""
    raise GenError("unknown export column %r" % kind)


def xl_value(E, kind, r):
    """The same cell as an Excel value: real dates and real numbers."""
    if kind in ("date", "pdate") and not E.get("dates_as_text"):
        d = r[kind]
        return dt.datetime(d.year, d.month, d.day)
    v = -r["amt"] if r["dir"] == "D" else r["amt"]
    if kind == "amount" and E["amount_style"] == "lead":
        return v / 100.0
    if kind == "amount_u":
        return r["amt"] / 100.0
    if kind in ("debit", "credit"):
        return None if r[kind] is None else r[kind] / 100.0
    if kind == "balance" and E["balance_style"] == "lead":
        return r["bal"] / 100.0
    s = export_cell(E, kind, r)
    return s if s != "" else None


def export_rows_back(E, table):
    """Read data rows back into (date, description, signed cents, balance cents),
    independently of how they were written."""
    kinds = [k for _, k in E["cols"]]
    out = []
    for cells in table:
        rec = dict(zip(kinds, cells))
        if any(isinstance(v, str) and "OPENING" in v.upper() for v in cells if v):
            continue
        dv = rec["date"]
        if isinstance(dv, dt.datetime):
            d = dv.date()
        else:
            hits = [x for x in daterange(dt.date(2024, 1, 1), dt.date(2028, 12, 31))
                    if fmt_date(x, E["date_fmt"]) == dv]
            if len(hits) != 1:
                raise GenError("%s: export date %r" % (E["id"], dv))
            d = hits[0]

        def money(x):
            if x is None or x == "":
                return None
            if isinstance(x, (int, float)):
                return int(round(x * 100))
            parts = str(x).split(" ")
            return parse_money(parts[0], parts[1] if len(parts) > 1 else "")
        if "amount" in rec:
            amt = money(rec["amount"])
        elif "amount_u" in rec:
            amt = money(rec["amount_u"]) * (-1 if rec["drcr"] == "DR" else 1)
        else:
            dv_, cv_ = money(rec.get("debit")), money(rec.get("credit"))
            amt = -dv_ if dv_ is not None else cv_
        text = [str(rec[k]) for k in kinds if k in TEXT_KINDS and rec[k] not in (None, "")]
        out.append((d, " ".join(text), amt, money(rec.get("balance"))))
    return out


def fixed_zip(raw, path):
    """Rewrite an .xlsx with fixed timestamps, so the file is byte-identical."""
    src = zipfile.ZipFile(io.BytesIO(raw))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = re.sub(rb"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\dZ", b"2026-01-01T00:00:00Z", data)
            zi = zipfile.ZipInfo(info.filename, date_time=(2026, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, data)


def build_export(E, k, spec, out_dir):
    name = "%s_%d" % (E["id"], k + 1)
    seed = SPLIT_SEED[E["split"]] + zlib.crc32(("export:%s:%d" % (E["id"], k)).encode("utf-8"))
    rng = random.Random(seed)
    kinds = [kd for _, kd in E["cols"]]
    ym = spec["ym"]
    start = dt.date(ym[0], ym[1], 1)
    end = add_months(start, 1) - dt.timedelta(days=1)
    st = {"start": start, "end": end}
    cfg = dict(catalog=E["catalog"], credits="normal", credit_p=E.get("credit_p", 0.17),
               balance_kind=E.get("balance_kind", "positive"), open_range=(300, 9000),
               scale=E.get("scale", 1.0), large=0, zero_rows=0, card=None,
               kinds=set(kinds), text_order=kinds)
    Lx = dict(id=E["id"], newest_first=E.get("newest_first", False), long_desc=0.0,
              wrap_indent=0.0, thousands=E["thousands"], bal_mode="every")
    n = spec["n"]
    code = E.get("code") or BANKS[E["bank"]]["codes"][0]
    acct = acct_number(rng, code)
    rows, opening, closing, _ = gen_rows(rng, Lx, cfg, st, [1] * n, [0] * n, [0] * n,
                                         lambda kd: None, None)
    for r in rows:
        r["bal_printed"] = "balance" in kinds
    p0, p1 = fmt_date(start, E["date_fmt"]), fmt_date(end, E["date_fmt"])
    fmtv = {"acct": acct, "p0": p0, "p1": p1, "open": mag(opening, False),
            "close": mag(closing, False), "name": rng.choice(CUSTOMERS["personal"])[0]}
    pre = [line.format(**fmtv) for line in E.get("preamble", [])]
    post = [line.format(**fmtv) for line in E.get("postamble", [])]
    headers = [h for h, _ in E["cols"]]
    table = [[export_cell(E, kd, r) for kd in kinds] for r in rows]
    if E.get("open_row"):
        orow = ["" for _ in kinds]
        orow[kinds.index("date")] = fmt_date(start, E["date_fmt"])
        orow[[i for i, kd in enumerate(kinds) if kd in TEXT_KINDS][0]] = "OPENING BALANCE"
        orow[kinds.index("balance")] = joined(*fmt_money(opening, E["balance_style"], E))
        table_out = [orow] + table if not Lx["newest_first"] else table + [orow]
    else:
        table_out = table
    fmt = E["fmt"]
    path = os.path.join(out_dir, name + "." + fmt)
    if fmt == "csv":
        sio = io.StringIO()
        w = csv.writer(sio, delimiter=E.get("delim", ","), lineterminator="\r\n",
                       quoting=csv.QUOTE_ALL if E.get("quote_all") else csv.QUOTE_MINIMAL)
        for line in pre:
            sio.write(line + "\r\n")
        w.writerow(headers)
        w.writerows(table_out)
        for line in post:
            sio.write(line + "\r\n")
        with open(path, "w", newline="", encoding="utf-8") as f:
            f.write(("\ufeff" if E.get("bom") else "") + sio.getvalue())
        # Read it back the way an importer would, from the header line on.
        with open(path, newline="", encoding="utf-8-sig") as f:
            lines = list(csv.reader(f, delimiter=E.get("delim", ",")))
        hi = next(i for i, l in enumerate(lines) if l == headers)
        back = [l for l in lines[hi + 1:hi + 1 + len(table_out)]]
    else:
        import openpyxl
        from openpyxl.styles import Font
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = E.get("sheet", "Transactions")
        ncol = len(kinds)
        ws.cell(row=1, column=1, value=E["title"].format(**fmtv)).font = Font(bold=True, size=13)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
        ws.cell(row=2, column=1, value="Period %s to %s" % (p0, p1))
        top = 4
        for j, h in enumerate(headers, 1):
            ws.cell(row=top, column=j, value=h).font = Font(bold=True)
        xrows = []
        for r in rows:
            xrows.append([xl_value(E, kd, r) for kd in kinds])
        for i, vals in enumerate(xrows, top + 1):
            for j, (v, kd) in enumerate(zip(vals, kinds), 1):
                c = ws.cell(row=i, column=j, value=v)
                if isinstance(v, dt.datetime):
                    c.number_format = E.get("xl_date", "dd/mm/yyyy")
                elif isinstance(v, float):
                    c.number_format = E.get("xl_money", "#,##0.00")
        notes = wb.create_sheet(E.get("other_sheet", "Notes"))
        notes.append(["Exported from internet banking", None])
        notes.append(["Interest rate", 2.5])
        notes.append(["Fees this period", 5.0])
        notes.append([SYNTHETIC, None])
        wb.properties.creator = GENERATOR
        wb.properties.created = dt.datetime(2026, 1, 1)
        wb.properties.modified = dt.datetime(2026, 1, 1)
        buf = io.BytesIO()
        wb.save(buf)
        fixed_zip(buf.getvalue(), path)
        wb2 = openpyxl.load_workbook(path)
        ws2 = wb2[E.get("sheet", "Transactions")]
        hdr_back = [ws2.cell(row=top, column=j).value for j in range(1, ncol + 1)]
        if hdr_back != headers:
            raise GenError("%s: header row reads back as %r" % (name, hdr_back))
        back = [[ws2.cell(row=i, column=j).value for j in range(1, ncol + 1)]
                for i in range(top + 1, top + 1 + len(rows))]
    got = export_rows_back(E, back)
    want = [(r["date"], r["desc"], -r["amt"] if r["dir"] == "D" else r["amt"],
             r["bal"] if "balance" in kinds else None) for r in rows]
    if got != want:
        raise GenError("%s: export does not read back to its truth" % name)
    truth = {
        "case": name,
        "generator": GENERATOR,
        "note": "%s split. %s internet-banking %s export (synthetic). %s"
                % (E["split"], BANKS[E["bank"]]["name"], fmt.upper(), E["note"]),
        "bank": BANKS[E["bank"]]["name"],
        "layout": E["id"],
        "product": E.get("product", "Everyday account"),
        "source_format": fmt,
        "account_bank_code": code if any("{acct}" in x for x in E.get("preamble", []) +
                                         [E.get("title", "")]) else None,
        "account_number": acct if any("{acct}" in x for x in E.get("preamble", []) +
                                      [E.get("title", "")]) else None,
        "account_redaction": None,
        "features": ["bank:" + E["bank"], "format:" + fmt, "date:" + E["date_fmt"],
                     "cols:" + "|".join(kinds)] + E.get("features", []) +
                    (["preamble"] if pre else []) + (["newest_first"] if Lx["newest_first"] else []) +
                    (["opening_row"] if E.get("open_row") else []),
        "row_order": "newest_first" if Lx["newest_first"] else "oldest_first",
        "opening_balance": opening / 100.0,
        "closing_balance": closing / 100.0,
        "removed_rows": 0,
        "row_count": len(rows),
        "rows": [{"date": r["date"].isoformat(), "description": r["desc"],
                  "debit": None if r["debit"] is None else r["debit"] / 100.0,
                  "credit": None if r["credit"] is None else r["credit"] / 100.0,
                  "balance": (r["bal"] / 100.0) if "balance" in kinds else None} for r in rows],
    }
    chain_check(name, truth["rows"], truth["opening_balance"], truth["closing_balance"],
                Lx["newest_first"], [False] * len(rows))
    write_json(os.path.join(out_dir, name + ".truth.json"), truth)
    return name, truth


def export(S, eid, bank, fmt, cols, date_fmt, count, **kw):
    E = dict(id=eid, split=S.split, bank=bank, fmt=fmt, cols=cols, date_fmt=date_fmt,
             thousands=False, dollar=False, amount_style="lead", balance_style="lead",
             catalog="bank", note="")
    E.update(kw)
    E["stmts"] = [dict(n=S.n(12, 40), ym=(S.v(2025, 2026), S.n(1, 12))) for _ in range(count)]
    return E


@export_design()
def anz_csv_export(S):
    return export(S, "anz_csv_export", "anz", "csv",
                  [("Type", "type"), ("Details", "desc"), ("Amt (NZD)", "amount"),
                   ("Txn Dt", "date"), ("ForeignCurrencyAmount", "blank")],
                  S.v("dd/mm/yyyy", "d/m/yyyy"), S.v(2, 1), catalog="typed",
                  note="Type column, signed amount, date AFTER the amount, an always-empty column.",
                  features=["signed", "blank_column"])


@export_design()
def asb_csv_export(S):
    return export(S, "asb_csv_export", "asb", "csv",
                  [("Date", "date"), ("Tran Type", "type"), ("Payee", "desc"),
                   ("Memo", "ref"), ("Amount", "amount")],
                  S.v("yyyy/mm/dd", "yyyymmdd"), S.v(1, 2), catalog="typed",
                  preamble=["Created date / time : 03 October 2026 / 10:15:00",
                            "Bank 12; Branch {acct}", "From date {p0}", "To date {p1}",
                            "Avail Bal : {close}", "Ledger Balance : {close}", ""],
                  note="Seven preamble lines (with balances in them) before the header.",
                  features=["signed"])


@export_design()
def bnz_xlsx_export(S):
    return export(S, "bnz_xlsx_export", "bnz", "xlsx",
                  [("Date", "date"), ("Payee", "payee"), ("Particulars", "part"), ("Code", "code"),
                   ("Reference", "ref"), ("Dr", "debit"), ("Cr", "credit"),
                   ("Running Bal", "balance")],
                  S.v("dd/mm/yyyy", "d/m/yyyy"), S.v(2, 1), catalog="payee_pcr",
                  title="BNZ - Transactions - {acct}", other_sheet="Account details",
                  note="Excel: merged title row, real date and number cells, Dr/Cr columns, "
                       "a second irrelevant sheet.",
                  features=["dr_cr_columns", "merged_title", "second_sheet"])


@export_design()
def westpac_csv_export(S):
    return export(S, "westpac_csv_export", "westpac", "csv",
                  [("Txn Date", "date"), ("Narrative", "desc"), ("Amount", "amount_u"),
                   ("DR/CR", "drcr"), ("Running Balance", "balance")],
                  S.v("dd-Mon-yyyy", "dd Mon yyyy"), S.v(1, 2), thousands=S.v(False, True),
                  preamble=["Account,{acct}", "Statement period,{p0} - {p1}", ""],
                  quote_all=S.v(False, True), open_row=True,
                  note="Unsigned amount with a separate DR/CR column; an OPENING BALANCE data row.",
                  features=["drcr_indicator"])


@export_design()
def kiwibank_xlsx_export(S):
    return export(S, "kiwibank_xlsx_export", "kiwibank", "xlsx",
                  [("Txn Date", "date"), ("Narrative", "desc"), ("Amt (NZD)", "amount"),
                   ("Running Bal", "balance")],
                  S.v("dd/mm/yyyy", "dd-mm-yyyy"), S.v(2, 1), newest_first=True,
                  title="Kiwibank account {acct} - {name}", sheet="Statement",
                  xl_money=S.v("#,##0.00", '"$"#,##0.00'),
                  note="Excel, newest first, signed amounts as numbers.",
                  features=["signed", "merged_title", "second_sheet"])


@export_design()
def tsb_csv_export(S):
    return export(S, "tsb_csv_export", "tsb", "csv",
                  [("Date", "date"), ("Transaction Description", "desc"), ("Paid Out", "debit"),
                   ("Paid In", "credit"), ("Bal", "balance")],
                  S.v("yyyymmdd", "yyyy-mm-dd"), S.v(1, 2), thousands=True, quote_all=True,
                  bom=True, preamble=["TSB account {acct}", ""],
                  note="Quoted fields with thousands separators, a byte-order mark.",
                  features=["dr_cr_columns", "bom"])


@export_design()
def coop_csv_export(S):
    return export(S, "coop_csv_export", "coop", "csv",
                  [("Txn Dt", "date"), ("Other Party", "payee"), ("Particulars", "part"),
                   ("Reference", "ref"), ("Amount", "amount"), ("Balance", "balance")],
                  S.v("d/m/yy", "dd/mm/yy"), S.v(2, 1), catalog="payee_pcr",
                  amount_style="drcr", balance_style="drcr",
                  postamble=["", "Closing balance,{close}"],
                  note="Amounts and balances with a DR/CR suffix; a closing line after the data.",
                  features=["drcr_suffix", "postamble"])


@export_design()
def rimu_xlsx_export(S):
    return export(S, "rimu_xlsx_export", "rimu", "xlsx",
                  [("Processed", "date"), ("Transaction Description", "desc"),
                   ("Money Out", "debit"), ("Money In", "credit"), ("Running Balance", "balance")],
                  "dd/mm/yyyy", S.v(1, 2), dates_as_text=S.v(True, False),
                  title="Rimu Bank transactions {acct}",
                  note="Excel with Money Out / Money In columns.",
                  features=["dr_cr_columns", "merged_title", "second_sheet"])


# ---------------------------------------------------------------------------
# THE DESIGNS, by bank. Each is a plausible statement design; S.v / S.ch / S.r
# give dev one set of parameters and holdout a disjoint one.
# ---------------------------------------------------------------------------

ARCH = []


def design(splits="both"):
    def deco(fn):
        ARCH.append((fn.__name__, splits, fn))
        return fn
    return deco


def pile(S, count, n_rng, first, months=1, tweaks=None):
    """`count` statements of one design: different months, sizes and mixes."""
    out = []
    y, m = first[0], first[1]
    d = first[2] if len(first) > 2 else 1
    for j in range(count):
        spec = dict(n=S.n(*n_rng), ym=(y, m, d), months=months)
        spec.update((tweaks or {}).get(j, {}))
        out.append(spec)
        m += months * S.rng.choice([1, 1, 2, 3])
        while m > 12:
            m -= 12
            y += 1
    return out


# ---- ANZ (01, 06) ------------------------------------------------------------

@design()
def anz_everyday(S):
    size = S.v(8, 8.5)
    left, right = S.r((36, 42), (48, 54)), S.r((553, 558), (540, 546))
    two = S.v(False, True)

    def hd(w):
        return (w, "($)") if two else (w,)
    cols = mkcols(left, right, [
        ("date", S.r((38, 42), (32, 36)), ("Date",)),
        ("desc", None, S.v(("Transaction type and details",), ("Transaction details",))),
        ("debit", S.r((64, 70), (72, 78)), hd("Withdrawals")),
        ("credit", S.r((62, 68), (66, 72)), hd("Deposits")),
        ("balance", S.r((72, 80), (80, 86)), hd("Balance")),
    ])
    sx = S.r((318, 326), (296, 304))
    return layout(
        S, "anz_everyday", bank="anz", product="Everyday account",
        note="Type line over details line, balance printed only on the last transaction of "
             "each day, page totals and a period totals line.",
        font="Helvetica", size=size, pitch=round(size * S.v(1.75, 1.6), 2), cols=cols,
        date_fmt=S.v("dd Mon", "d Mon"), period_fmt="long", catalog="bank", wrap=0.4,
        bal_mode="last_of_day", opening=dict(label="OPENING BALANCE", date=S.v(True, False)),
        page_totals="Totals at end of page", trailer=["totals_close"],
        totals_label="Totals at end of period", head_style="rule",
        hdr=dict(mast="bar", color=DARK, title="Account statement", left=left, right=right,
                 addr=(left, 100), det=(sx, 96, 84), det_rows=["product", "acct", "period", "stno"],
                 summary=dict(style="box", x=sx, y=154, w=right - sx - 8, title="Account summary",
                              rows=[("Opening balance", "open"), ("Total withdrawals", "tot_d"),
                                    ("Total deposits", "tot_c"), ("Closing balance", "close")]),
                 table_title="Account transactions"),
        table_top1=272, table_top2=84, bottom=S.v(70, 74),
        footer=dict(pos="right", other="ANZ Bank New Zealand Limited"),
        stmts=pile(S, 4, (18, 40), S.v((2025, 9), (2026, 1)),
                   tweaks={1: dict(short_last=True), 2: dict(code="06")}))


@design()
def anz_visa(S):
    size = S.v(8, 7.5)
    left, right = S.r((40, 44), (30, 34)), S.r((550, 556), (562, 566))
    two = S.v(True, False)
    cols = mkcols(left, right, [
        ("date", S.r((52, 56), (40, 44)), ("Transaction", "date") if two else ("Date",)),
        ("pdate", S.r((46, 50), (44, 48)), ("Processed", "date") if two else ("Processed",)),
        ("desc", None, ("Transaction details",)),
        ("amount", S.r((74, 80), (64, 70)), ("Amount", "NZ$") if two else ("Amount $",)),
    ])
    return layout(
        S, "anz_visa", bank="anz", product="Visa credit card",
        note="Credit card: transaction and processed dates, purchases plain, payments marked "
             "CR, foreign-currency line under overseas purchases, no running balance.",
        font="Helvetica", size=size, cols=cols, date_fmt=S.v("dd Mon", "dd MON"),
        period_fmt=S.v("long", "short"), catalog="card", credits="card", balance_kind="card",
        bal_mode="none", amount_style="card_cr", hang=S.v((), ("amount",)), wrap=0.3,
        opening=S.v(dict(label="Opening balance"), None), trailer=["close"],
        closing_label="Closing balance", head_style=S.v("bar", "rule"),
        card=dict(limit=S.ch([500000, 800000], [1000000, 1200000])), rate_text="20.95% p.a.",
        hdr=dict(mast="bar", color=DARK, title="Credit card statement", left=left, right=right,
                 addr=(left, 104), det=(S.r((300, 310), (330, 336)), 100, 76),
                 det_rows=["card", "period", "issued"],
                 summary=dict(style="box", x=left, y=S.v(170, 176), w=S.r((220, 236), (240, 250)),
                              title="Account summary",
                              rows=[("Opening balance", "open"), ("Purchases & debits", "tot_d"),
                                    ("Payments & credits", "tot_c"), ("Closing balance", "close")]),
                 sidebar=dict(style="shaded", x=S.r((330, 340), (344, 350)), y=S.v(170, 176),
                              w=S.r((200, 210), (206, 212)), title="Payment due",
                              rows=[("Minimum payment", "minpay"), ("Due date", "due"),
                                    ("Credit limit", "limit"), ("Available credit", "avail"),
                                    ("Purchase rate", "rate")])),
        table_top1=S.v(286, 292), table_top2=S.v(86, 92),
        footer=dict(pos=S.v("center", "left")),
        stmts=pile(S, 3, (10, 34), S.v((2025, 10, 17), (2026, 2, 9))))


@design()
def anz_homeloan(S):
    size = S.v(9, 8.5)
    left, right = S.r((40, 44), (48, 52)), S.r((552, 556), (544, 548))
    cols = mkcols(left, right, [
        ("date", S.r((88, 92), (64, 68)), ("Date",)),
        ("desc", None, ("Transaction details",)),
        ("debit", S.r((72, 78), (76, 82)), ("Debit",)),
        ("credit", S.r((72, 78), (76, 82)), ("Credit",)),
        ("balance", S.r((92, 98), (96, 102)), ("Balance",)),
    ])
    return layout(
        S, "anz_homeloan", bank="anz", product="Home loan",
        note="Loan account: the balance is owed throughout and printed with a DR token; "
             "repayments are credits, interest is a debit; large figures over a quarter.",
        font="Helvetica", size=size, cols=cols, date_fmt=S.v("d Month yyyy", "dd Mon yyyy"),
        period_fmt="long", catalog="loan", credits="loan", balance_kind="loan",
        balance_style="dr_only", opening=dict(label="Opening balance"), trailer=["close"],
        head_style="bar", rate_text=S.v("6.24% p.a.", "5.99% p.a."),
        hdr=dict(mast="bar", color=DARK, title="Home loan statement", left=left, right=right,
                 addr=(left, 100), det=(S.r((320, 330), (300, 310)), 96, 84),
                 det_rows=["acct", "period", "issued"],
                 summary=dict(style="box", x=S.r((320, 330), (300, 310)), y=146,
                              w=S.r((200, 208), (226, 236)), title="Loan summary",
                              rows=[("Opening balance", "open"), ("Interest and fees", "tot_d"),
                                    ("Repayments", "tot_c"), ("Closing balance", "close"),
                                    ("Interest rate", "rate")])),
        table_top1=S.v(266, 272), table_top2=84,
        footer=dict(pos="right"),
        stmts=pile(S, 3, (8, 14), S.v((2025, 4), (2025, 7)), months=3,
                   tweaks={1: dict(code="06")}))


@design()
def anz_business_visa(S):
    size = S.v(8, 7.5)
    left, right = S.r((40, 44), (48, 52)), S.r((552, 556), (544, 548))
    cols = mkcols(left, right, [
        ("date", S.r((40, 44), (34, 38)), ("Date",)),
        ("pdate", S.r((48, 52), (42, 46)), ("Processed",)),
        ("desc", None, ("Transaction details",)),
        ("amount", S.r((72, 78), (76, 82)), ("Amount",)),
    ])
    return layout(
        S, "anz_business_visa", bank="anz", product="Business Visa",
        note="Business credit card with a section per cardholder: a cardholder heading and a "
             "card sub-total line inside the table (neither is a transaction).",
        font="Helvetica", size=size, cols=cols, date_fmt=S.v("dd Mon", "d Mon"),
        period_fmt="short", catalog="card", credits="card", balance_kind="card",
        bal_mode="none", amount_style="card_cr", hang=S.v(("amount",), ()), wrap=0.2,
        customers="biz", sections=["J SAMPLE  {card}", "K EXAMPLE  {card}"],
        trailer=["close"], closing_label="Closing balance", head_style="bar",
        card=dict(limit=2000000), rate_text="18.95% p.a.",
        hdr=dict(mast="bar", color=DARK, title="Business card statement", left=left,
                 right=right, addr=(left, 100), addr_bold=True,
                 det=(S.r((316, 324), (300, 306)), 100, 78), det_rows=["card", "period", "issued"],
                 summary=dict(style="box", x=S.r((316, 324), (300, 306)), y=144,
                              w=S.r((200, 208), (226, 236)), title="Account summary",
                              rows=[("Opening balance", "open"), ("Purchases", "tot_d"),
                                    ("Payments", "tot_c"), ("Closing balance", "close"),
                                    ("Credit limit", "limit")])),
        table_top1=S.v(260, 254), table_top2=84,
        footer=dict(pos="right"),
        stmts=pile(S, 3, (24, 40), S.v((2025, 11, 5), (2026, 3, 5))))


@design()
def anz_combined(S):
    size = S.v(8, 8.5)
    left, right = S.r((40, 44), (48, 52)), S.r((552, 556), (544, 548))
    cols = mkcols(left, right, [
        ("date", S.r((40, 44), (44, 48)), ("Date",)),
        ("desc", None, ("Transaction details",)),
        ("debit", S.r((66, 72), (70, 76)), S.v(("Debits",), ("Paid Out",))),
        ("credit", S.r((66, 72), (66, 72)), S.v(("Credits",), ("Paid In",))),
        ("balance", S.r((74, 80), (78, 84)), S.v(("Balance",), ("Running Balance",))),
    ])
    return layout(
        S, "anz_combined", bank="anz", product="Combined statement",
        note="Several accounts in one statement: each its own sub-table with its own opening "
             "and closing, and an interest-and-fees table printed between them.",
        font="Helvetica", size=size, cols=cols, date_fmt=S.v("dd Mon", "dd MON"),
        period_fmt="long", catalog="bank", wrap=0.15, opening=dict(label="Opening balance"),
        trailer=["close"], head_style="rule",
        segments=[dict(product="Everyday account", share=0.65),
                  dict(product="Online saver", catalog="savings", credit_p=0.35,
                       open_range=(2000, 15000), rate_text=S.v("3.40% p.a.", "3.15% p.a."), share=0.35)],
        segblock=dict(title="Interest and fees - {product}", x=S.r((60, 70), (300, 310)),
                      w=S.r((200, 210), (190, 196)),
                      rows=[("Interest rate", "rate"), ("Credit interest", "interest"),
                            ("Account fees", "fees")]),
        hdr=dict(mast="bar", color=DARK, title="Combined statement", left=left, right=right,
                 addr=(left, 100), det=(S.r((318, 326), (296, 304)), 96, 84),
                 det_rows=["name", "period", "stno"],
                 summary=dict(style="box", x=S.r((318, 326), (296, 304)), y=146,
                              w=S.r((200, 210), (226, 236)), title="Your accounts",
                              rows=[("Everyday account", "close:0"),
                                    ("Online saver", "close:1")])),
        table_top1=240, table_top2=84,
        footer=dict(pos="right"),
        stmts=pile(S, 3, (30, 48), S.v((2025, 8), (2026, 5)),
                   tweaks={2: dict(code="06")}))


# ---- ASB (12) ----------------------------------------------------------------

@design()
def asb_everyday(S):
    font, size = S.v(("Helvetica", 7.5), ("Times", 8))
    left, right = S.r((30, 34), (40, 46)), S.r((562, 566), (552, 556))
    cols = mkcols(left, right, [
        ("date", S.r((44, 48), (36, 40)), ("Date",)),
        ("part", None, ("Particulars",)),
        ("code", S.r((58, 62), (64, 70)), ("Code",)),
        ("ref", S.r((74, 80), (66, 72)), ("Reference",)),
        ("debit", S.r((60, 64), (66, 70)), S.v(("Withdrawals",), ("Debit",))),
        ("credit", S.r((58, 62), (62, 66)), S.v(("Deposits",), ("Credit",))),
        ("balance", S.r((66, 72), (70, 76)), ("Balance",)),
    ])
    return layout(
        S, "asb_everyday", bank="asb", product="Everyday account",
        note="NZ layout with Particulars, Code and Reference as separate text columns "
             "(12-character fields), shaded alternate rows.",
        font=font, size=size, cols=cols, date_fmt=S.v("dd/mm/yyyy", "dd/mm/yy"),
        period_fmt="num", period_sep=" - ", catalog="pcr", zebra=True,
        head_every_page=S.v(True, False), head_style="dark", head_fill=(0.15, 0.15, 0.15),
        opening=dict(label="Opening Balance"), trailer=["close"], closing_label="Closing Balance",
        hdr=dict(mast="bar", color=LIGHT, ink=BLACK, title="Statement of account",
                 left=left, right=right, addr=(left, S.v(96, 104)),
                 det=(S.r((320, 330), (290, 300)), S.v(96, 104), 80),
                 det_rows=["name", "acct", "period", "branch"],
                 summary=dict(style="band", x=left, y=S.v(166, 172), w=right - left,
                              rows=[("Opening balance", "open"), ("Total debits", "tot_d"),
                                    ("Total credits", "tot_c"), ("Closing balance", "close")])),
        table_top1=S.v(236, 244), table_top2=S.v(80, 86),
        footer=dict(pos="right", other="Branch {branch}"),
        stmts=pile(S, 4, (12, 48), S.v((2025, 12), (2026, 2))))


@design()
def asb_cheque(S):
    size = S.v(9, 8.5)
    left, right = S.r((40, 44), (48, 52)), S.r((554, 558), (546, 550))
    cols = mkcols(left, right, [
        ("date", S.r((62, 66), (60, 64)), ("Date",), None, "center"),
        ("desc", None, S.v(("Description",), ("Transaction Description",))),
        ("debit", S.r((70, 76), (74, 80)), ("Debit",), "center"),
        ("credit", S.r((70, 76), (74, 80)), ("Credit",), "center"),
        ("balance", S.r((80, 86), (82, 88)), ("Balance",), "center"),
    ])
    return layout(
        S, "asb_cheque", bank="asb", product="Cheque account",
        note="Money centred in its columns rather than right-aligned; long descriptions.",
        font="Helvetica", size=size, cols=cols, date_fmt=S.v("d Mon yyyy", "dd/mm/yyyy"),
        period_fmt="short", catalog="bank", long_desc=0.3, head_style="dark",
        head_fill=(0.2, 0.2, 0.2), opening=dict(label="Opening balance"), trailer=["close"],
        hdr=dict(mast="bar", color=LIGHT, ink=BLACK, title="Account statement", left=left,
                 right=right, addr=(left, 98), det=(S.r((316, 324), (300, 306)), 98, 84),
                 det_rows=["acct", "period", "stno"],
                 summary=dict(style="band", x=left, y=158, w=right - left,
                              rows=[("Opening", "open"), ("Debits", "tot_d"),
                                    ("Credits", "tot_c"), ("Closing", "close")])),
        table_top1=226, table_top2=84,
        footer=dict(pos="right"),
        stmts=pile(S, 3, (20, 34), S.v((2025, 10), (2026, 8))))


@design("dev")
def asb_business(S):
    size = 7.5
    left, right = 36, 560
    cols = mkcols(left, right, [
        ("date", 46, ("Date",)),
        ("payee", None, ("Payee / payer",)),
        ("part", 66, ("Particulars",)),
        ("code", 56, ("Code",)),
        ("ref", 62, ("Reference",)),
        ("debit", 62, ("Payments",)),
        ("credit", 62, ("Receipts",)),
        ("balance", 70, ("Balance",)),
    ])
    return layout(
        S, "asb_business", bank="asb", product="Business current account",
        note="Portrait business statement with Payee plus Particulars / Code / Reference.",
        font="Helvetica", size=size, cols=cols, date_fmt="dd/mm/yyyy", period_fmt="num",
        period_sep=" - ", catalog="biz_pcr", customers="biz", scale=0.5, zebra=True,
        head_style="dark", head_fill=(0.15, 0.15, 0.15), opening=dict(label="Opening balance"),
        trailer=["totals_close"], totals_label="Totals",
        hdr=dict(mast="bar", color=LIGHT, ink=BLACK, title="Business statement", left=left,
                 right=right, addr=(left, 96), det=(320, 96, 80),
                 det_rows=["name", "acct", "period", "branch"]),
        table_top1=176, table_top2=80,
        footer=dict(pos="right", other="Branch {branch}"),
        stmts=pile(S, 3, (30, 50), (2026, 3)))


@design("holdout")
def asb_personal_cheque(S):
    size = 8
    left, right = 44, 552
    cols = mkcols(left, right, [
        ("date", 52, ("Date",)),
        ("part", None, ("Particulars",)),
        ("code", 64, ("Code",)),
        ("ref", 70, ("Reference",)),
        ("debit", 64, ("Payments", "$")),
        ("credit", 64, ("Receipts", "$")),
        ("balance", 72, ("Balance", "$")),
    ])
    return layout(
        S, "asb_personal_cheque", bank="asb", product="Personal cheque account",
        note="Particulars / Code / Reference with two-line money headings; balance with OD.",
        font="Times", size=size, cols=cols, date_fmt="d/mm/yyyy", period_fmt="long",
        catalog="pcr", balance_style="od", hang=("balance",), balance_kind="overdraft",
        opening=dict(label="Opening balance"), trailer=["close"], head_style="rule",
        head_every_page=False,
        hdr=dict(mast="right", title="Statement", left=left, right=right, addr=(left, 92),
                 det=(left, 150, 90), det_rows=["acct", "period", "stno"],
                 summary=dict(style="box", x=330, y=88, w=200, title="Summary",
                              rows=[("Opening balance", "open"), ("Payments", "tot_d"),
                                    ("Receipts", "tot_c"), ("Closing balance", "close")])),
        table_top1=228, table_top2=72,
        footer=dict(pos="center"),
        stmts=pile(S, 3, (26, 42), (2026, 4)))


@design()
def asb_online_printout(S):
    size = S.v(8, 8.5)
    left, right = S.r((36, 40), (46, 50)), S.r((556, 560), (546, 550))
    cols = mkcols(left, right, [
        ("date", S.r((50, 54), (44, 48)), S.v(("Txn Date",), ("Date",))),
        ("desc", None, S.v(("Narrative",), ("Transaction Description",))),
        ("debit", S.r((62, 68), (58, 62)), S.v(("Debit",), ("Dr",))),
        ("credit", S.r((62, 68), (58, 62)), S.v(("Credit",), ("Cr",))),
        ("balance", S.r((76, 82), (72, 78)), S.v(("Running Balance",), ("Bal",))),
    ])
    return layout(
        S, "asb_online_printout", bank="asb", product="Internet banking printout",
        note="Internet-banking printout: no masthead, no bank name anywhere (only the 12 "
             "account code says ASB), newest transaction first, opening balance at the bottom.",
        font="Helvetica", size=size, cols=cols, date_fmt=S.v("dd/mm/yyyy", "d/mm/yyyy"),
        period_fmt="num", catalog="bank", newest_first=True, row_rules=True,
        opening=dict(label="Closing balance", value="closing"), trailer=["open_bal"],
        head_style="bar", head_fill=(0.92, 0.92, 0.92),
        hdr=dict(mast="ib", title=S.v("Account transactions", "Transaction history"),
                 left=left, right=right, addr=None, det=(left, 70, 90),
                 det_rows=["name", "acct", "period", "printed"], cont="minimal"),
        table_top1=150, table_top2=70,
        footer=dict(pos="right", small="none"),
        stmts=pile(S, 4, (14, 40), S.v((2025, 7), (2026, 1))))


@design()
def asb_business_landscape(S):
    size = S.v(8, 7.5)
    left = S.r((34, 38), (44, 48))
    right = S.r((636, 640), (626, 630))
    cols = mkcols(left, right, [
        ("date", S.r((48, 52), (42, 46)), ("Date",)),
        ("payee", None, S.v(("Payee",), ("Other party",))),
        ("part", S.r((72, 76), (68, 72)), ("Particulars",)),
        ("code", S.r((64, 68), (66, 70)), ("Code",)),
        ("ref", S.r((68, 72), (70, 74)), ("Reference",)),
        ("debit", S.r((70, 74), (66, 70)), S.v(("Payments",), ("Withdrawals",))),
        ("credit", S.r((70, 74), (66, 70)), S.v(("Receipts",), ("Deposits",))),
        ("balance", S.r((76, 80), (74, 78)), ("Balance",)),
    ])
    sx = right + S.r((22, 26), (28, 32))
    return layout(
        S, "asb_business_landscape", bank="asb", product="Business cheque account",
        note="Landscape business statement: Payee plus Particulars / Code / Reference, "
             "Payments and Receipts, and a sidebar beside the table on page 1.",
        page="A4L", font="Helvetica", size=size, cols=cols, date_fmt=S.v("dd/mm/yyyy", "dd/mm/yy"),
        period_fmt="long", catalog="biz_pcr", customers="biz", scale=0.8,
        col_rules=S.v(True, False), head_style="box", trailer=["totals_close"],
        totals_label="Closing totals", opening=dict(label="Opening balance"),
        rate_text="5.10% p.a.",
        hdr=dict(mast="bar", color=MID, title="Business account statement", left=left,
                 right=842 - left, addr=(left, 92), det=(S.r((330, 340), (360, 370)), 92, 90),
                 det_rows=["name", "acct", "period", "stno"],
                 sidebar=dict(style="box", x=sx, y=S.v(176, 184), w=842 - left - sx,
                              title="Your banker", size=7,
                              rows=[("Overdraft rate", "rate"), ("Fees", "fees"),
                                    ("Transactions", "count")],
                              notes=["Business line 0800 000 001", "Mon-Fri 8am-6pm"])),
        table_top1=S.v(176, 184), table_top2=S.v(74, 80),
        footer=dict(pos="right", other="{branch}", small="about"),
        stmts=pile(S, 3, (30, 56), S.v((2026, 1), (2025, 11))))


# ---- BNZ (02) ----------------------------------------------------------------

@design()
def bnz_everyday(S):
    size = S.v(9, 8)
    left, right = S.r((46, 52), (36, 40)), S.r((548, 552), (556, 560))
    cols = mkcols(left, right, [
        ("date", S.r((58, 62), (60, 64)), ("Date",)),
        ("desc", None, S.v(("Transactions",), ("Description",))),
        ("debit", S.r((70, 76), (64, 68)), ("Withdrawals",)),
        ("credit", S.r((66, 72), (62, 66)), ("Deposits",)),
        ("balance", S.r((76, 82), (72, 76)), ("Balance",)),
    ])
    return layout(
        S, "bnz_everyday", bank="bnz", product="Everyday account",
        note="Date printed once per day, balance brought forward / carried forward at page "
             "breaks, summary band across the page.",
        font="Helvetica", size=size, cols=cols, date_fmt=S.v("d Mon yyyy", "dd Mon yyyy"),
        period_fmt="short", catalog="bank", wrap=0.2, date_once=True,
        cf_bf=S.v(("Balance carried forward", "Balance brought forward"),
                  ("Carried forward", "Brought forward")),
        opening=dict(label="Balance brought forward"), trailer=["close"],
        pitch=round(size * S.v(2.0, 2.2), 2), max_rows=S.v(None, 26), head_style="rule2",
        hdr=dict(mast="logo", title="Statement", left=left, right=right,
                 addr=(left, 92), det=(S.r((310, 316), (330, 336)), 92, S.v(86, 80)),
                 det_rows=["acct", "period", "stno"],
                 summary=dict(style="band", x=left, y=S.v(150, 156), w=right - left,
                              rows=[("Opening balance", "open"), ("Withdrawals", "tot_d"),
                                    ("Deposits", "tot_c"), ("Closing balance", "close")]),
                 cont="line", top_page=True),
        table_top1=S.v(222, 230), table_top2=S.v(92, 96),
        footer=dict(pos="left", fmt="Page {p}"),
        stmts=pile(S, 4, (25, 58), S.v((2025, 6), (2026, 3)), tweaks={1: dict(short_last=True)}))


@design()
def bnz_saver(S):
    font, size = S.v(("Times", 10), ("Helvetica", 9.5))
    left, right = S.r((56, 60), (46, 50)), S.r((540, 544), (550, 554))
    cols = mkcols(left, right, [
        ("date", S.r((50, 54), (54, 58)), ("Date",)),
        ("desc", None, ("Details",)),
        ("debit", S.r((70, 76), (76, 80)), ("Withdrawals",)),
        ("credit", S.r((64, 70), (64, 68)), ("Deposits",)),
        ("balance", S.r((74, 80), (78, 84)), ("Balance",)),
    ])
    return layout(
        S, "bnz_saver", bank="bnz", product="Online saver",
        note="Savings account where the Deposits column is nearly empty: one or two "
             "deposits, none on page 1 in some statements.",
        font=font, size=size, cols=cols, date_fmt=S.v("dd Mon", "dd MON"), period_fmt="short",
        catalog="savings", credits="sparse", opening=dict(label="Opening balance"),
        trailer=["close"], max_rows=S.v(15, 13), head_style="rule2",
        rate_text=S.v("3.10% p.a.", "2.85% p.a."), open_range=(4000, 20000),
        hdr=dict(mast="logo", title="Savings statement", left=left, right=right, addr=(left, 96),
                 det=(S.r((320, 330), (300, 310)), 96, 84), det_rows=["acct", "period"],
                 sidebar=dict(style="shaded", x=S.r((320, 330), (300, 310)), y=136,
                              w=S.r((190, 200), (220, 230)),
                              rows=[("Interest rate", "rate"), ("Fees this period", "fees")])),
        table_top1=S.v(216, 210), table_top2=86,
        footer=dict(pos="left", fmt="Page {p} of {n}", other="Online saver"),
        stmts=pile(S, 3, (10, 24), S.v((2025, 8), (2026, 2)),
                   tweaks={0: dict(credits="sparse_none_p1", n=S.n(18, 24))}))


@design()
def bnz_business(S):
    font, size = S.v(("Helvetica", 8), ("Times", 8.5))
    left, right = S.r((34, 38), (42, 46)), S.r((558, 562), (550, 554))
    cols = mkcols(left, right, [
        ("date", S.r((48, 52), (42, 46)), ("Date",)),
        ("type", S.r((36, 40), (36, 40)), ("Type",)),
        ("desc", None, ("Details",)),
        ("debit", S.r((66, 72), (68, 74)), ("Debit",)),
        ("credit", S.r((66, 72), (68, 74)), ("Credit",)),
        ("balance", S.r((74, 80), (78, 84)), ("Balance",)),
    ])
    return layout(
        S, "bnz_business", bank="bnz", product="Business account",
        note="Busy business account over several pages: type column plus separate Debit / "
             "Credit, $ figures with thousands separators.",
        font=font, size=size, cols=cols, date_fmt=S.v("dd/mm/yyyy", "dd/mm/yy"),
        period_fmt="num", catalog="typed_biz", customers="biz", dollar=True, scale=S.v(1.0, 0.7),
        open_range=(8000, 60000), max_rows=S.v(22, 20), head_style="bar",
        head_fill=(0.88, 0.9, 0.95), opening=dict(label="Opening balance"),
        trailer=["totals", "close"], totals_label="Period totals",
        page_dx=S.v(None, ("jitter", 2.5)), head_dx=S.v(None, {"debit": -6.0, "credit": -6.0}),
        hdr=dict(mast="logo", title="Business statement", left=left,
                 right=right, addr=(left, 94), addr_bold=True,
                 det=(S.r((330, 336), (310, 316)), 94, 80),
                 det_rows=["acct", "period", "stno", "branch"],
                 summary=dict(style="band", x=left, y=158, w=right - left, size=8,
                              rows=[("Opening balance", "open"), ("Total debits", "tot_d"),
                                    ("Total credits", "tot_c"), ("Closing balance", "close")])),
        table_top1=226, table_top2=S.v(80, 86),
        footer=dict(pos="right", other="{branch}"),
        stmts=pile(S, 3, (40, 60), S.v((2026, 1), (2025, 9))))


@design()
def bnz_visa_fx(S):
    size = S.v(7.5, 8)
    left, right = S.r((36, 40), (44, 48)), S.r((556, 560), (548, 552))
    cols = mkcols(left, right, [
        ("date", S.r((40, 44), (44, 48)), ("Date",)),
        ("pdate", S.r((44, 48), (46, 50)), S.v(("Processed",), ("Posted",))),
        ("desc", None, ("Transaction details",)),
        ("fxcur", S.r((40, 44), (36, 40)), S.v(("Currency",), ("Ccy",))),
        ("fxamt", S.r((60, 64), (62, 66)), S.v(("Foreign", "amount"), ("Amount", "foreign"))),
        ("amount", S.r((66, 70), (70, 74)), S.v(("Amount", "NZD"), ("NZ$", "amount"))),
    ])
    return layout(
        S, "bnz_visa_fx", bank="bnz", product="Visa credit card",
        note="Credit card with separate Currency and Foreign-amount columns beside the NZ$ "
             "amount; payments with a leading minus.",
        font="Helvetica", size=size, cols=cols, date_fmt=S.v("dd/mm/yy", "dd Mon yy"),
        period_fmt="short", catalog="card_fx", credits="card", balance_kind="card",
        bal_mode="none", amount_style="card_minus", trailer=["close"],
        closing_label="New balance", head_style="rule", card=dict(limit=600000),
        rate_text="19.95% p.a.",
        hdr=dict(mast="plain", title="Visa statement", left=left, right=right, addr=(left, 98),
                 det=(S.r((318, 326), (300, 306)), 98, 76), det_rows=["card", "period", "issued"],
                 summary=dict(style="box", x=S.r((318, 326), (300, 306)), y=144,
                              w=S.r((200, 208), (226, 236)), title="Summary",
                              rows=[("Previous balance", "open"), ("Purchases", "tot_d"),
                                    ("Payments", "tot_c"), ("New balance", "close"),
                                    ("Payment due", "due")])),
        table_top1=S.v(250, 256), table_top2=84,
        footer=dict(pos="center"),
        stmts=pile(S, 3, (16, 36), S.v((2025, 9, 12), (2026, 4, 12))))


@design()
def bnz_bundle(S):
    size = S.v(8.5, 9)
    left, right = S.r((42, 46), (50, 54)), S.r((550, 554), (542, 546))
    cols = mkcols(left, right, [
        ("date", S.r((44, 48), (56, 60)), ("Date",)),
        ("desc", None, S.v(("Particulars",), ("Details",))),
        ("debit", S.r((66, 72), (68, 74)), S.v(("Money Out",), ("Withdrawals",))),
        ("credit", S.r((66, 72), (64, 70)), S.v(("Money In",), ("Deposits",))),
        ("balance", S.r((74, 80), (76, 82)), ("Balance",)),
    ])
    return layout(
        S, "bnz_bundle", bank="bnz", product="Everyday account",
        note="Two or three consecutive monthly statements of one account bundled in one PDF, "
             "each with its own header and page numbering; balances chain across them.",
        font="Helvetica", size=size, cols=cols, date_fmt=S.v("d Mon", "dd Mon"),
        period_fmt="long", catalog="bank", wrap=0.1, opening=dict(label="Opening balance"),
        trailer=["close"], head_style="rule", bundle=(2, 3), open_range=(6000, 15000),
        hdr=dict(mast="logo", title="Statement", left=left, right=right, addr=(left, 94),
                 det=(S.r((316, 322), (300, 306)), 94, 80), det_rows=["acct", "period", "stno"]),
        table_top1=S.v(180, 186), table_top2=86,
        footer=dict(pos="right"),
        stmts=pile(S, 3, (10, 18), S.v((2025, 3), (2025, 10))))


@design("holdout")
def bnz_value_date(S):
    size = 8
    left, right = 38, 558
    cols = mkcols(left, right, [
        ("date", 48, ("Date",)),
        ("pdate", 48, ("Value", "date")),
        ("desc", None, ("Details",)),
        ("debit", 66, ("Debit",)),
        ("credit", 66, ("Credit",)),
        ("balance", 78, ("Balance",)),
    ])
    return layout(
        S, "bnz_value_date", bank="bnz", product="Call account",
        note="Transaction date and value date side by side; truth date is the first.",
        font="Helvetica", size=size, cols=cols, date_fmt="dd-Mon-yy", period_fmt="short",
        catalog="bank", balance_style="drcr", opening=dict(label="Opening balance"),
        trailer=["close"], head_style="rule2",
        hdr=dict(mast="bar", color=MID, title="Account statement", left=left, right=right,
                 addr=(left, 98), det=(310, 98, 84), det_rows=["acct", "period", "stno"],
                 summary=dict(style="box", x=310, y=146, w=220, title="Summary",
                              rows=[("Opening balance", "open"), ("Debits", "tot_d"),
                                    ("Credits", "tot_c"), ("Closing balance", "close")])),
        table_top1=252, table_top2=84,
        footer=dict(pos="right"),
        stmts=pile(S, 3, (16, 30), (2026, 3)))


# ---- Westpac (03) --------------------------------------------------------------

@design()
def westpac_everyday(S):
    font, size = S.v(("Helvetica", 8.5), ("Helvetica", 9))
    left, right = S.r((40, 44), (50, 56)), S.r((552, 556), (544, 548))
    cols = mkcols(left, right, [
        ("date", S.r((52, 56), (48, 52)), ("Date",)),
        ("desc", None, S.v(("Transaction description",), ("Details",))),
        ("debit", S.r((62, 66), (58, 62)), ("Debit",)),
        ("credit", S.r((62, 66), (58, 62)), ("Credit",)),
        ("balance", S.r((82, 88), (84, 90)), ("Balance",)),
    ])
    return layout(
        S, "westpac_everyday", bank="westpac", product="Everyday account",
        note="Debit / Credit columns, balance carries a CR or DR token, account dips into "
             "overdraft, long descriptions run up to the debit column.",
        font=font, size=size, cols=cols, date_fmt=S.v("dd/mm/yyyy", "dd-mm-yy"),
        period_fmt=S.v("long", "num"), catalog="bank", long_desc=0.35,
        balance_style="drcr", hang=S.v((), ("balance",)), balance_kind="overdraft",
        trailer=["close"], head_style="rule",
        hdr=dict(mast="plain", title="Statement of account",
                 tagline="Westpac New Zealand Limited", left=left, right=right, addr=(left, 104),
                 det=(S.r((330, 338), (300, 306)), 100, 80), det_rows=["acct", "period", "issued"],
                 summary=dict(style="box", x=S.r((330, 338), (300, 306)), y=150,
                              w=S.r((200, 210), (226, 236)), title="Summary",
                              rows=[("Opening balance", "open"), ("Debits", "tot_d"),
                                    ("Credits", "tot_c"), ("Closing balance", "close")])),
        table_top1=258, table_top2=S.v(88, 80),
        footer=dict(pos="center"),
        stmts=pile(S, 4, (10, 46), S.v((2025, 11), (2026, 2))))


@design()
def westpac_overdraft(S):
    size = S.v(8.5, 8)
    left, right = S.r((44, 48), (34, 38)), S.r((550, 554), (558, 562))
    cols = mkcols(left, right, [
        ("date", S.r((40, 44), (34, 38)), ("Date",)),
        ("desc", None, ("Particulars",)),
        ("debit", S.r((64, 70), (62, 66)), ("Withdrawals",)),
        ("credit", S.r((64, 70), (62, 66)), ("Deposits",)),
        ("balance", S.r((78, 84), (82, 88)), ("Balance",)),
    ])
    return layout(
        S, "westpac_overdraft", bank="westpac", product="Overdraft account",
        note="Overdrawn balances carry an OD token after the figure.",
        font="Helvetica", size=size, cols=cols, date_fmt=S.v("dd Mon", "d Mon"),
        period_fmt="long", catalog="bank", wrap=0.2, balance_style="od",
        hang=S.v(("balance",), ()), balance_kind="overdraft", opening=dict(label="Opening balance"),
        trailer=["close"], head_style="rule",
        hdr=dict(mast="plain", title="Statement", left=left, right=right, addr=(left, 100),
                 det=(S.r((316, 324), (296, 304)), 100, 84), det_rows=["acct", "period", "stno"],
                 summary=dict(style="box", x=S.r((316, 324), (296, 304)), y=146,
                              w=S.r((204, 210), (232, 240)), title="Summary",
                              rows=[("Opening balance", "open"), ("Closing balance", "close"),
                                    ("Overdraft limit", "odlimit")])),
        table_top1=232, table_top2=86,
        footer=dict(pos="right", other="Westpac New Zealand Limited"),
        stmts=pile(S, 3, (16, 30), S.v((2025, 5), (2026, 6))))


@design()
def westpac_type(S):
    font, size = S.v(("Helvetica", 8), ("Courier", 7.5))
    left, right = S.r((36, 40), (46, 50)), S.r((556, 560), (548, 552))
    cols = mkcols(left, right, [
        ("date", S.r((50, 54), (50, 54)), ("Date",)),
        ("type", S.r((44, 48), (40, 44)), ("Type",)),
        ("desc", None, ("Details",)),
        ("amount", S.r((70, 76), (84, 90)), S.v(("Amount",), ("Amount NZD",))),
        ("balance", S.r((74, 80), (84, 90)), ("Balance",)),
    ])
    return layout(
        S, "westpac_type", bank="westpac", product="Everyday plus",
        note="Transaction-type column (EFTPOS, AP, DD, BP, DC ...) before the details, one "
             "signed Amount column.",
        font=font, size=size, cols=cols, date_fmt=S.v("dd/mm/yyyy", "d/mm/yyyy"),
        period_fmt="num", catalog="typed", amount_style=S.v("lead", "paren"),
        balance_style="lead", trailer=["close"], head_style=S.v("rule", "box"),
        row_rules=S.v(True, False), head_shift=S.v(0.0, 3.0),
        hdr=dict(mast="bar", color=DARK, title="Statement", left=left, right=right,
                 addr=(left, 98), det=(S.r((330, 340), (300, 310)), 98, 80),
                 det_rows=["acct", "period", "stno"],
                 sidebar=dict(style="box", x=S.r((330, 340), (300, 310)), y=146,
                              w=S.r((200, 210), (230, 240)), title="Important information",
                              rows=[("Interest rate", "rate"), ("Fees this period", "fees")],
                              notes=["Fees are charged on the last business day."])),
        rate_text=S.v("0.05% p.a.", "0.50% p.a."),
        table_top1=S.v(236, 244), table_top2=80,
        footer=dict(pos="left", other="Ref {stno}"),
        stmts=pile(S, 3, (8, 36), S.v((2025, 3), (2025, 12))))


@design()
def westpac_staggered(S):
    size = S.v(8, 8.5)
    left, right = S.r((40, 44), (48, 52)), S.r((552, 556), (544, 548))
    cols = mkcols(left, right, [
        ("date", S.r((44, 48), (52, 56)), S.v(("Date",), ("Txn Date",))),
        ("desc", None, ("Transaction Description",)),
        ("debit", S.r((64, 70), (66, 72)), S.v(("Debits",), ("Dr",))),
        ("credit", S.r((64, 70), (62, 66)), S.v(("Credits",), ("Cr",))),
        ("balance", S.r((76, 82), (78, 84)), S.v(("Running Balance",), ("Bal",))),
    ])
    return layout(
        S, "westpac_staggered", bank="westpac", product="Cheque account",
        note="Staggered rows: on a two-line transaction the date is on the first line and "
             "the amount and balance on the second.",
        font="Helvetica", size=size, cols=cols, date_fmt=S.v("dd/mm/yy", "dd-mm-yyyy"),
        period_fmt="long", catalog="bank", wrap=0.6, money_line="last",
        opening=dict(label="Opening balance"), trailer=["close"], head_style="rule2",
        hdr=dict(mast="plain", title="Account statement", left=left, right=right,
                 addr=(left, 100), det=(S.r((318, 326), (300, 306)), 100, 82),
                 det_rows=["acct", "period", "stno"]),
        table_top1=S.v(184, 190), table_top2=84,
        footer=dict(pos="right"),
        stmts=pile(S, 3, (14, 34), S.v((2025, 12), (2026, 5))))


@design("dev")
def westpac_trail(S):
    size = 9
    left, right = 46, 550
    cols = mkcols(left, right, [
        ("date", 52, ("Date",)),
        ("desc", None, ("Description",)),
        ("amount", 78, ("Amount",)),
        ("balance", 80, ("Balance",)),
    ])
    return layout(
        S, "westpac_trail", bank="westpac", product="Flexi account",
        note="Signed amounts with a trailing minus in a proportional font; balance "
             "brought / carried forward lines.",
        font="Helvetica", size=size, cols=cols, date_fmt="dd MON yy", period_fmt="short",
        catalog="bank", long_desc=0.3, amount_style="trail", balance_style="trail",
        cf_bf=("Carried forward", "Brought forward"), opening=dict(label="Brought forward"),
        trailer=["close"], max_rows=22, head_style="bar",
        hdr=dict(mast="bar", color=DARK, title="Statement", left=left, right=right,
                 addr=(left, 96), det=(320, 96, 80), det_rows=["acct", "period"],
                 summary=dict(style="box", x=320, y=132, w=200, title="Summary",
                              rows=[("Opening balance", "open"), ("Closing balance", "close")])),
        table_top1=214, table_top2=80,
        footer=dict(pos="right"),
        stmts=pile(S, 3, (16, 44), (2026, 2)))


@design("holdout")
def westpac_business_landscape(S):
    size = 8.5
    left, right = 48, 720
    cols = mkcols(left, right, [
        ("date", 64, ("DATE",)),
        ("type", 56, ("TYPE",)),
        ("desc", None, ("TRANSACTION DETAILS",)),
        ("amount", 96, ("AMOUNT",)),
        ("balance", 104, ("BALANCE",)),
    ])
    return layout(
        S, "westpac_business_landscape", bank="westpac", product="Business online",
        note="Landscape page, upper-case headings, type column, signed amount with trailing "
             "minus, balance with OD.",
        page="A4L", font="Helvetica", size=size, cols=cols, date_fmt="dd MON yy",
        period_fmt="upper", catalog="typed_biz", customers="biz", scale=0.5,
        amount_style="trail", balance_style="od", balance_kind="overdraft",
        opening=dict(label="OPENING BALANCE"), trailer=["close"], closing_label="CLOSING BALANCE",
        head_style="bar", max_rows=20,
        hdr=dict(mast="bar", color=DARK, title="Business statement", left=left, right=794,
                 addr=(left, 92), det=(420, 92, 90), det_rows=["name", "acct", "period"]),
        table_top1=168, table_top2=80,
        footer=dict(pos="right", other="Westpac New Zealand Limited"),
        stmts=pile(S, 3, (30, 50), (2026, 2)))


# ---- Kiwibank (38) -------------------------------------------------------------

@design()
def kiwibank_everyday(S):
    font, size = S.v(("Helvetica", 9), ("Times", 9))
    left, right = S.r((50, 56), (40, 44)), S.r((545, 550), (553, 557))
    cols = mkcols(left, right, [
        ("date", S.r((62, 66), (60, 64)), ("Date",)),
        ("desc", None, ("Description",)),
        ("debit", S.r((70, 76), (74, 80)), S.v(("Money out",), ("Paid out",))),
        ("credit", S.r((70, 76), (72, 78)), S.v(("Money in",), ("Paid in",))),
        ("balance", S.r((76, 82), (80, 86)), ("Balance",)),
    ])
    return layout(
        S, "kiwibank_everyday", bank="kiwibank", product="Everyday account",
        note="Weekday dates with no year, $ figures, shaded alternate rows, an interest and "
             "fees sidebar.",
        font=font, size=size, cols=cols, date_fmt=S.v("Dow dd Mon", "Dow d Mon"),
        period_fmt="long", catalog="bank", wrap=0.15, dollar=True, zebra=True,
        head_every_page=S.v(False, True), head_style="bar", head_fill=(0.9, 0.9, 0.9),
        trailer=["close"], opening=S.v(None, dict(label="Opening balance")),
        rate_text=S.v("0.10% p.a.", "0.25% p.a."),
        hdr=dict(mast="logo", title="Your statement", left=left, right=right,
                 addr=(left, 96), det=(left, 160, 90), det_rows=["acct", "period"],
                 sidebar=dict(style="box", x=S.r((350, 360), (330, 340)), y=88,
                              w=S.r((180, 190), (210, 216)), title="Your account at a glance",
                              rows=[("Opening balance", "open"), ("Money out", "tot_d"),
                                    ("Money in", "tot_c"), ("Closing balance", "close"),
                                    ("Interest rate", "rate"), ("Fees this period", "fees")])),
        table_top1=S.v(232, 226), table_top2=S.v(70, 84),
        footer=dict(pos="right"),
        stmts=pile(S, 4, (14, 46), S.v((2025, 12, 20), (2026, 1, 12))))


@design()
def kiwibank_duplex(S):
    size = S.v(8, 8.5)
    left, right = S.r((40, 44), (44, 48)), S.r((546, 550), (540, 544))
    cols = mkcols(left, right, [
        ("date", S.r((38, 42), (40, 44)), ("Date",)),
        ("desc", None, ("Details",)),
        ("debit", S.r((66, 70), (68, 72)), ("Withdrawals",)),
        ("credit", S.r((62, 66), (64, 68)), ("Deposits",)),
        ("balance", S.r((72, 78), (74, 80)), ("Balance",)),
    ])
    return layout(
        S, "kiwibank_duplex", bank="kiwibank", product="Cheque account",
        note="Printed for duplex binding: even pages shifted right by a gutter, so the table "
             "sits at a different x on alternate pages; balance on the last row of each day.",
        font="Helvetica", size=size, cols=cols, date_fmt="dd Mon",
        period_fmt="long", catalog="bank", wrap=0.25, bal_mode="last_of_day",
        page_dx=("gutter", S.v(7.0, 5.0)), max_rows=S.v(16, 18),
        opening=dict(label="Opening balance"), trailer=["close"], head_style="rule",
        hdr=dict(mast="logo", title="Statement", left=left, right=right,
                 addr=(left, 96), det=(S.r((316, 324), (300, 306)), 96, 80),
                 det_rows=["acct", "period"],
                 summary=dict(style="box", x=S.r((316, 324), (300, 306)), y=126,
                              w=S.r((196, 204), (210, 218)), title="Summary",
                              rows=[("Opening balance", "open"), ("Withdrawals", "tot_d"),
                                    ("Deposits", "tot_c"), ("Closing balance", "close")])),
        table_top1=226, table_top2=80,
        footer=dict(pos=S.v("right", "center"), small="about"),
        stmts=pile(S, 3, (36, 60), S.v((2025, 9), (2026, 10))))


@design()
def kiwibank_card(S):
    font, size = S.v(("Helvetica", 7), ("Times", 8))
    left, right = S.r((36, 40), (46, 50)), S.r((556, 560), (548, 552))
    cols = mkcols(left, right, [
        ("date", S.r((48, 52), (46, 50)), ("Date",)),
        ("pdate", S.r((50, 54), (48, 52)), ("Processed",)),
        ("desc", None, ("Details",)),
        ("amount", S.r((60, 66), (66, 72)), ("Amount",)),
    ])
    return layout(
        S, "kiwibank_card", bank="kiwibank", product="Mastercard",
        note="Credit card where payments and refunds print with a leading minus; "
             "processed date column; heading on every page.",
        font=font, size=size, cols=cols, date_fmt=S.v("dd/mm/yyyy", "dd-mm-yy"),
        period_fmt="num", catalog="card", credits="card", balance_kind="card",
        bal_mode="none", amount_style="card_minus", wrap=0.25, trailer=["close"],
        closing_label="New balance", head_style="rule", card=dict(limit=S.ch([300000], [450000])),
        rate_text="13.95% p.a.",
        hdr=dict(mast="right", title="Credit card statement", left=left, right=right,
                 addr=(left, 96), det=(left, 160, 76), det_rows=["card", "period"],
                 summary=dict(style="box", x=S.r((330, 340), (310, 320)), y=90,
                              w=S.r((200, 210), (226, 236)), title="Statement summary",
                              rows=[("Previous balance", "open"), ("Purchases", "tot_d"),
                                    ("Payments and credits", "tot_c"), ("New balance", "close"),
                                    ("Minimum payment", "minpay"), ("Payment due", "due")])),
        table_top1=S.v(230, 236), table_top2=S.v(80, 76), max_rows=S.v(None, 22),
        footer=dict(pos="center"),
        stmts=pile(S, 3, (20, 56), S.v((2025, 5, 3), (2026, 7, 21))))


@design()
def kiwibank_balance_left(S):
    size = S.v(8.5, 9)
    left, right = S.r((44, 48), (38, 42)), S.r((550, 554), (556, 560))
    cols = mkcols(left, right, [
        ("date", S.r((52, 56), (64, 68)), S.v(("Date",), ("Txn Date",))),
        ("desc", None, S.v(("Transaction",), ("Narrative",))),
        ("balance", S.r((76, 82), (80, 86)), S.v(("Balance",), ("Bal",))),
        ("amount", S.r((72, 78), (76, 82)), S.v(("Amount",), ("Amount NZD",))),
    ])
    return layout(
        S, "kiwibank_balance_left", bank="kiwibank", product="Online call account",
        note="The running balance column printed LEFT of the signed amount column.",
        font="Helvetica", size=size, cols=cols, date_fmt=S.v("dd-mm-yyyy", "d Mon yyyy"),
        period_fmt="short", catalog="bank", amount_style="lead", dollar=S.v(False, True),
        opening=dict(label="Opening balance"), trailer=["close"], head_style="rule",
        hdr=dict(mast="logo", title="Account statement", left=left, right=right,
                 addr=(left, 96), det=(S.r((316, 324), (300, 306)), 96, 82),
                 det_rows=["acct", "period", "stno"]),
        table_top1=S.v(178, 184), table_top2=82,
        footer=dict(pos="right"),
        stmts=pile(S, 3, (14, 40), S.v((2026, 3), (2025, 6))))


@design("dev")
def kiwibank_youth(S):
    size = 10
    left, right = 50, 546
    cols = mkcols(left, right, [
        ("date", 82, ("Date",)),
        ("desc", None, ("What for",)),
        ("debit", 70, ("Money out",)),
        ("credit", 70, ("Money in",)),
        ("balance", 76, ("Balance",)),
    ])
    return layout(
        S, "kiwibank_youth", bank="kiwibank", product="Youth account",
        note="Large type, few rows, $ figures, a single page with a sidebar of tips.",
        font="Helvetica", size=size, cols=cols, date_fmt="Dow dd Mon", period_fmt="long",
        catalog="bank", dollar=True, thousands=True, credit_p=0.3, zebra=True,
        trailer=["close"], head_style="bar", head_fill=(0.9, 0.9, 0.9),
        rate_text="1.00% p.a.", open_range=(50, 600),
        hdr=dict(mast="logo", title="Your statement", left=left,
                 right=right, addr=(left, 96), det=(left, 160, 96), det_rows=["acct", "period"],
                 sidebar=dict(style="shaded", x=330, y=90, w=190, title="Good to know",
                              rows=[("Interest rate", "rate"), ("Fees this period", "fees")])),
        table_top1=226, table_top2=84,
        footer=dict(pos="center"),
        stmts=pile(S, 3, (8, 14), (2026, 4)))


@design("holdout")
def kiwibank_app(S):
    size = 9
    left, right = 50, 546
    cols = mkcols(left, right, [
        ("date", 76, ("Date",)),
        ("desc", None, ("Description",)),
        ("amount", 84, ("Amount",)),
        ("balance", 84, ("Balance",)),
    ])
    return layout(
        S, "kiwibank_app", bank="kiwibank", product="Spend account",
        note="App-style print: explicit + and - on every amount, $ figures, no rules.",
        font="Helvetica", size=size, pitch=17, cols=cols, date_fmt="dd Mon yyyy",
        period_fmt="long", catalog="bank", dollar=True, amount_style="plus",
        balance_style="lead", wrap=0.2, trailer=["close"], head_style="none",
        hdr=dict(mast="logo", title="Monthly statement", left=left,
                 right=right, addr=(left, 96), det=(left, 156, 96), det_rows=["acct", "period"],
                 summary=dict(style="shaded", x=330, y=90, w=190, title="This month",
                              rows=[("Started with", "open"), ("Spent", "tot_d"),
                                    ("Received", "tot_c"), ("Ended with", "close")])),
        table_top1=232, table_top2=72,
        footer=dict(pos="center", fmt="{p}/{n}"),
        stmts=pile(S, 3, (14, 34), (2026, 8)))


# ---- TSB (15) ------------------------------------------------------------------

@design()
def tsb_everyday(S):
    font, size = S.v(("Times", 9), ("Helvetica", 8))
    left, right = S.r((54, 60), (40, 44)), S.r((540, 546), (552, 556))
    cols = mkcols(left, right, [
        ("date", S.r((48, 52), (56, 60)), ("Date",)),
        ("desc", None, ("Details",)),
        ("amount", S.r((70, 76), (80, 86)), ("Amount",)),
        ("balance", S.r((76, 82), (80, 86)), ("Balance",)),
    ])
    return layout(
        S, "tsb_everyday", bank="tsb", product="Everyday account",
        note="One signed Amount column (leading minus on money out) beside a running balance; "
             "opening balance as the first line, closing balance as the last.",
        font=font, size=size, cols=cols, date_fmt=S.v("dd-mm-yy", "dd-mm-yyyy"),
        period_fmt=S.v("short", "long"), catalog="bank", wrap=S.v(0.0, 0.2), wrap_indent=S.v(0, 6),
        amount_style="lead", opening=dict(label="Opening balance"), trailer=["close"],
        head_style="rule2",
        hdr=dict(mast="right", title="Account statement", left=left, right=right,
                 addr=(left, 92), det=(left, 160, 96), det_rows=["acct", "period", "stno"],
                 summary=dict(style="box", x=S.r((320, 330), (340, 350)), y=104,
                              w=S.r((190, 200), (190, 196)), title="Account summary",
                              rows=[("Opening balance", "open"), ("Total paid out", "tot_d"),
                                    ("Total paid in", "tot_c"), ("Closing balance", "close")])),
        table_top1=S.v(232, 236), table_top2=80,
        footer=dict(pos="right", other="Statement {stno}"),
        stmts=pile(S, 4, (16, 52), S.v((2025, 7), (2026, 9))))


@design()
def tsb_cheque(S):
    font, size = S.v(("Times", 9.5), ("Helvetica", 9))
    left, right = S.r((48, 52), (40, 44)), S.r((546, 550), (550, 554))
    cols = mkcols(left, right, [
        ("date", S.r((48, 52), (60, 64)), ("Date",)),
        ("desc", None, ("Particulars",)),
        ("debit", S.r((66, 72), (70, 76)), S.v(("Debits",), ("Withdrawals",))),
        ("credit", S.r((66, 72), (66, 72)), S.v(("Credits",), ("Lodgements",))),
        ("balance", S.r((74, 80), (76, 82)), ("Balance",)),
    ])
    return layout(
        S, "tsb_cheque", bank="tsb", product="Cheque account",
        note="Each page's table lands a few points left or right of the last; continuation "
             "pages are short.",
        font=font, size=size, cols=cols, date_fmt=S.v("dd-mm-yy", "dd-Mon-yy"),
        period_fmt="short", catalog="bank", wrap=0.1, page_dx=("jitter", 4.0),
        max_rows=S.v(18, 15), opening=dict(label="Opening balance"), trailer=["close"],
        head_style="rule",
        hdr=dict(mast="right", title="Statement", left=left, right=right, addr=(left, 92),
                 det=(left, 150, 90), det_rows=["acct", "period"], cont="line"),
        table_top1=S.v(210, 204), table_top2=86,
        footer=dict(pos="center"),
        stmts=pile(S, 3, (20, 44), S.v((2025, 4), (2026, 5)), tweaks={0: dict(short_last=True)}))


@design("dev")
def tsb_saver(S):
    size = 8.5
    left, right = 40, 552
    cols = mkcols(left, right, [
        ("date", 50, ("Date",)),
        ("desc", None, ("Transaction",)),
        ("amount", 80, ("Amount",)),
        ("balance", 84, ("Balance",)),
    ])
    return layout(
        S, "tsb_saver", bank="tsb", product="Smart saver",
        note="Parenthesised money out with a $ prefix, account summary box.",
        font="Helvetica", size=size, cols=cols, date_fmt="dd-mm-yy", period_fmt="short",
        catalog="bank", dollar=True, amount_style="paren", balance_style="paren",
        wrap=0.25, wrap_indent=10, opening=dict(label="Opening balance"), trailer=["close"],
        head_style="rule",
        hdr=dict(mast="right", title="Statement", left=left, right=right, addr=(left, 96),
                 det=(left, 156, 90), det_rows=["acct", "period", "stno"],
                 summary=dict(style="box", x=330, y=96, w=200, title="Account summary",
                              rows=[("Opening balance", "open"), ("Money out", "tot_d"),
                                    ("Money in", "tot_c"), ("Closing balance", "close")])),
        table_top1=226, table_top2=80,
        footer=dict(pos="left"),
        stmts=pile(S, 3, (24, 52), (2026, 6)))


@design()
def tsb_notice_saver(S):
    size = S.v(9, 10)
    left, right = S.r((54, 58), (44, 48)), S.r((540, 544), (548, 552))
    cols = mkcols(left, right, [
        ("date", S.r((96, 100), (66, 70)), ("Date",)),
        ("desc", None, ("Transaction",)),
        ("debit", S.r((68, 74), (72, 78)), S.v(("Payments",), ("Withdrawals",))),
        ("credit", S.r((68, 74), (66, 72)), S.v(("Receipts",), ("Deposits",))),
        ("balance", S.r((76, 82), (80, 86)), ("Balance",)),
    ])
    return layout(
        S, "tsb_notice_saver", bank="tsb", product="Notice saver",
        note="Quiet notice-saver account: few transactions over a quarter, a two-page "
             "statement whose second page holds one or two rows.",
        font="Times", size=size, cols=cols, date_fmt=S.v("d Month yyyy", "d Mon yyyy"),
        period_fmt="long", catalog="savings", credits="sparse", max_rows=S.v(9, 7),
        opening=dict(label="Opening balance", date=True), trailer=["close"], head_style="rule2",
        rate_text=S.v("4.15% p.a.", "3.95% p.a."), open_range=(6000, 40000),
        hdr=dict(mast="plain", title="Statement", tagline="Notice saver",
                 left=left, right=right, addr=(left, 104), addr_size=10,
                 det=(S.r((316, 322), (300, 306)), 104, 84), det_size=9,
                 det_rows=["acct", "period"],
                 summary=dict(style="box", x=S.r((316, 322), (300, 306)), y=140,
                              w=S.r((200, 206), (236, 244)), size=9, title="Your savings",
                              rows=[("Opening balance", "open"), ("Closing balance", "close"),
                                    ("Interest rate", "rate")])),
        table_top1=244, table_top2=90,
        footer=dict(pos="center"),
        stmts=pile(S, 3, (8, 12), S.v((2025, 1), (2025, 4)), months=3,
                   tweaks={0: dict(short_last=True), 2: dict(short_last=True)}))


@design()
def tsb_desc_first(S):
    size = S.v(8.5, 8)
    left, right = S.r((44, 48), (36, 40)), S.r((550, 554), (558, 562))
    cols = mkcols(left, right, [
        ("desc", None, S.v(("Transaction",), ("Transaction Description",))),
        ("date", S.r((54, 58), (60, 64)), S.v(("Date",), ("Txn Date",))),
        ("debit", S.r((66, 72), (64, 70)), S.v(("Paid Out",), ("Withdrawals",))),
        ("credit", S.r((66, 72), (64, 70)), S.v(("Paid In",), ("Deposits",))),
        ("balance", S.r((74, 80), (76, 82)), S.v(("Balance",), ("Running Balance",))),
    ])
    return layout(
        S, "tsb_desc_first", bank="tsb", product="Everyday account",
        note="The description column comes BEFORE the date column.",
        font="Helvetica", size=size, cols=cols, date_fmt=S.v("dd/mm/yyyy", "d Mon yyyy"),
        period_fmt="short", catalog="bank", wrap=0.15, opening=dict(label="Opening balance"),
        trailer=["close"], head_style="bar", head_fill=(0.9, 0.9, 0.9),
        hdr=dict(mast="plain", title="Statement", left=left, right=right, addr=(left, 98),
                 det=(S.r((316, 324), (300, 306)), 98, 84), det_rows=["acct", "period", "stno"],
                 summary=dict(style="box", x=S.r((316, 324), (300, 306)), y=144,
                              w=S.r((200, 208), (226, 236)), title="Summary",
                              rows=[("Opening balance", "open"), ("Closing balance", "close")])),
        table_top1=226, table_top2=84,
        footer=dict(pos="right"),
        stmts=pile(S, 3, (16, 40), S.v((2026, 2), (2025, 8))))


# ---- Co-operative Bank (02 -- shared with BNZ) ---------------------------------------

@design()
def coop_cheque(S):
    size = S.v(8.5, 9.5)
    left, right = S.r((40, 44), (52, 58)), S.r((552, 556), (540, 546))
    cols = mkcols(left, right, [
        ("date", S.r((62, 66), (88, 92)), ("Date",)),
        ("desc", None, ("Description",)),
        ("debit", S.r((64, 70), (68, 74)), S.v(("Paid out", "($)"), ("Payments", "$"))),
        ("credit", S.r((64, 70), (68, 74)), S.v(("Paid in", "($)"), ("Receipts", "$"))),
        ("balance", S.r((72, 78), (76, 82)), S.v(("Balance", "($)"), ("Balance", "$"))),
    ])
    return layout(
        S, "coop_cheque", bank="coop", product="Cheque account",
        note="Two-line money headings, headings printed on page 1 only, multi-line "
             "descriptions with an indented second line.",
        font="Helvetica", size=size, cols=cols, date_fmt=S.v("d Mon yyyy", "d Month yyyy"),
        period_fmt="short", catalog="bank", wrap=0.35, wrap_indent=S.v(8, 0),
        head_every_page=False, head_style="rule", trailer=["totals", "close"],
        totals_label="Totals", opening=dict(label="Opening balance", date=True),
        hdr=dict(mast="plain", title="Statement of account", left=left, right=right,
                 addr=(left, 100), det=(S.r((310, 320), (300, 306)), 100, 84),
                 det_rows=["acct", "period", "stno", "branch"],
                 cont="minimal", top_page=True),
        table_top1=S.v(190, 196), table_top2=S.v(70, 76),
        footer=dict(pos="right", small="first"),
        stmts=pile(S, 3, (28, 52), S.v((2025, 9), (2026, 2)), tweaks={0: dict(short_last=True)}))


@design()
def coop_everyday(S):
    size = S.v(8.5, 8)
    left, right = S.r((44, 48), (36, 40)), S.r((548, 552), (556, 560))
    cols = mkcols(left, right, [
        ("date", S.r((62, 66), (58, 62)), ("Date",)),
        ("desc", None, ("Transaction details",)),
        ("debit", S.r((66, 72), (70, 76)), ("Withdrawals",)),
        ("credit", S.r((66, 72), (62, 68)), ("Deposits",)),
        ("balance", S.r((72, 78), (76, 82)), ("Balance",)),
    ])
    return layout(
        S, "coop_everyday", bank="coop", product="Everyday account",
        note="Fees-and-interest summary block printed under the table (figures that are "
             "not transactions), weekday dates with no year.",
        font="Helvetica", size=size, cols=cols, date_fmt="Dow dd Mon",
        period_fmt="long", catalog="bank", wrap=0.2, opening=dict(label="Opening balance"),
        trailer=["close"], head_style="rule", row_rules=S.v(False, True),
        after=dict(title="Fees and interest this period", x=S.r((60, 70), (300, 310)),
                   w=S.r((200, 210), (190, 200)),
                   rows=[("Account fees", "fees"), ("Credit interest", "interest"),
                         ("Transactions", "count")]),
        hdr=dict(mast="plain", title="Account statement", left=left, right=right,
                 addr=(left, 98), det=(S.r((320, 330), (300, 310)), 98, 84),
                 det_rows=["acct", "period", "stno"],
                 summary=dict(style="lines", x=S.r((320, 330), (300, 310)), y=134,
                              w=S.r((190, 200), (232, 240)),
                              rows=[("Opening balance", "open"), ("Closing balance", "close")])),
        table_top1=214, table_top2=82,
        footer=dict(pos="left", other="The Co-operative Bank Limited", small="first"),
        stmts=pile(S, 4, (24, 50), S.v((2025, 11), (2025, 12))))


@design("dev")
def coop_reference(S):
    size = 9
    left, right = 42, 554
    cols = mkcols(left, right, [
        ("date", 92, ("Date",)),
        ("desc", None, ("Details",)),
        ("ref", 74, ("Reference",)),
        ("debit", 64, ("Withdrawals",)),
        ("credit", 60, ("Deposits",)),
        ("balance", 74, ("Balance",)),
    ])
    return layout(
        S, "coop_reference", bank="coop", product="Savings account",
        note="A separate Reference text column between the details and the money.",
        font="Times", size=size, cols=cols, date_fmt="d Month yyyy", period_fmt="long",
        catalog="bank", opening=dict(label="Opening balance"), trailer=["close"],
        head_style="rule",
        hdr=dict(mast="plain", title="Statement", left=left, right=right, addr=(left, 100),
                 det=(318, 100, 84), det_rows=["acct", "period", "stno"]),
        table_top1=186, table_top2=84,
        footer=dict(pos="center"),
        stmts=pile(S, 3, (14, 30), (2026, 2)))


@design()
def coop_business(S):
    size = S.v(8, 7.5)
    left, right = S.r((38, 42), (46, 50)), S.r((556, 560), (548, 552))
    cols = mkcols(left, right, [
        ("date", S.r((46, 50), (50, 54)), S.v(("Date",), ("Txn Date",))),
        ("desc", None, S.v(("Details",), ("Narrative",))),
        ("debit", S.r((76, 82), (78, 84)), S.v(("Dr",), ("Debits",))),
        ("credit", S.r((76, 82), (78, 84)), S.v(("Cr",), ("Credits",))),
        ("balance", S.r((80, 86), (84, 90)), S.v(("Balance",), ("Bal",))),
    ])
    return layout(
        S, "coop_business", bank="coop", product="Business account",
        note="Business account with a property settlement in seven figures (1,234,567.89), "
             "fee-waived 0.00 rows, and page sub-totals.",
        font="Helvetica", size=size, cols=cols, date_fmt=S.v("dd/mm/yyyy", "dd-mm-yyyy"),
        period_fmt="num", catalog="biz", customers="biz", scale=0.6, large=1, zero_rows=1,
        open_range=(20000, 80000), page_totals="Page totals", trailer=["totals_close"],
        totals_label="Totals for period", opening=dict(label="Opening balance"),
        head_style="rule", max_rows=S.v(24, 22),
        hdr=dict(mast="plain", title="Business statement", left=left, right=right,
                 addr=(left, 98), addr_bold=True, det=(S.r((316, 324), (300, 306)), 98, 84),
                 det_rows=["name", "acct", "period"]),
        table_top1=184, table_top2=84,
        footer=dict(pos="right", other="The Co-operative Bank Limited", small="none"),
        stmts=pile(S, 3, (30, 56), S.v((2025, 7), (2026, 6))))


@design()
def coop_online_printout(S):
    size = S.v(8, 7.5)
    left, right = S.r((34, 38), (42, 46)), S.r((558, 562), (550, 554))
    cols = mkcols(left, right, [
        ("date", S.r((48, 52), (44, 48)), S.v(("Txn Date",), ("Date",))),
        ("payee", None, S.v(("Payee",), ("Other Party",))),
        ("part", S.r((62, 66), (60, 64)), ("Particulars",)),
        ("code", S.r((54, 58), (52, 56)), ("Code",)),
        ("ref", S.r((60, 64), (58, 62)), ("Reference",)),
        ("amount", S.r((66, 70), (68, 72)), S.v(("Amount NZD",), ("Amt (NZD)",))),
        ("balance", S.r((70, 74), (72, 76)), S.v(("Running Balance",), ("Balance",))),
    ])
    return layout(
        S, "coop_online_printout", bank="coop", product="Internet banking printout",
        note="Internet-banking printout with no masthead: the bank is named only in the "
             "footer, and its 02 code is shared with BNZ, so both cues are needed.",
        font="Helvetica", size=size, cols=cols, date_fmt=S.v("dd/mm/yy", "d/mm/yyyy"),
        period_fmt="num", catalog="payee_pcr", amount_style="lead",
        opening=dict(label="Opening balance"), trailer=["close"], head_style="bar",
        head_fill=(0.92, 0.92, 0.92),
        hdr=dict(mast="ib", title="Transaction search results", left=left, right=right,
                 addr=None, det=(left, 70, 90), det_rows=["name", "acct", "period", "printed"],
                 cont="minimal"),
        table_top1=146, table_top2=70,
        footer=dict(pos="right", other="The Co-operative Bank - internet banking", small="none"),
        stmts=pile(S, 3, (14, 44), S.v((2026, 1), (2025, 9))))


# ---- Rimu Bank (99, fictional) -------------------------------------------------

@design()
def rimu_mainframe(S):
    size = S.v(8, 7.5)
    left, right = S.r((36, 40), (44, 48)), S.r((556, 560), (548, 552))
    cols = mkcols(left, right, [
        ("date", S.r((50, 54), (40, 44)), ("DATE",)),
        ("desc", None, ("DESCRIPTION",)),
        ("amount", S.r((84, 90), (90, 96)), ("AMOUNT",)),
        ("balance", S.r((90, 96), (96, 102)), ("BALANCE",)),
    ])
    return layout(
        S, "rimu_mainframe", bank="rimu", product="Savings plus",
        note="Monospaced mainframe print: trailing minus on amounts and balances, dashed "
             "heading underline, BALANCE B/F and C/F lines.",
        font="Courier", size=size, pitch=round(size * S.v(1.5, 1.45), 2), cols=cols,
        date_fmt=S.v("dd MON yy", "dd MON"), period_fmt="upper", period_sep=" TO ",
        period_label="PERIOD", catalog="bank", thousands=S.v(False, True),
        amount_style="trail", balance_style="trail", balance_kind="overdraft",
        head_style="dashes", cf_bf=("BALANCE C/F", "BALANCE B/F"),
        opening=dict(label="BALANCE B/F"), trailer=["close"], closing_label="CLOSING BALANCE",
        hdr=dict(mast="courier", title="Statement of account", left=left, right=right,
                 addr=(left, 90), addr_size=8, det=(S.r((316, 324), (300, 306)), 90, 82),
                 det_size=8, det_rows=["acct", "period", "stno"], cont="courier"),
        table_top1=S.v(170, 176), table_top2=80, max_rows=S.v(36, 40),
        footer=dict(pos="center", fmt="PAGE {p} OF {n}", size=7, small="about"),
        stmts=pile(S, 3, (16, 58), S.v((2026, 1), (2026, 4)), tweaks={0: dict(short_last=True)}))


@design()
def rimu_credit_union(S):
    size = S.v(8, 7)
    left, right = S.r((40, 44), (32, 36)), S.r((552, 556), (560, 564))
    cols = mkcols(left, right, [
        ("date", S.r((52, 56), (48, 52)), ("DATE",)),
        ("desc", None, ("DETAILS",)),
        ("debit", S.r((80, 86), (76, 82)), ("WITHDRAWALS",)),
        ("credit", S.r((70, 76), (66, 72)), ("DEPOSITS",)),
        ("balance", S.r((76, 82), (72, 78)), ("BALANCE",)),
    ])
    return layout(
        S, "rimu_credit_union", bank="rimu", product="Everyday access",
        note="Monospaced print with vertical column rules, no thousands separators, headings "
             "on page 1 only.",
        font="Courier", size=size, pitch=round(size * 1.6, 2), cols=cols,
        date_fmt=S.v("dd-mm-yy", "dd/mm/yy"), period_fmt="upper", period_sep=" TO ",
        period_label="PERIOD", catalog="bank", thousands=False, wrap=0.15,
        head_every_page=False, head_style="dashes", col_rules=True,
        opening=dict(label="OPENING BALANCE"), trailer=["close"], closing_label="CLOSING BALANCE",
        hdr=dict(mast="courier", title="Member statement", left=left, right=right,
                 addr=(left, 88), addr_size=8, det=(S.r((320, 330), (334, 340)), 88, 72),
                 det_size=8, det_rows=["acct", "period", "branch"], cont="courier"),
        table_top1=162, table_top2=78,
        footer=dict(pos="right", fmt="PAGE {p}", small="about"),
        stmts=pile(S, 3, (12, 40), S.v((2025, 10), (2026, 2))))


@design()
def rimu_business(S):
    size = S.v(7, 7.5)
    left, right = S.r((44, 48), (34, 38)), S.r((548, 552), (558, 562))
    cols = mkcols(left, right, [
        ("date", S.r((50, 54), (46, 50)), ("Date",)),
        ("desc", None, ("Details",)),
        ("amount", S.r((70, 76), (76, 82)), ("Amount",)),
        ("balance", S.r((74, 80), (80, 86)), ("Balance",)),
    ])
    return layout(
        S, "rimu_business", bank="rimu", product="Business edge",
        note="Signed Amount with a separate DR / CR token after the figure, and the balance "
             "carrying its own CR / DR token.",
        font="Helvetica", size=size, cols=cols, date_fmt=S.v("d Mon yyyy", "dd Mon yy"),
        period_fmt="short", catalog="biz", customers="biz", scale=S.v(1.0, 0.6),
        amount_style="drcr", balance_style="drcr", hang=("amount", "balance"),
        balance_kind="overdraft", wrap=0.2, trailer=["close"], head_style="bar",
        opening=dict(label="Opening balance"),
        hdr=dict(mast="right", title="Business statement", left=left, right=right,
                 addr=(left, 88), addr_bold=True, det=(left, 150, 84),
                 det_rows=["name", "acct", "period", "branch"],
                 summary=dict(style="box", x=S.r((340, 350), (320, 330)), y=96,
                              w=S.r((180, 190), (200, 210)), title="Account summary", size=7.5,
                              rows=[("Opening balance", "open"), ("Total debits", "tot_d"),
                                    ("Total credits", "tot_c"), ("Closing balance", "close")])),
        table_top1=S.v(222, 214), table_top2=S.v(76, 72),
        footer=dict(pos="right", other="Rimu Bank of Aotearoa - Branch {branch}", small="about"),
        stmts=pile(S, 3, (20, 55), S.v((2026, 2), (2026, 10))))


@design()
def rimu_letter(S):
    size = S.v(10, 9)
    left, right = S.r((50, 54), (40, 44)), S.r((560, 564), (570, 574))
    cols = mkcols(left, right, [
        ("date", S.r((96, 100), (92, 96)), ("Date",)),
        ("desc", None, ("Particulars",)),
        ("amount", S.r((82, 88), (78, 84)), ("Amount", "($)")),
        ("balance", S.r((82, 88), (84, 90)), ("Balance", "($)")),
    ])
    return layout(
        S, "rimu_letter", bank="rimu", product="Investment share account",
        note="US Letter page: money out in parentheses, full month names, two-line headings "
             "repeated on every page, a quarterly period.",
        page="Letter", font="Times", size=size, cols=cols, date_fmt="d Month yyyy",
        period_fmt="long", catalog=S.v("savings", "bank"), credit_p=S.v(0.4, 0.3),
        amount_style="paren", balance_style="paren",
        opening=dict(label="Opening balance", date=True), trailer=["close"], head_style="rule",
        hdr=dict(mast="plain", title="Statement", tagline="A mutual society",
                 left=left, right=right, addr=(left, 104), addr_size=10,
                 det=(S.r((320, 330), (340, 350)), 104, 88), det_size=9,
                 det_rows=["acct", "period", "stno"],
                 summary=dict(style="shaded", x=left, y=176, w=S.r((230, 240), (250, 260)),
                              title="Summary", size=9,
                              rows=[("Opening balance", "open"), ("Withdrawals", "tot_d"),
                                    ("Deposits", "tot_c"), ("Closing balance", "close")])),
        table_top1=S.v(296, 300), table_top2=S.v(84, 90),
        footer=dict(pos="center", fmt="Page {p} of {n}"),
        stmts=pile(S, 3, (14, 26), S.v((2025, 1), (2025, 7)), months=S.v(3, 2)))


@design()
def rimu_offsets(S):
    size = S.v(8.5, 8)
    left, right = S.r((50, 54), (56, 60)), S.r((528, 532), (522, 526))
    cols = mkcols(left, right, [
        ("date", S.r((50, 54), (56, 60)), ("Date",)),
        ("desc", None, S.v(("Transaction Description",), ("Narrative",))),
        ("debit", S.r((64, 70), (66, 72)), S.v(("Debits",), ("Money Out",))),
        ("credit", S.r((64, 70), (66, 72)), S.v(("Credits",), ("Money In",))),
        ("balance", S.r((74, 80), (76, 82)), S.v(("Balance",), ("Running Bal",))),
    ])
    return layout(
        S, "rimu_offsets", bank="rimu", product="Transaction account",
        note="The table moves by a different amount on every page, the money headings sit "
             "off their columns, the whole heading row is offset from the data, and the "
             "page drifts a little from page to page.",
        font="Helvetica", size=size, cols=cols, date_fmt=S.v("dd/mm/yyyy", "dd Mon yyyy"),
        period_fmt="short", catalog="bank", wrap=0.15,
        table_dx=("list", S.v([0, 18, -10, 24], [0, -14, 20, 8])),
        page_dx=("drift", S.v(1.5, -1.0)), head_shift=S.v(4.0, -5.0),
        head_dx=S.v({"debit": -12.0, "credit": -12.0, "balance": -8.0},
                    {"debit": 8.0, "credit": -10.0, "balance": 6.0}),
        max_rows=S.v(18, 16), opening=dict(label="Opening balance"), trailer=["close"],
        head_style="rule",
        hdr=dict(mast="logo", color=TEAL, title="Statement", left=left, right=right,
                 addr=(left, 96), det=(S.r((316, 322), (300, 306)), 96, 82),
                 det_rows=["acct", "period"], cont="line"),
        table_top1=176, table_top2=86,
        footer=dict(pos="center", small="about"),
        stmts=pile(S, 3, (30, 60), S.v((2025, 8), (2026, 3))))


@design("holdout")
def rimu_card_letter(S):
    size = 8.5
    left, right = 44, 568
    cols = mkcols(left, right, [
        ("date", 92, ("Trans date",)),
        ("pdate", 92, ("Posted",)),
        ("desc", None, ("Description",)),
        ("amount", 84, ("Amount",)),
    ])
    return layout(
        S, "rimu_card_letter", bank="rimu", product="Rewards Visa",
        note="US Letter credit card: full month dates, payments marked CR in a token slot.",
        page="Letter", font="Times", size=size, cols=cols, date_fmt="d Month yyyy",
        period_fmt="long", catalog="card", credits="card", balance_kind="card",
        bal_mode="none", amount_style="card_cr", hang=("amount",), wrap=0.3,
        opening=dict(label="Previous balance"), trailer=["close"], closing_label="New balance",
        head_style="rule", card=dict(limit=700000), rate_text="22.95% p.a.",
        hdr=dict(mast="plain", title="Card statement", left=left, right=right, addr=(left, 98),
                 det=(330, 98, 76), det_rows=["card", "period"],
                 sidebar=dict(style="box", x=330, y=132, w=220, title="Pay by",
                              rows=[("Payment due", "due"), ("Minimum payment", "minpay"),
                                    ("Credit limit", "limit")])),
        table_top1=232, table_top2=84,
        footer=dict(pos="right", small="about"),
        stmts=pile(S, 3, (16, 34), (2026, 5, 11)))


# ---------------------------------------------------------------------------

REDACTIONS = [
    dict(mode="remove", desc=1, amount=1, balance=1, row=1),
    dict(mode="overlay", desc=2, amount=1, balance=1),
    dict(mode="remove", desc=2, amount=1, row=1),
    dict(mode="overlay", desc=1, amount=2),
]
ACCOUNT_CASES = ["masked", "removed", "overlay", "masked"]


def sprinkle(split, Ls):
    """Spread redactions and account-identity cases over the piles, deterministically:
    about one statement in eight gets a redaction, one in twelve an account case."""
    rng = random.Random(SPLIT_SEED[split] + 99)
    ri = ai = 0
    for L in Ls:
        for spec in L["stmts"]:
            x = rng.random()
            if L["bundle"]:
                continue
            if x < 0.13:
                spec["redact"] = dict(REDACTIONS[ri % len(REDACTIONS)])
                ri += 1
            elif x < 0.22 and acct_printed(L):
                spec["acct"] = ACCOUNT_CASES[ai % len(ACCOUNT_CASES)]
                ai += 1


def designs_for(split):
    out = []
    for name, splits, fn in ARCH:
        if splits in ("both", split):
            L = fn(Split(split, name))
            if L["id"] != name:
                raise GenError("design %s returned id %s" % (name, L["id"]))
            out.append(L)
    sprinkle(split, out)
    return out


def exports_for(split):
    out = []
    for name, splits, fn in EXPORTS:
        if splits in ("both", split):
            E = fn(Split(split, "export:" + name))
            out.append(E)
    return out


def scan_picks(names):
    """About fifteen statements per split, spread evenly across the designs."""
    names = sorted(names)
    step = max(1, len(names) // 15)
    return set(names[::step][:15])


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--out", help="where to write the files")
    ap.add_argument("--split", choices=sorted(SPLIT_SEED), default="dev")
    ap.add_argument("--only", default=None, help="build just the cases whose name contains this")
    ap.add_argument("--list", action="store_true")
    ap.add_argument("--no-scans", action="store_true", help="skip the image-only copies")
    a = ap.parse_args()

    try:
        Ls = designs_for(a.split)
        Es = exports_for(a.split)
    except GenError as e:
        sys.exit("GENERATOR SELF-CHECK FAILED: %s" % e)
    pdf_names = ["%s_%d" % (L["id"], k + 1) for L in Ls for k in range(len(L["stmts"]))]
    picks = scan_picks(pdf_names)
    if a.list:
        for L in Ls:
            for k, _ in enumerate(L["stmts"]):
                name = "%s_%d" % (L["id"], k + 1)
                if not a.only or a.only in name:
                    print("%-30s %-18s %s%s" % (name, BANKS[L["bank_id"]]["name"], L["note"][:70],
                                                "  [+scan]" if name in picks else ""))
        for E in Es:
            for k, _ in enumerate(E["stmts"]):
                name = "%s_%d" % (E["id"], k + 1)
                if not a.only or a.only in name:
                    print("%-30s %-18s %s export" % (name, BANKS[E["bank"]]["name"], E["fmt"]))
        return 0
    if not a.out:
        ap.error("--out is required (or use --list to see the cases)")
    os.makedirs(a.out, exist_ok=True)
    try:
        import openpyxl  # noqa: F401
        have_xlsx = True
    except ImportError:
        have_xlsx = False
        print("NOTE: openpyxl is not installed, so the .xlsx exports are skipped "
              "(python3 -m pip install openpyxl).")

    index = []
    name = "?"
    try:
        for L in Ls:
            for k, spec in enumerate(L["stmts"]):
                name = "%s_%d" % (L["id"], k + 1)
                if a.only and a.only not in name:
                    continue
                name, truth, npg = build_case(L, k, spec, a.out)
                print("%-30s %2d rows %d pg  %s" % (name, truth["row_count"], npg,
                                                    " ".join(x for x in truth["features"]
                                                             if x.startswith(("id:", "redacted",
                                                                              "bundle", "multi")))))
                index.append({"case": name, "file": name + ".pdf", "bank": truth["bank"],
                              "layout": L["id"], "product": L["product"],
                              "source_format": "pdf", "rows": truth["row_count"], "pages": npg,
                              "features": truth["features"]})
                if name in picks and not a.no_scans:
                    sname = name + "_scan"
                    settings = make_scan(os.path.join(a.out, name + ".pdf"),
                                         os.path.join(a.out, sname + ".pdf"),
                                         SPLIT_SEED[a.split] + zlib.crc32(sname.encode()))
                    st = scan_truth(truth, settings)
                    write_json(os.path.join(a.out, sname + ".truth.json"), st)
                    print("%-30s    image-only scan, %d dpi, skew %.2f" % (sname, settings["dpi"],
                                                                          settings["skew"]))
                    index.append({"case": sname, "file": sname + ".pdf", "bank": st["bank"],
                                  "layout": L["id"], "product": L["product"],
                                  "source_format": "scan", "rows": st["row_count"], "pages": npg,
                                  "features": st["features"]})
        for E in Es:
            for k, spec in enumerate(E["stmts"]):
                name = "%s_%d" % (E["id"], k + 1)
                if a.only and a.only not in name:
                    continue
                if E["fmt"] == "xlsx" and not have_xlsx:
                    continue
                name, truth = build_export(E, k, spec, a.out)
                print("%-30s %2d rows  %s export" % (name, truth["row_count"], E["fmt"]))
                index.append({"case": name, "file": name + "." + E["fmt"], "bank": truth["bank"],
                              "layout": E["id"], "product": truth["product"],
                              "source_format": E["fmt"], "rows": truth["row_count"], "pages": None,
                              "features": truth["features"]})
    except GenError as e:
        sys.exit("GENERATOR SELF-CHECK FAILED in %s: %s" % (name, e))
    with open(os.path.join(a.out, "index.json"), "w") as f:
        json.dump(index, f, indent=1)
    by = {}
    for x in index:
        by[x["source_format"]] = by.get(x["source_format"], 0) + 1
    print("\n%s split: %s; %d banks, %d layouts -> %s"
          % (a.split, ", ".join("%d %s" % (v, k) for k, v in sorted(by.items())),
             len({x["bank"] for x in index}), len({x["layout"] for x in index}), a.out))
    return 0


if __name__ == "__main__":
    sys.exit(main())
