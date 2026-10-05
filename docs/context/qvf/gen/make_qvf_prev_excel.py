#!/usr/bin/env python3
"""make_qvf_prev_excel.py -- lookalikes of the workbook the QVF's "N/A - Previous
conversion" type reads back in (script.qvs lines 1858-2030, the "Excel" tab): the
QVF's own "Statement Data" table, downloaded to Excel by an analyst, corrected, and
uploaded again. Analysts hold archives of these.

Shape (see cards/OUTPUT.md section 1 and cards/qvf_prev_excel.md):
  * one sheet, headings on row 1, the Statement Data table's 23 columns in its
    order: File Name, Row ID, Sort Number, Bank, Account Name, Account Number,
    Other Party Account Name, Other Party Account Number, Code Description,
    Transaction Category, Transaction Type, Date, Transaction Time,
    Description as per bank statement, Transaction Code, Amount, Balance,
    Doc Reference Bank Statement, Doc Reference Bank Voucher, Year, Tax Year,
    Balance Check, Balance Pass;
  * Amount is UNSIGNED; Transaction Type (Deposit / Withdrawal) carries the sign;
  * row "<statement>-1" of each statement is the QVF's own "Opening Balance" row
    (Transaction Type Deposit, code 100, Amount = Balance = the opening balance);
  * the Qlik table sorts on Row ID as TEXT, so a download of a statement with 10 or
    more rows comes out 1-1, 1-10, 1-11, ..., 1-2, 1-20, 1-3 ...; "Sort Number" is
    the column that puts it back.

The answer key holds the transactions only (never the Opening Balance rows), signed,
in PRINTED order.
Run: python3 scratchpad/qvf/gen/make_qvf_prev_excel.py
"""
import datetime as dt
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import qvf_excel_common as C  # noqa: E402

GEN = "scratchpad/qvf/gen/make_qvf_prev_excel.py"
HEADS = ["File Name", "Row ID", "Sort Number", "Bank", "Account Name", "Account Number",
         "Other Party Account Name", "Other Party Account Number", "Code Description",
         "Transaction Category", "Transaction Type", "Date", "Transaction Time",
         "Description as per bank statement", "Transaction Code", "Amount", "Balance",
         "Doc Reference Bank Statement", "Doc Reference Bank Voucher", "Year", "Tax Year",
         "Balance Check", "Balance Pass"]
CASES = [
    dict(k=1, per=(2024, 3, 1, 1), stmts=[dict(n=8, opening=150000, bank="Westpac - Excel", code="03")],
         order="text", note="One statement of 8 rows: Row ID text order is also numeric order."),
    dict(k=2, per=(2024, 6, 1, 1), stmts=[dict(n=20, opening=82050, bank="BNZ - Excel", code="02")],
         order="text", note="One statement of 20 rows, downloaded in the table's Row ID TEXT order "
                            "(1-1, 1-10, 1-11 ... 1-2, 1-20, 1-21, 1-3 ...)."),
    dict(k=3, per=(2025, 1, 1, 1), stmts=[dict(n=14, opening=230000, bank="Kiwibank - Excel", code="38"),
                                          dict(n=11, opening=900000, bank="Kiwibank - Excel", code="38")],
         order="sort", note="Two statements (two accounts), put back in Sort Number order by the analyst."),
]


def tax_year(d):
    return d.year - 1 if d.month <= 3 else d.year


def build(case, out_dir):
    import openpyxl
    name = "qvf_prev_excel_%d" % case["k"]
    rng = C.rng_for(name)
    y, m, d, months = case["per"]
    start, end = C.period(y, m, d, months)
    first, last = rng.choice(C.HOLDERS)
    recs, truth_rows, accounts = [], [], []
    for si, s in enumerate(case["stmts"], 1):
        acct = C.acct_number(rng, s["code"])
        rows, op, cl = C.gen_rows(rng, start, end, s["n"], s["opening"], "positive", interest=False)
        accounts.append({"account_index": si - 1, "product": "Everyday account", "account_number": acct,
                         "opening_balance": op / 100.0, "closing_balance": cl / 100.0})
        base = dict(bank=s["bank"], name="%s %s" % (first, last), acct=acct)
        recs.append(dict(base, rid=(si, 1), date=rows[0]["date"], desc="Opening Balance", ttype="Deposit",
                         code="100", cdesc="Opening Balance", cat="Opening Balance", amt=op, bal=op,
                         oacct="", time=""))
        for j, t in enumerate(rows, 2):
            dep = t["dir"] == "C"
            desc = "  ".join(x for x in (t["ptype"], t["part"], t["code"], t["ref"], t["party"]))  # Qlik's & of blanks
            cdesc, cat, code = ("To be done", "To be done", "200" if dep else "400")
            if t["kind"] == "BP" and not dep:
                cat = "Transfers to third parties"
            recs.append(dict(base, rid=(si, j), date=t["date"], desc=desc, ttype="Deposit" if dep else "Withdrawal",
                             code=code, cdesc=cdesc, cat=cat, amt=t["amt"], bal=t["bal"],
                             oacct=t["oacct"] if s["bank"] != "BNZ - Excel" else "",
                             time=t["time"] if s["bank"] == "Kiwibank - Excel" else "", t=t, si=si - 1))
    for n, r in enumerate(recs, 1):
        r["sort"] = n
    if case["order"] == "text":
        recs.sort(key=lambda r: "%d-%d" % r["rid"])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for j, h in enumerate(HEADS, 1):
        ws.cell(row=1, column=j, value=h)
    for i, r in enumerate(recs, 2):
        vals = {"File Name": "statement_upload_%s" % name, "Row ID": "%d-%d" % r["rid"], "Sort Number": r["sort"],
                "Bank": r["bank"], "Account Name": r["name"], "Account Number": r["acct"],
                "Other Party Account Name": None, "Other Party Account Number": r["oacct"] or None,
                "Code Description": r["cdesc"], "Transaction Category": r["cat"], "Transaction Type": r["ttype"],
                "Date": dt.datetime(r["date"].year, r["date"].month, r["date"].day),
                "Transaction Time": r["time"] or None, "Description as per bank statement": r["desc"],
                "Transaction Code": r["code"], "Amount": r["amt"] / 100.0, "Balance": r["bal"] / 100.0,
                "Doc Reference Bank Statement": None, "Doc Reference Bank Voucher": None,
                "Year": r["date"].year, "Tax Year": tax_year(r["date"]),
                "Balance Check": r["bal"] / 100.0, "Balance Pass": "Pass"}
        for j, h in enumerate(HEADS, 1):
            c = ws.cell(row=i, column=j, value=vals[h])
            if h == "Date":
                c.number_format = "d/mm/yyyy"
            elif h in ("Amount", "Balance", "Balance Check"):
                c.number_format = '"$"#,##0.00;-"$"#,##0.00'
        if r["desc"] != "Opening Balance":
            truth_rows.append(C.truth_row(r["t"], r["desc"], account_index=r["si"] if len(case["stmts"]) > 1 else None))
    ws.cell(row=len(recs) + 3, column=1, value=C.SYNTHETIC)
    path = os.path.join(out_dir, name + ".xlsx")
    C.save_xlsx(wb, path)
    multi = len(case["stmts"]) > 1
    feats = ["qvf:previous_conversion", "format:xlsx", "unsigned_amount_with_type_words", "opening_balance_rows",
             "row_order:%s" % ("row_id_text" if case["order"] == "text" else "sort_number")]
    truth = C.truth_doc(
        name, GEN, "QVF lookalike. A Statement Converter 'Statement Data' download (synthetic). " + case["note"],
        "QVF export", "qvf_prev_excel", "Everyday account", "xlsx", None, None, feats, False,
        accounts[0]["opening_balance"], accounts[0]["closing_balance"], [], None)
    # printed order may not be date order, so the chain is checked on Sort Number order, not printed order
    truth["rows"] = truth_rows
    truth["row_count"] = len(truth_rows)
    if multi:
        truth["accounts"] = accounts
        truth["opening_balance"] = truth["closing_balance"] = None
    else:
        truth["opening_balance"], truth["closing_balance"] = accounts[0]["opening_balance"], accounts[0]["closing_balance"]
    C.write_json(os.path.join(out_dir, name + ".truth.json"), truth)
    return name, len(truth_rows)


def main():
    os.makedirs(C.OUT, exist_ok=True)
    for case in CASES:
        name, n = build(case, C.OUT)
        print("%-17s rows=%2d" % (name, n))


if __name__ == "__main__":
    main()
