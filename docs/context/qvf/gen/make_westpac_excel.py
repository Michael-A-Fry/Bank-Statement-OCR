#!/usr/bin/env python3
"""make_westpac_excel.py -- lookalikes of the "Westpac - Excel" statement the QVF
reads (script.qvs lines 2414-2519), with answer keys.

The shape the script expects (see cards/westpac_excel.md):
  * an .xlsx (Qlik `ooxml`) whose FIRST sheet is named "Sheet1";
  * B1 = account name, B2 = account number (A1 / A2 hold their labels);
  * the heading row is row 3 ("header is 2 lines");
  * hyphenated column names, including the one with an underscore:
    Date, Amount, Other-Party-Account-Number, Other-Party_Name, Source-Type,
    This-Party-Desc, This-Party-Code, This-Party-Reference, Running-Balance;
  * oldest first; Date is a real Excel date; Amount is signed;
  * a zero running balance can print as "." -- the script turns '.' into 0.00. Two
    ways that can reach a file are drawn: a text cell holding "." and a numeric 0
    under the number format "#,###.##" (which Excel and Qlik both print as ".").

Writes sets/excel/westpac_excel_<n>.xlsx + .truth.json.
Run: python3 scratchpad/qvf/gen/make_westpac_excel.py
"""
import datetime as dt
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import qvf_excel_common as C  # noqa: E402

GEN = "scratchpad/qvf/gen/make_westpac_excel.py"
HEADS = ["Date", "Amount", "Other-Party-Account-Number", "Other-Party_Name", "Source-Type",
         "This-Party-Desc", "This-Party-Code", "This-Party-Reference", "Running-Balance"]

CASES = [
    dict(k=1, per=(2024, 4, 1, 1), n=20, opening=312045, mode="positive", zero=None,
         note="Plain statement."),
    dict(k=2, per=(2024, 7, 1, 1), n=22, opening=84050, mode="zero", zero_at=(6, 15), zero="text",
         note="The balance is brought to exactly zero twice; each zero balance is a TEXT cell '.'."),
    dict(k=3, per=(2024, 9, 1, 1), n=18, opening=-340010, mode="od", zero=None,
         note="Overdrawn throughout: negative running balances."),
    dict(k=4, per=(2024, 12, 5, 1), n=24, opening=150075, mode="positive", zero=None,
         note="Period crosses a new year (5 Dec - 4 Jan); interest and RWT lines."),
    dict(k=5, per=(2025, 2, 1, 1), n=16, opening=0, mode="zero", zero_at=(9,), zero="fmt",
         note="A new account opened at zero; the running balance column uses the number format "
              "'#,###.##', so a zero balance PRINTS as '.' while the cell holds 0."),
    dict(k=6, per=(2025, 5, 1, 1), n=0, opening=None, mode="positive", zero=None,
         note="No transactions: the two account rows and the heading row only."),
    dict(k=7, per=(2025, 6, 1, 1), n=22, opening=25010, mode="cross", zero=None, holder="O'SAMPLE & CO LTD",
         note="Into overdraft and back; an account name with an apostrophe and an ampersand."),
]

SOURCE = {"EP": "EFTPOS", "AP": "AP", "DD": "DD", "BP": "BP", "ATM": "ATM", "FEE": "FEE",
          "TFR": "TFR", "DC": "DC", "DEP": "DEP", "INT": "INT", "RWT": "RWT"}


def build(case, out_dir):
    import openpyxl
    name = "westpac_excel_%d" % case["k"]
    rng = C.rng_for(name)
    y, m, d, months = case["per"]
    start, end = C.period(y, m, d, months)
    acct = C.acct_number(rng, "03")
    first, last = rng.choice(C.HOLDERS)
    holder = case.get("holder") or ("%s %s" % (first, last))
    if case["n"]:
        rows, opening, closing = C.gen_rows(rng, start, end, case["n"], case["opening"], case["mode"],
                                            zero_at=case.get("zero_at", ()))
    else:
        rows, opening, closing = [], None, None
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "Account Name"
    ws["B1"] = holder
    ws["A2"] = "Account Number"
    ws["B2"] = acct
    for j, h in enumerate(HEADS, 1):
        ws.cell(row=3, column=j, value=h)
    truth_rows = []
    r = 3
    for t in rows:
        r += 1
        src = SOURCE[t["kind"]]
        vals = {"Date": dt.datetime(t["date"].year, t["date"].month, t["date"].day),
                "Amount": (-t["amt"] if t["dir"] == "D" else t["amt"]) / 100.0,
                "Other-Party-Account-Number": t["oacct"] or None,
                "Other-Party_Name": t["party"] or None, "Source-Type": src,
                "This-Party-Desc": t["part"] or None, "This-Party-Code": t["code"] or None,
                "This-Party-Reference": t["ref"] or None, "Running-Balance": t["bal"] / 100.0}
        if t["bal"] == 0 and case["zero"] == "text":
            vals["Running-Balance"] = "."
        for j, h in enumerate(HEADS, 1):
            c = ws.cell(row=r, column=j, value=vals[h])
            if h == "Date":
                c.number_format = "dd/mm/yyyy"
            elif h == "Amount":
                c.number_format = "#,##0.00"
            elif h == "Running-Balance" and not isinstance(vals[h], str):
                c.number_format = "#,###.##" if case["zero"] == "fmt" else "#,##0.00"
        words = [t["oacct"], t["party"], src, t["part"], t["code"], t["ref"]]
        desc = " ".join(w for w in words if w)
        truth_rows.append(C.truth_row(t, desc))
    other = wb.create_sheet("Sheet2")
    other["A1"] = C.SYNTHETIC
    path = os.path.join(out_dir, name + ".xlsx")
    C.save_xlsx(wb, path)
    feats = ["bank:westpac", "format:xlsx", "qvf:westpac_excel", "header_row:3", "signed",
             "mode:%s" % case["mode"]]
    if case["zero"] == "text":
        feats.append("zero_balance_as_text_dot")
    if case["zero"] == "fmt":
        feats.append("zero_balance_formatted_as_dot")
    if not rows:
        feats.append("no_transactions")
    if rows and start.year != end.year:
        feats.append("crosses_new_year")
    truth = C.truth_doc(
        name, GEN, "QVF lookalike. Westpac - Excel export (synthetic). " + case["note"],
        "Westpac", "qvf_westpac_excel", "Everyday account", "xlsx", "03", acct, feats, False,
        None if opening is None else opening / 100.0, None if closing is None else closing / 100.0,
        truth_rows)
    C.write_json(os.path.join(out_dir, name + ".truth.json"), truth)
    return name, len(truth_rows)


def main():
    os.makedirs(C.OUT, exist_ok=True)
    for case in CASES:
        name, n = build(case, C.OUT)
        print("%-17s rows=%2d" % (name, n))


if __name__ == "__main__":
    main()
