#!/usr/bin/env python3
"""make_kiwibank_excel.py -- lookalikes of the "Kiwibank - Excel" statement the QVF
reads (script.qvs lines 2196-2413), with answer keys.

The shape the script expects (see cards/kiwibank_excel.md):
  * an .xlsx (Qlik `ooxml`), first sheet, a transaction table whose FIRST column is
    headed "Customer";
  * the heading row is on row 1, 3, 4 or 8 -- the only four shapes the script knows:
      row 1: no preamble, no account holder name;
      row 3: the holder's first and last names in B2 and C2;
      row 4: names in B2 and C2, and NO ThisPartyParticulars / ThisPartyCode /
             ThisPartyReference columns;
      row 8: seven preamble rows, names in B7 and C7;
  * columns (by name): Customer, KiwiAcc, ProcessDate ('YYYY-MM-DD'), ReceiptTime,
    Narration1, PayingBankDRN, PayeeDetails, ThisPartyParticulars, ThisPartyCode,
    ThisPartyReference, OtherPartyAcc, Amount (signed), Running Balance;
  * NEWEST FIRST (the script reverses the file);
  * one file may hold several accounts (KiwiAcc changes); each is its own statement.

Writes sets/excel/kiwibank_excel_<n>.xlsx + .truth.json.
Run: python3 scratchpad/qvf/gen/make_kiwibank_excel.py
"""
import datetime as dt
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import qvf_excel_common as C  # noqa: E402

GEN = "scratchpad/qvf/gen/make_kiwibank_excel.py"

CASES = [
    dict(k=1, hdr=1, per=(2024, 2, 1, 1), accts=[dict(n=18, opening=245010, mode="positive")],
         note="Heading on row 1: no preamble, no holder name."),
    dict(k=2, hdr=3, per=(2024, 5, 1, 1), accts=[dict(n=24, opening=118075, mode="positive")],
         note="Heading on row 3; holder's names in B2 / C2."),
    dict(k=3, hdr=4, per=(2024, 8, 1, 1), accts=[dict(n=20, opening=67340, mode="positive")],
         note="Heading on row 4; no ThisParty particulars / code / reference columns."),
    dict(k=4, hdr=8, per=(2024, 10, 1, 1), accts=[dict(n=22, opening=356000, mode="positive")],
         note="Heading on row 8 after seven preamble rows; names in B7 / C7."),
    dict(k=5, hdr=3, per=(2025, 1, 1, 1),
         accts=[dict(n=16, opening=152500, mode="positive"),
                dict(n=12, opening=1250000, mode="positive", product="Savings account")],
         note="Two accounts in one file (KiwiAcc changes part-way): two statements."),
    dict(k=6, hdr=1, per=(2024, 12, 15, 1), accts=[dict(n=26, opening=30500, mode="cross")],
         note="Period crosses a new year; goes into overdraft (negative running balance) and back."),
    dict(k=7, hdr=8, per=(2025, 3, 1, 1), accts=[dict(n=20, opening=88020, mode="positive")],
         excel_dates=True,
         note="ProcessDate stored as real Excel dates (yyyy-mm-dd format) and ReceiptTime as "
              "Excel times, not text."),
]


def build(case, out_dir):
    import openpyxl
    name = "kiwibank_excel_%d" % case["k"]
    rng = C.rng_for(name)
    y, m, d, months = case["per"]
    start, end = C.period(y, m, d, months)
    first, last = rng.choice(C.HOLDERS)
    custno = rng.randint(1000000, 9999999)
    with_tp = case["hdr"] != 4
    accts = []
    for a in case["accts"]:
        acct = C.acct_number(rng, "38")
        rows, op, cl = C.gen_rows(rng, start, end, a["n"], a["opening"], a["mode"])
        accts.append(dict(acct=acct, rows=rows, opening=op, closing=cl,
                          product=a.get("product", "Everyday account")))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Transactions"
    pre = []
    if case["hdr"] == 3:
        pre = [["Customer No.", "First Name", "Last Name"], [custno, first, last]]
    elif case["hdr"] == 4:
        pre = [["Customer No.", "First Name", "Last Name"], [custno, first, last],
               ["Accounts listed: %d" % len(accts)]]
    elif case["hdr"] == 8:
        pre = [["KIWIBANK - TRANSACTION DETAIL"], ["Request reference: REQ-%06d" % rng.randint(0, 999999)],
               ["Date range: %s to %s" % (start.isoformat(), end.isoformat())],
               ["Prepared: %s" % (end + dt.timedelta(days=5)).isoformat()],
               ["Accounts: %s" % ", ".join(a["acct"] for a in accts)],
               ["Customer No.", "First Name", "Last Name"], [custno, first, last]]
    for i, vals in enumerate(pre, 1):
        for j, v in enumerate(vals, 1):
            ws.cell(row=i, column=j, value=v)
    heads = ["Customer", "KiwiAcc", "ProcessDate", "ReceiptTime", "Narration1", "PayingBankDRN",
             "PayeeDetails"]
    if with_tp:
        heads += ["ThisPartyParticulars", "ThisPartyCode", "ThisPartyReference"]
    heads += ["OtherPartyAcc", "Amount", "Running Balance"]
    r = case["hdr"]
    for j, h in enumerate(heads, 1):
        ws.cell(row=r, column=j, value=h)
    truth_rows, accounts = [], []
    for ai, a in enumerate(accts):
        accounts.append({"account_index": ai, "product": a["product"], "account_number": a["acct"],
                         "opening_balance": a["opening"] / 100.0, "closing_balance": a["closing"] / 100.0})
        for t in reversed(a["rows"]):                 # newest first
            r += 1
            drn = ("%010d" % rng.randint(0, 9999999999)) if (t["dir"] == "C" and t["kind"] in ("DC", "BP")) else ""
            vals = {"Customer": custno, "KiwiAcc": a["acct"], "Narration1": t["narr"],
                    "PayingBankDRN": drn or None, "PayeeDetails": t["party"] or None,
                    "ThisPartyParticulars": t["part"] or None, "ThisPartyCode": t["code"] or None,
                    "ThisPartyReference": t["ref"] or None, "OtherPartyAcc": t["oacct"] or None,
                    "Amount": (-t["amt"] if t["dir"] == "D" else t["amt"]) / 100.0,
                    "Running Balance": t["bal"] / 100.0}
            if case.get("excel_dates"):
                vals["ProcessDate"] = dt.datetime(t["date"].year, t["date"].month, t["date"].day)
                hh, mm, ss = (int(x) for x in t["time"].split(":"))
                vals["ReceiptTime"] = dt.time(hh, mm, ss)
            else:
                vals["ProcessDate"] = t["date"].isoformat()
                vals["ReceiptTime"] = t["time"]
            for j, h in enumerate(heads, 1):
                c = ws.cell(row=r, column=j, value=vals[h])
                if h == "ProcessDate" and case.get("excel_dates"):
                    c.number_format = "yyyy-mm-dd"
                elif h == "ReceiptTime" and case.get("excel_dates"):
                    c.number_format = "hh:mm:ss"
                elif h in ("Amount", "Running Balance"):
                    c.number_format = "0.00"
            words = [t["narr"], drn, t["party"]] + ([t["part"], t["code"], t["ref"]] if with_tp else []) + \
                    [t["oacct"]]
            desc = " ".join([a["acct"], t["time"]] + [w for w in words if w])
            truth_rows.append(C.truth_row(t, desc, account_index=ai if len(accts) > 1 else None))
    ws.cell(row=r + 2, column=1, value=C.SYNTHETIC)
    path = os.path.join(out_dir, name + ".xlsx")
    C.save_xlsx(wb, path)
    feats = ["bank:kiwibank", "format:xlsx", "qvf:kiwibank_excel", "header_row:%d" % case["hdr"],
             "newest_first", "signed"] + (["no_this_party_columns"] if not with_tp else []) + \
            (["accounts:%d" % len(accts)] if len(accts) > 1 else []) + \
            (["excel_dates", "excel_times"] if case.get("excel_dates") else ["text_iso_dates"]) + \
            (["overdrawn_rows"] if any(t["bal"] < 0 for a in accts for t in a["rows"]) else []) + \
            (["crosses_new_year"] if start.year != end.year else [])
    multi = len(accts) > 1
    truth = C.truth_doc(
        name, GEN, "QVF lookalike. Kiwibank - Excel export (synthetic). " + case["note"],
        "Kiwibank", "qvf_kiwibank_excel", "Everyday account", "xlsx", "38",
        accts[0]["acct"], feats, True, accts[0]["opening"] / 100.0, accts[0]["closing"] / 100.0,
        truth_rows, accounts=accounts if multi else None)
    C.write_json(os.path.join(out_dir, name + ".truth.json"), truth)
    return name, len(truth_rows)


def main():
    os.makedirs(C.OUT, exist_ok=True)
    for case in CASES:
        name, n = build(case, C.OUT)
        print("%-18s rows=%2d" % (name, n))


if __name__ == "__main__":
    main()
