#!/usr/bin/env python3
"""make_bnz_excel.py -- lookalikes of the "BNZ - Excel" statement the QVF reads
(script.qvs lines 2033-2195), with answer keys.

The shape the script expects (see cards/bnz_excel.md):
  * a legacy Excel 97-2003 workbook (Qlik `biff`), first sheet;
  * row 1 is the label row of the first load: B1 and E1 are blank (Qlik names those
    columns F2 and F5); the account holder's name is in E2;
  * a VARIABLE number of blank rows (the reason the script searches for the header);
  * the account number in column B of the row just above the header;
  * the header row: A "Date", B "Payment Type", C/D/E blank (F3/F4/F5), F
    "Particulars", G "Withdrawals", H "Deposits", I "Balance";
  * oldest first; dates are real Excel dates; Withdrawals / Deposits are positive
    numbers in their own columns; Balance is a positive number and an overdrawn
    balance is the SAME positive number with a number format that prints " OD";
  * optionally, after the statement rows, an "unstatemented" (pending) block: a
    sub-heading row with "Date" in A and "Name of Other Party" in E, then rows whose
    Date is TEXT 'DD Mon YY', whose only words are in E, and that have no balance.
    The QVF outputs those rows as transactions; the answer key here does NOT,
    because the new tool's product owner ruled pending items are not transactions
    (docs/context/auto-reading-spec.md section 2). They are named in the truth's
    note and features.

Writes sets/excel/bnz_excel_<n>.xlsx (+ .truth.json), the shape the scorer reads,
and sets/excel/xls/bnz_excel_<n>.xls, the real BIFF shape, for the front-door test.

Run: python3 scratchpad/qvf/gen/make_bnz_excel.py
"""
import datetime as dt
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import qvf_excel_common as C  # noqa: E402

GEN = "scratchpad/qvf/gen/make_bnz_excel.py"
MON = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
FMT_POS = "#,##0.00"
FMT_OD = '#,##0.00" OD"'

CASES = [
    # name, period (y, m, d, months), n, opening (cents), mode, blank rows, pending rows, note
    dict(k=1, per=(2024, 3, 1, 1), n=18, opening=184250, mode="positive", blanks=1, pending=0,
         note="Plain statement: one blank row above the account number, balances in credit."),
    dict(k=2, per=(2024, 6, 1, 1), n=22, opening=96010, mode="positive", blanks=3, pending=3,
         note="Three blank rows; an 'unstatemented' block of 3 pending rows after the statement "
              "(text dates 'DD Mon YY', words only in column E, no balance)."),
    dict(k=3, per=(2024, 9, 1, 1), n=20, opening=-215075, mode="od", blanks=2, pending=0,
         note="Overdrawn throughout: every balance is a POSITIVE number formatted '#,##0.00\" OD\"'."),
    dict(k=4, per=(2024, 11, 1, 1), n=24, opening=40500, mode="cross", blanks=1, pending=0,
         note="Goes into overdraft and back: OD only by number format on the middle rows."),
    dict(k=5, per=(2024, 12, 10, 1), n=26, opening=532090, mode="positive", blanks=4, pending=0,
         note="Period crosses a new year (10 Dec - 9 Jan); interest and RWT lines on the last day."),
    dict(k=6, per=(2025, 2, 1, 1), n=0, opening=None, mode="positive", blanks=2, pending=0,
         note="No transactions: the header row and nothing under it."),
    dict(k=7, per=(2025, 4, 1, 1), n=16, opening=-88015, mode="od", blanks=2, pending=2, title=False,
         note="Overdrawn throughout AND an unstatemented block of 2 pending rows, with no title row: "
              "only the repeated sub-heading ('Date' ... 'Name of Other Party') marks where it starts."),
]


def build(case, out_dir, xls_dir):
    import openpyxl
    name = "bnz_excel_%d" % case["k"]
    rng = C.rng_for(name)
    y, m, d, months = case["per"]
    start, end = C.period(y, m, d, months)
    acct = C.acct_number(rng, "02")
    first, last = rng.choice(C.HOLDERS)
    holder = "%s %s" % (first, last)
    if case["n"]:
        rows, opening, closing = C.gen_rows(rng, start, end, case["n"], case["opening"], case["mode"])
    else:
        rows, opening, closing = [], None, None
    pend = []
    for j in range(case["pending"]):
        e = C.pick(rng, C.EVERYDAY, "D" if j % 2 == 0 else None)
        t = C.txn(rng, e, end + dt.timedelta(days=1 + j), C.amount_of(rng, e))
        t["bal"] = None
        pend.append(t)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "Account Transactions"
    ws["A2"] = "Account Name:"
    ws["E2"] = holder
    r = 3 + case["blanks"]          # rows 3 .. 2+blanks stay empty
    ws.cell(row=r, column=1, value="Account Number:")
    ws.cell(row=r, column=2, value=acct)
    r += 1
    hdr = r
    for j, h in enumerate(["Date", "Payment Type", None, None, None, "Particulars",
                           "Withdrawals", "Deposits", "Balance"], 1):
        if h is not None:
            ws.cell(row=r, column=j, value=h)
    truth_rows = []
    for t in rows:
        r += 1
        c = ws.cell(row=r, column=1, value=dt.datetime(t["date"].year, t["date"].month, t["date"].day))
        c.number_format = "dd/mm/yyyy"
        ws.cell(row=r, column=2, value=t["ptype"])
        ws.cell(row=r, column=3, value=t["code"] or None)
        ws.cell(row=r, column=4, value=t["ref"] or None)
        ws.cell(row=r, column=5, value=t["party"] or None)
        ws.cell(row=r, column=6, value=t["part"] or None)
        col = 7 if t["dir"] == "D" else 8
        c = ws.cell(row=r, column=col, value=t["amt"] / 100.0)
        c.number_format = FMT_POS
        c = ws.cell(row=r, column=9, value=abs(t["bal"]) / 100.0)
        c.number_format = FMT_OD if t["bal"] < 0 else FMT_POS
        desc = " ".join(x for x in (t["ptype"], t["code"], t["ref"], t["party"], t["part"]) if x)
        truth_rows.append(C.truth_row(t, desc))
    if pend:
        r += 2                       # one blank row
        if case.get("title", True):
            ws.cell(row=r, column=2, value="Unstatemented Transactions")
            r += 1
        ws.cell(row=r, column=1, value="Date")
        ws.cell(row=r, column=5, value="Name of Other Party")
        ws.cell(row=r, column=7, value="Withdrawals")
        ws.cell(row=r, column=8, value="Deposits")
        for t in pend:
            r += 1
            ws.cell(row=r, column=1, value="%02d %s %02d" % (t["date"].day, MON[t["date"].month - 1],
                                                            t["date"].year % 100))
            words = " ".join(x for x in (t["party"] or t["narr"], t["part"]) if x)
            ws.cell(row=r, column=5, value=words)
            c = ws.cell(row=r, column=7 if t["dir"] == "D" else 8, value=t["amt"] / 100.0)
            c.number_format = FMT_POS
            # NOT a truth row: the product owner's rule (auto-reading-spec section 2,
            # 3 Oct 2026) is that pending items are not transactions. The QVF does
            # output them (see cards/bnz_excel.md); that difference is reported as an
            # output-field gap, not scored as a reading fault.
    r += 2
    ws.cell(row=r, column=1, value=C.SYNTHETIC)
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["E"].width = 28
    path = os.path.join(out_dir, name + ".xlsx")
    C.save_xlsx(wb, path)

    feats = ["bank:bnz", "format:xlsx", "qvf:bnz_excel", "header_row:%d" % hdr,
             "blank_rows:%d" % case["blanks"], "mode:%s" % case["mode"]]
    if pend:
        feats += ["unstatemented_block", "pending_rows_not_in_truth:%d" % len(pend), "mixed_date_cells"]
    if any(t["bal"] < 0 for t in rows):
        feats += ["od_by_number_format"]
    if not rows:
        feats += ["no_transactions"]
    truth = C.truth_doc(
        name, GEN, "QVF lookalike. BNZ - Excel export (synthetic). " + case["note"],
        "BNZ", "qvf_bnz_excel", "Everyday account", "xlsx", "02", acct, feats, False,
        None if opening is None else opening / 100.0, None if closing is None else closing / 100.0,
        truth_rows)
    C.write_json(os.path.join(out_dir, name + ".truth.json"), truth)
    xls = C.to_xls(path, xls_dir)
    return name, len(truth_rows), xls


def main():
    out = C.OUT
    os.makedirs(out, exist_ok=True)
    xls_dir = os.path.join(out, "xls")
    for case in CASES:
        name, n, xls = build(case, out, xls_dir)
        print("%-14s rows=%2d  xls=%s" % (name, n, "yes" if xls else "NO (LibreOffice missing)"))


if __name__ == "__main__":
    main()
