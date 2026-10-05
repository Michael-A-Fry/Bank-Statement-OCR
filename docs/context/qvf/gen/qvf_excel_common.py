"""qvf_excel_common.py -- shared pieces of the three QVF Excel lookalike generators
(make_bnz_excel.py, make_kiwibank_excel.py, make_westpac_excel.py).

Everything here is invented: names, account digits, merchants, figures. Account
numbers have the NZ shape BB-bbbb-AAAAAAA-SSS with the issuing bank's real 2-digit
code, exactly as tools/synth/make_layouts.py does, so bank identification can work.

The truth file is the one tools/synth/make_layouts.py writes for an export
(case, generator, note, bank, layout, product, source_format, account_bank_code,
account_number, account_redaction, features, row_order, opening_balance,
closing_balance, removed_rows, row_count, rows[{date, description, debit, credit,
balance}], plus `accounts` + rows[].account_index for a file holding several
accounts). debit = money out, credit = money in, both positive; balance = the
running balance printed on that row, else null. Rows in PRINTED order.

Deterministic: the per-case seed is zlib.crc32 of the case name, and every .xlsx is
re-zipped with fixed timestamps (make_layouts.fixed_zip).
"""

import datetime as dt
import io
import json
import os
import random
import re


import sys
import zlib

REPO = "/home/user/Bank-Statement-OCR"
HERE = os.path.dirname(os.path.abspath(__file__))
QVF = os.path.dirname(HERE)
OUT = os.path.join(QVF, "sets", "excel")
sys.path.insert(0, os.path.join(REPO, "tools", "synth"))
import make_layouts as ML  # noqa: E402  -- fixed_zip, write_json (read-only reuse)

SEED = 20261005
SYNTHETIC = "SYNTHETIC TEST DOCUMENT - NOT A REAL STATEMENT"


class GenError(Exception):
    pass


def rng_for(name):
    return random.Random(SEED + zlib.crc32(name.encode("utf-8")))


# ---- invented words ---------------------------------------------------------------
WORDS = dict(
    grocer=["SAMPLE FRESH MARKET", "EXAMPLE GROCERS", "DEMO FOODS", "TEST VALUE MART"],
    fuel=["TEST FUEL STOP", "SAMPLE SERVICE STN", "DEMO GAS"],
    cafe=["CAFE EXEMPLAR", "THE MOCK BEAN", "SPECIMEN BAKERY"],
    employer=["SAMPLE TRADING LTD", "EXAMPLE HOLDINGS LTD", "DEMO LOGISTICS"],
    landlord=["DEMO PROPERTY MGMT", "TEST RENTALS LTD"],
    utility=["SAMPLE POWER CO", "EXAMPLE WATER", "MOCK GAS SUPPLY"],
    telco=["MOCK MOBILE", "SAMPLE TELECOM"],
    person=["A EXAMPLE", "B SAMPLE", "C TESTER", "D DEMO", "E PLACEHOLDER"],
    council=["SAMPLETOWN DISTRICT COUNCIL", "EXAMPLE CITY COUNCIL"],
    insurer=["EXAMPLE INSURANCE", "DEMO MUTUAL"],
    online=["SAMPLESHOP ONLINE", "DEMO STREAMING"],
    atm=["SAMPLETOWN MALL", "EXAMPLE ST", "DEMO PLAZA"],
)

# Account holders (invented). First/last split for the Kiwibank preamble.
HOLDERS = [("JORDAN", "SAMPLE"), ("ALEX", "EXAMPLE"), ("PAT", "TESTER"),
           ("RIVER", "DEMO"), ("SAM", "PLACEHOLDER")]


def acct_number(rng, code):
    return "%s-%04d-%07d-%03d" % (code, rng.randint(1, 9999), rng.randint(0, 9999999),
                                  rng.randint(0, 30))


def other_acct(rng):
    code = rng.choice(["01", "02", "03", "06", "12", "15", "38"])
    return acct_number(rng, code)


def _num(rng, key):
    if key == "card4":
        return "%04d" % rng.randint(1000, 9999)
    if key == "ref":
        return "%06d" % rng.randint(0, 999999)
    if key == "inv":
        return "%05d" % rng.randint(10000, 99999)
    if key == "custno":
        return "%07d" % rng.randint(1000000, 9999999)
    if key == "wk":
        return "%d" % rng.randint(1, 52)
    raise GenError("unknown placeholder %s" % key)


def fill(rng, s):
    def sub(m):
        k = m.group(1)
        return rng.choice(WORDS[k]) if k in WORDS else _num(rng, k)
    return re.sub(r"\{(\w+)\}", sub, s)


# ---- the generic NZ everyday-account catalogue -------------------------------------
# kind  = the bank's short source/payment code (BP = bill payment, AP = automatic
#         payment, DC = direct credit, DD = direct debit, EP = eftpos ...)
# ptype = the same, spelt out the way a "Payment Type" column prints it
# party = other party's name; part/code/ref = NZ particulars / code / reference
def E(d, w, lo, hi, kind, ptype, narr, party, part, code, ref, xfer=False, mult=0, fixed=None):
    return dict(dir=d, w=w, lo=lo, hi=hi, kind=kind, ptype=ptype, narr=narr, party=party,
                part=part, code=code, ref=ref, xfer=xfer, mult=mult, fixed=fixed)


EVERYDAY = [
    E("D", 7, 12, 310, "EP", "Eft-Pos", "EFTPOS PURCHASE", "{grocer}", "EFTPOS", "{card4}", "{ref}"),
    E("D", 3, 35, 160, "EP", "Eft-Pos", "EFTPOS PURCHASE", "{fuel}", "EFTPOS", "{card4}", "{ref}"),
    E("D", 2, 4, 38, "EP", "Eft-Pos", "EFTPOS PURCHASE", "{cafe}", "EFTPOS", "{card4}", "{ref}"),
    E("D", 2, 380, 760, "AP", "Automatic Payment", "AUTOMATIC PAYMENT", "{landlord}", "RENT", "", "WK {wk}"),
    E("D", 2, 60, 340, "DD", "Direct Debit", "DIRECT DEBIT", "{utility}", "{custno}", "POWER", "INV {inv}"),
    E("D", 2, 90, 1650, "BP", "Bill Payment", "BILL PAYMENT", "{council}", "RATES", "{custno}", "INV {inv}"),
    E("D", 2, 20, 400, "ATM", "ATM Withdrawal", "ATM WITHDRAWAL", "ATM {atm}", "", "", "", mult=20),
    E("D", 1, 0, 0, "FEE", "Fee", "ACCOUNT FEE", "", "MONTHLY FEE", "", "", fixed=500),
    E("D", 2, 50, 1500, "TFR", "Transfer", "TRANSFER TO", "{person}", "TFR", "", "{ref}", xfer=True),
    E("D", 1, 35, 190, "BP", "Bill Payment", "BILL PAYMENT", "{telco}", "MOBILE", "", "REF {ref}"),
    E("C", 3, 1800, 6200, "DC", "Direct Credit", "DIRECT CREDIT", "{employer}", "SALARY", "", "PAY {wk}"),
    E("C", 2, 50, 1500, "TFR", "Transfer", "TRANSFER FROM", "{person}", "TFR", "", "{ref}", xfer=True),
    E("C", 1, 20, 900, "DEP", "Deposit", "DEPOSIT", "", "BRANCH DEPOSIT", "", ""),
    E("C", 1, 50, 900, "BP", "Bill Payment", "BILL PAYMENT", "{person}", "BOARD", "", "WK {wk}", xfer=True),
]

INTEREST = E("C", 0, 0.15, 24, "INT", "Interest", "CREDIT INTEREST", "", "INTEREST", "", "")
RWT = E("D", 0, 0.02, 3, "RWT", "Tax", "RWT ON INTEREST", "", "RWT", "", "")


def pick(rng, cat, d=None):
    cands = [e for e in cat if d is None or e["dir"] == d]
    tot = sum(e["w"] for e in cands)
    x = rng.uniform(0, tot)
    for e in cands:
        x -= e["w"]
        if x <= 0:
            return e
    return cands[-1]


def amount_of(rng, e):
    if e["fixed"] is not None:
        return e["fixed"]
    if e["mult"]:
        return rng.randint(e["lo"] // e["mult"], e["hi"] // e["mult"]) * e["mult"] * 100
    return rng.randint(int(round(e["lo"] * 100)), int(round(e["hi"] * 100)))


def txn(rng, e, date, amt):
    return dict(date=date, dir=e["dir"], amt=amt, kind=e["kind"], ptype=e["ptype"],
                narr=e["narr"], party=fill(rng, e["party"]), part=fill(rng, e["part"]),
                code=fill(rng, e["code"]), ref=fill(rng, e["ref"]),
                oacct=other_acct(rng) if e["xfer"] or e["kind"] == "DC" else "",
                time="%02d:%02d:%02d" % (rng.randint(6, 22), rng.randint(0, 59), rng.randint(0, 59)))


def dates_in(rng, start, end, n):
    days = (end - start).days
    ds = sorted(start + dt.timedelta(days=rng.randint(0, days)) for _ in range(n))
    return ds


def gen_rows(rng, start, end, n, opening, mode="positive", zero_at=(), interest=True):
    """n transactions between start and end, oldest first, with running balances
    (cents). mode: positive | od (overdrawn throughout) | cross (into overdraft and
    back) | zero (positive; the balance is brought to exactly 0.00 at each index in
    zero_at by a transfer of the whole balance, and the next row is a credit)."""
    ds = dates_in(rng, start, end, n)
    rows, bal = [], opening
    third = max(1, n // 3)
    for i, d in enumerate(ds):
        m = mode
        if mode == "cross":
            m = "positive" if i < third else ("od" if i < 2 * third else "positive")
        forced = None
        if mode == "zero" and (i in zero_at) and bal > 0:
            e = E("D", 0, 0, 0, "TFR", "Transfer", "TRANSFER TO", "{person}", "SAVINGS", "", "{ref}", xfer=True)
            forced = txn(rng, e, d, bal)
        elif mode == "zero" and (i - 1) in zero_at:
            forced = txn(rng, EVERYDAY[10], d, amount_of(rng, EVERYDAY[10]))
        elif mode == "cross" and i == third and bal >= 0:
            e = EVERYDAY[3]  # rent: into the overdraft
            forced = txn(rng, e, d, bal + rng.randint(5000, 40000))
        elif mode == "cross" and i == 2 * third and bal < 0:
            e = EVERYDAY[10]  # salary: back out of it
            forced = txn(rng, e, d, -bal + rng.randint(20000, 150000))
        if forced is not None:
            t = forced
        else:
            e = pick(rng, EVERYDAY)
            a = amount_of(rng, e)
            if m == "positive" and e["dir"] == "D" and bal - a < 2000:
                if bal - 2000 > 500 and e["fixed"] is None and not e["mult"]:
                    a = rng.randint(100, bal - 2000)
                else:
                    e = pick(rng, EVERYDAY, "C")
                    a = amount_of(rng, e)
            if m == "od" and e["dir"] == "C" and bal + a > -100:
                if -bal - 100 > 500:
                    a = rng.randint(100, -bal - 100)
                else:
                    e = pick(rng, EVERYDAY, "D")
                    a = amount_of(rng, e)
            if mode == "zero" and e["dir"] == "D" and bal - a < 0:
                e = pick(rng, EVERYDAY, "C")
                a = amount_of(rng, e)
            t = txn(rng, e, d, a)
        bal = bal - t["amt"] if t["dir"] == "D" else bal + t["amt"]
        t["bal"] = bal
        rows.append(t)
    if interest and n >= 10 and mode != "od":
        # interest and its withholding tax on the last day: real-world "interest lines"
        d = end
        for e in (INTEREST, RWT):
            a = amount_of(rng, e)
            t = txn(rng, e, d, a)
            bal = bal - a if e["dir"] == "D" else bal + a
            t["bal"] = bal
            rows.append(t)
    # same-day rows in time order
    for i in range(1, len(rows)):
        if rows[i]["date"] == rows[i - 1]["date"] and rows[i]["time"] < rows[i - 1]["time"]:
            rows[i]["time"] = rows[i - 1]["time"][:6] + "%02d" % min(59, int(rows[i - 1]["time"][6:]) + 1)
    check_mode(rows, opening, mode)
    return rows, opening, bal


def check_mode(rows, opening, mode):
    b = opening
    for t in rows:
        b = b - t["amt"] if t["dir"] == "D" else b + t["amt"]
        if b != t["bal"]:
            raise GenError("balance chain broken")
        if t["amt"] < 0:
            raise GenError("negative amount")
    if mode == "od" and any(t["bal"] >= 0 for t in rows):
        raise GenError("od statement has a non-negative balance")
    if mode == "cross" and not (any(t["bal"] < 0 for t in rows) and any(t["bal"] > 0 for t in rows)):
        raise GenError("cross statement does not cross")
    if mode == "zero" and not any(t["bal"] == 0 for t in rows):
        raise GenError("zero statement never reaches zero")


def truth_row(t, desc, balance_printed=True, account_index=None):
    r = {"date": t["date"].isoformat(), "description": desc,
         "debit": t["amt"] / 100.0 if t["dir"] == "D" else None,
         "credit": t["amt"] / 100.0 if t["dir"] == "C" else None,
         "balance": (t["bal"] / 100.0) if (balance_printed and t.get("bal") is not None) else None}
    if account_index is not None:
        r["account_index"] = account_index
    return r


def verify_truth(truth):
    """The truth's own arithmetic, per account, in date order (rows with no printed
    balance carry the chain on)."""
    groups = {}
    for r in truth["rows"]:
        groups.setdefault(r.get("account_index", 0), []).append(r)
    accts = truth.get("accounts") or [{"account_index": 0, "opening_balance": truth["opening_balance"],
                                       "closing_balance": truth["closing_balance"]}]
    newest = truth["row_order"] == "newest_first"
    for a in accts:
        rows = groups.get(a["account_index"], [])
        if not rows:
            continue
        seq = rows[::-1] if newest else rows
        b = round(a["opening_balance"] * 100)
        for r in seq:
            if (r["debit"] is None) == (r["credit"] is None):
                raise GenError("%s: a row needs exactly one of debit/credit" % truth["case"])
            b += -round(r["debit"] * 100) if r["debit"] is not None else round(r["credit"] * 100)
            if r["balance"] is not None and round(r["balance"] * 100) != b:
                raise GenError("%s: printed balance %s, arithmetic %s" % (truth["case"], r["balance"], b / 100))
            if not re.fullmatch(r"\d{4}-\d\d-\d\d", r["date"]):
                raise GenError("%s: bad date" % truth["case"])
        last_printed = [r for r in seq if r["balance"] is not None]
        if last_printed and a["closing_balance"] is not None and \
                round(last_printed[-1]["balance"] * 100) != round(a["closing_balance"] * 100):
            raise GenError("%s: closing balance is not the last printed balance" % truth["case"])
    if truth["row_count"] != len(truth["rows"]):
        raise GenError("%s: row_count" % truth["case"])


def truth_doc(case, generator, note, bank, layout, product, fmt, code, acct, features,
              newest, opening, closing, rows, accounts=None):
    t = {
        "case": case, "generator": generator, "note": note, "bank": bank, "layout": layout,
        "product": product, "source_format": fmt, "account_bank_code": code,
        "account_number": acct, "account_redaction": None, "features": features,
        "row_order": "newest_first" if newest else "oldest_first",
        "opening_balance": None if accounts else opening,
        "closing_balance": None if accounts else closing,
        "removed_rows": 0, "row_count": len(rows), "rows": rows,
    }
    if accounts:
        t["accounts"] = accounts
    verify_truth(t)
    return t


def save_xlsx(wb, path):
    wb.properties.creator = "scratchpad/qvf/gen"
    wb.properties.created = dt.datetime(2026, 1, 1)
    wb.properties.modified = dt.datetime(2026, 1, 1)
    buf = io.BytesIO()
    wb.save(buf)
    ML.fixed_zip(buf.getvalue(), path)


def write_json(path, obj):
    ML.write_json(path, obj)


def to_xls(xlsx_path, out_dir):
    """Re-save an .xlsx as a legacy BIFF8 .xls (the shape Qlik's `biff` reads):
    same cells, same number formats (so an OD balance is still a positive number
    printed with " OD"). Uses xlwt, installed into scratchpad/qvf/.pylib (dev-time
    only). Returns the .xls path, or None when xlwt is not available."""
    sys.path.insert(0, os.path.join(QVF, ".pylib"))
    try:
        import xlwt
    except ImportError:
        return None
    import openpyxl
    src = openpyxl.load_workbook(xlsx_path)
    book = xlwt.Workbook()
    styles = {}

    def style(fmt):
        if fmt not in styles:
            st = xlwt.XFStyle()
            st.num_format_str = fmt
            styles[fmt] = st
        return styles[fmt]
    for ws in src.worksheets:
        sh = book.add_sheet(ws.title)
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                v = c.value
                if isinstance(v, (dt.datetime, dt.date)):
                    sh.write(c.row - 1, c.column - 1, v, style(c.number_format or "dd/mm/yyyy"))
                elif isinstance(v, (int, float)) and not isinstance(v, bool):
                    sh.write(c.row - 1, c.column - 1, v, style(c.number_format or "General"))
                else:
                    sh.write(c.row - 1, c.column - 1, v)
    os.makedirs(out_dir, exist_ok=True)
    p = os.path.join(out_dir, os.path.splitext(os.path.basename(xlsx_path))[0] + ".xls")
    book.save(p)
    return p


def period(y, m, d=1, months=1):
    start = dt.date(y, m, d)
    mm = m - 1 + months
    end = dt.date(y + mm // 12, mm % 12 + 1, d) - dt.timedelta(days=1)
    return start, end
